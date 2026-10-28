# -*- coding: utf-8 -*-
"""ДС: должность в родительном падеже при вводе, подпись в именительном, названия юрлиц в файле, пересоздание на актуальном шаблоне (14.10.2026)."""
import io
from datetime import date
from types import SimpleNamespace

import pytest
from fastapi import HTTPException

import app.model_registry  # noqa: F401
from app.routers import annexes as A
from app.sales import annex, annex_docx, rugram


def test_position_typed_in_genitive_is_not_declined_twice():
    """Владелец 14.10.2026: «Генерального директоры» — должность, введённая как звучит в «в лице …», склонялась повторно."""
    assert rugram.gen_position("Генерального директора") == "Генерального директора"
    assert rugram.gen_position("Генеральный директор") == "Генерального директора"      # обычный ввод по-прежнему склоняется
    assert rugram.gen_position("Исполнительный директор по развитию") == "Исполнительного директора по развитию"


def test_signature_position_is_nominative_even_if_the_card_holds_genitive():
    assert rugram.nom_position("Генерального директора") == "Генеральный директор"
    assert rugram.nom_position("Генеральный директор") == "Генеральный директор"
    assert rugram.nom_position("Главного руководителя проекта") == "Главный руководитель проекта"      # зависимое слово не трогаем
    assert rugram.nom_position("Первого заместителя генерального директора") == "Первый заместитель генерального директора"
    assert rugram.nom_position(None) is None


def test_party_gives_both_forms_and_the_abbreviated_name():
    cp = SimpleNamespace(id=1, name="ГРИНДЕКС РУС ООО", inn="1", kpp=None, address=None, director_name="Кондеев Дмитрий Петрович",
                         signer_position="Генерального директора", signer_basis="Устава")
    p = annex.party(cp, cp)          # подписант — отдельным аргументом (28.10.2026); у образца поля те же
    assert p["position_gen"] == "Генерального директора" and p["position_nom"] == "Генеральный директор"
    assert p["legal_name"] == "ООО «ГРИНДЕКС РУС»"
    assert "должность подписанта" not in p["missing"]


def test_word_file_uses_abbreviated_names_and_a_nominative_signature():
    from docx import Document
    party = {"name": "ГРИНДЕКС РУС ООО", "legal_name": "ООО «ГРИНДЕКС РУС»", "position": "Генерального директора",
             "position_nom": "Генеральный директор", "position_gen": "Генерального директора", "director_gen": "Кондеева Дмитрия Петровича",
             "basis_gen": "Устава", "acting": "действующего", "short_fio": "Кондеев Д.П."}
    doc = {"number": "Приложение № 1", "contract": {"number": "Т-1", "date": date(2026, 1, 1)}, "date": date(2026, 9, 30),
           "period_from": date(2026, 10, 1), "period_to": date(2026, 10, 31), "executor": dict(party, legal_name="ООО «МЫ»", name="МЫ ООО"),
           "customer": party, "body": "Текст.", "rows": [], "vat_rate": 22, "amount_words": "x", "vat_words": "y"}
    d = Document(io.BytesIO(annex_docx.build_docx(doc)))
    text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "ООО «ГРИНДЕКС РУС», именуемое в дальнейшем «Заказчик», в лице Генерального директора Кондеева Дмитрия Петровича" in text
    assert "Генерального директоры" not in text
    sign = d.tables[-1].rows[0].cells[1].text
    assert "ООО «ГРИНДЕКС РУС»" in sign and "Генеральный директор" in sign and "Генерального директора" not in sign


# ── пересоздание на актуальном шаблоне ───────────────────────────────────────────────────────

def _rebuild_with(monkeypatch, annex_row, *, frozen=None, sent=False):
    monkeypatch.setattr(A, "_annex", lambda db, i, *a: annex_row)
    monkeypatch.setattr(A, "_contract", lambda db, i: SimpleNamespace(id=2))
    monkeypatch.setattr(A, "_freeze_for_confirm", lambda db, a, c, no: frozen or {"template": {"version_id": 9}, "amount": 1.0})
    monkeypatch.setattr(A, "_sent_to_onec", lambda a: sent)
    monkeypatch.setattr(A, "log_action", lambda *a, **k: None)
    monkeypatch.setattr(A, "_out", lambda a: {"id": a.id})
    monkeypatch.setattr(A.tpl_doc, "version_row", lambda db, vid: SimpleNamespace(version=vid) if vid else None)
    db = SimpleNamespace(commit=lambda: None)
    return A.rebuild_annex(1, db, SimpleNamespace(id=1))


def test_rebuild_replaces_the_frozen_document_and_the_version(monkeypatch):
    a = SimpleNamespace(id=1, no=5, number="Приложение № 5", contract_id=2, template_version_id=None, blocks=None)
    _rebuild_with(monkeypatch, a)
    assert a.template_version_id == 9 and a.blocks["template"]["version_id"] == 9 and a.no == 5      # номер сохраняется


def test_rebuild_refuses_a_draft_and_one_already_sent_to_onec(monkeypatch):
    draft = SimpleNamespace(id=1, no=None, number=None, contract_id=2, template_version_id=None, blocks=None)
    with pytest.raises(HTTPException) as e:
        _rebuild_with(monkeypatch, draft)
    assert e.value.status_code == 400
    sent = SimpleNamespace(id=1, no=5, number="Приложение № 5", contract_id=2, template_version_id=3, blocks={"old": 1})
    with pytest.raises(HTTPException) as e2:
        _rebuild_with(monkeypatch, sent, sent=True)
    assert e2.value.status_code == 409 and sent.blocks == {"old": 1}                  # ничего не изменилось


def test_client_file_has_the_closing_basis_page_only_when_asked():
    """«Для клиента»: ДС и основание закрывающих — одним документом на согласование; обычная выгрузка основания не содержит."""
    from docx import Document
    party = {"name": "МЫ ООО", "legal_name": "ООО «МЫ»", "position": "Генеральный директор", "position_nom": "Генеральный директор",
             "position_gen": "Генерального директора", "director_gen": "Иванова И.И.", "basis_gen": "Устава", "acting": "действующего", "short_fio": "Иванов И.И."}
    doc = {"number": "Приложение № 1", "contract": {"number": "Т-1", "date": date(2026, 1, 1)}, "date": date(2026, 9, 30), "period_from": date(2026, 10, 1),
           "period_to": date(2026, 10, 31), "executor": party, "customer": party, "body": "Текст.", "rows": [], "vat_rate": 22, "amount_words": "x",
           "vat_words": "y", "closing_basis": "Размещение рекламных материалов бренда «Икс», в сети SIMB-AD."}

    def text(data):
        d = Document(io.BytesIO(data))
        return "\n".join(p.text for p in d.paragraphs)
    assert "Основание для закрывающих документов" not in text(annex_docx.build_docx(doc))
    full = text(annex_docx.build_docx(doc, with_basis=True))
    assert "Основание для закрывающих документов" in full and "Размещение рекламных материалов бренда «Икс»" in full
    assert full.index("Подписи Сторон") < full.index("Основание для закрывающих документов")           # основание — после ДС
    assert A.file_name(no=5, doc_date=date(2026, 9, 30), brand="Икс", payer="ООО Ромашка", basis=True).endswith("с основанием для закрывающих.pdf")


# ── дата ДС: по умолчанию начало РК, смена — отдельной ручкой с записью в журнал (владелец 14.10.2026) ─────────

def _set_date_with(monkeypatch, annex_row, new, *, sent=False):
    logs = []
    monkeypatch.setattr(A, "_annex", lambda db, i, *a: annex_row)
    monkeypatch.setattr(A, "_contract", lambda db, i: SimpleNamespace(id=2))
    monkeypatch.setattr(A, "_sent_to_onec", lambda a: sent)
    monkeypatch.setattr(A, "log_action", lambda db, u, action, et, eid, details: logs.append((action, details)))
    monkeypatch.setattr(A, "_out", lambda a: {"id": a.id, "date": a.date})
    monkeypatch.setattr(A.tpl_doc, "version_row", lambda db, vid: SimpleNamespace(version=vid))
    monkeypatch.setattr(A, "_freeze_for_confirm", lambda db, a, c, no, version=None: {"date": a.date, "template": {"version_id": 3}})
    db = SimpleNamespace(commit=lambda: None)
    return A.set_annex_date(1, A.DateIn(doc_date=new), db, SimpleNamespace(id=1)), logs


def test_draft_date_changes_and_the_log_keeps_old_and_new(monkeypatch):
    a = SimpleNamespace(id=1, contract_id=2, no=None, number=None, date=date(2026, 10, 1), blocks=None, template_version_id=3)
    out, logs = _set_date_with(monkeypatch, a, date(2026, 9, 30))
    assert a.date == date(2026, 9, 30) and out["date"] == date(2026, 9, 30)
    assert logs and logs[0][0] == "annex_date_change" and "01.10.2026" in logs[0][1] and "30.09.2026" in logs[0][1]


def test_confirmed_date_change_refreezes_the_document(monkeypatch):
    a = SimpleNamespace(id=1, contract_id=2, no=5, number="Приложение № 5", date=date(2026, 10, 1), blocks={"blocks": [], "template": {"version_id": 3}},
                        template_version_id=3)
    _set_date_with(monkeypatch, a, date(2026, 10, 3))
    assert a.blocks["date"] == date(2026, 10, 3)


def test_date_cannot_change_after_onec_or_on_pre_constructor_doc(monkeypatch):
    sent = SimpleNamespace(id=1, contract_id=2, no=5, number="Приложение № 5", date=date(2026, 10, 1), blocks={"blocks": [], "template": {"version_id": 3}},
                           template_version_id=3)
    with pytest.raises(HTTPException) as e:
        _set_date_with(monkeypatch, sent, date(2026, 10, 3), sent=True)
    assert e.value.status_code == 409 and sent.date == date(2026, 10, 1)
    old = SimpleNamespace(id=1, contract_id=2, no=5, number="Приложение № 5", date=date(2026, 10, 1), blocks=None, template_version_id=None)
    with pytest.raises(HTTPException) as e2:
        _set_date_with(monkeypatch, old, date(2026, 10, 3))
    assert e2.value.status_code == 409
