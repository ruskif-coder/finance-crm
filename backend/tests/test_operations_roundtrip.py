"""
Выгрузка операций должна грузиться обратно импортом.

До 2026-08-25 не грузилась. Файл, скачанный кнопкой «Скачать» на /operations,
импорт отвергал сообщением «Worksheet named 'CF BEST' not found»: выгрузка пишет
лист «Операции», а парсер требовал лист с именем ровно «CF BEST». После
переименования листа файл падал второй раз — заголовок «НДС %» не совпадал с
ожидаемым «НДС», и обязательное поле vat_rate считалось отсутствующим. Ещё две
колонки расходились молча, без ошибки: «Документ» вместо «Ссылка на документ»
терял бы все ссылки на первичку, «НДС сумма» вместо «НДС факт» — сумму налога.

Оба формата описывают одни и те же поля, поэтому расхождение видно только тому,
кто откроет оба файла рядом. Тесты ниже держат три вещи:

  * состав выгрузки покрывает все обязательные поля импорта;
  * лист выбирается по содержимому, а не по имени;
  * лист «Инструкция» из шаблона не может быть принят за лист с данными —
    в нём есть строка с текстом «Статус», и поиск по одной этой ячейке
    (как было раньше) выбрал бы именно её.
"""
import ast
import io
from pathlib import Path

import pytest

pd = pytest.importorskip("pandas")
from openpyxl import Workbook                                    # noqa: E402

from app.routers.operations import (                             # noqa: E402
    _CF_BEST_COLUMN_MAP,
    _CF_BEST_REQUIRED_FIELDS,
    _SYNC_COMPARE_FIELDS,
    _find_import_sheet,
    _map_header_row,
    _parse_cf_best_rows,
)

OPERATIONS_PY = Path(__file__).resolve().parents[1] / "app" / "routers" / "operations.py"


def _export_columns() -> list:
    """Заголовки выгрузки — из исходника, без запуска эндпоинта.

    Эндпоинт требует сессию БД и фильтры; здесь важен только состав колонок,
    поэтому список читается из AST. Так тест ловит переименование колонки
    в момент правки, а не в момент, когда кто-то попробует загрузить файл.
    """
    tree = ast.parse(io.open(OPERATIONS_PY, encoding="utf-8").read())
    for node in ast.walk(tree):
        if not isinstance(node, ast.Assign):
            continue
        names = {t.id for t in node.targets if isinstance(t, ast.Name)}
        if "export_columns" not in names:
            continue
        return [el.elts[0].value for el in node.value.elts]
    raise AssertionError("не найден список export_columns в operations.py")


def _sheet(rows, title="Операции"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


EXPORT_HEADER = ['ID', 'Дата', 'Статус', 'Поступления', 'Списания', 'Банк', 'Период',
                 'Статья', 'Контрагент', 'ИНН', 'НДС %', 'НДС сумма', '№ ДС', 'Счет',
                 'Счет от дата', 'Документ', 'Назначение', 'Статус ДЗ']
EXPORT_ROW = [3072, None, 'ПЛАН ОПЛАТ', None, 1879.53, None, '2026-08', 'АПТЕКА',
              'ДИАЛОГ СТОЛИЦА ООО', '7722319430', 22, 338.93, None, '189',
              '2026-08-31', 'https://diadoc.kontur.ru/x', None, None]


# ---- состав выгрузки покрывает импорт ----

def test_every_export_column_is_understood_by_the_importer():
    """Ни одна колонка выгрузки не должна быть импорту незнакома.

    Незнакомая колонка не роняет импорт — она молча игнорируется. Именно так
    пропадали ссылки на документы: колонка в файле есть, глазом видна, а до базы
    не доезжает.
    """
    known = set(_CF_BEST_COLUMN_MAP)
    unknown = [c for c in _export_columns()
               if c.strip().lower() not in known
               and c != 'Статус ДЗ']                 # см. следующий тест
    assert not unknown, (
        'колонки выгрузки, которые импорт не распознаёт (данные из них потеряются): '
        f'{unknown}')


def test_receivable_status_is_never_imported():
    """«Статус ДЗ» вычисляется из даты и отсрочки.

    Принять его из файла значит позволить файлу переписать расчёт — колонка
    выгружается только для чтения человеком.
    """
    assert 'Статус ДЗ' in _export_columns()
    assert 'статус дз' not in set(_CF_BEST_COLUMN_MAP), (
        'Статус ДЗ стал читаться импортом — это решение, а не правка формата')


def test_id_is_a_hint_not_a_stored_field():
    """ID читается импортом, но подсказкой, а не данными.

    Он не входит ни в обязательные поля (в шаблоне колонки ID нет вовсе), ни в
    сравниваемые — иначе номер из чужого стенда показывался бы «изменением»
    и, что хуже, мог бы попытаться уехать в базу. Правило выбора по нему —
    в tests/test_import_matching.py.
    """
    assert 'ID' in _export_columns()
    assert _CF_BEST_COLUMN_MAP.get('id') == 'op_id'
    assert 'op_id' not in _CF_BEST_REQUIRED_FIELDS
    assert 'op_id' not in _SYNC_COMPARE_FIELDS


def test_export_covers_all_required_import_fields():
    """Обязательные поля импорта должны быть в выгрузке все до одного."""
    mapped = {_CF_BEST_COLUMN_MAP.get(c.strip().lower()) for c in _export_columns()}
    missing = _CF_BEST_REQUIRED_FIELDS - mapped
    assert not missing, f'выгрузка не даёт обязательных полей импорта: {sorted(missing)}'


def test_inn_travels_with_the_export():
    """Без ИНН импорт сопоставляет контрагента по имени и плодит дубли."""
    mapped = {_CF_BEST_COLUMN_MAP.get(c.strip().lower()) for c in _export_columns()}
    assert 'inn' in mapped


# ---- лист выбирается по содержимому ----

def test_export_sheet_name_is_accepted():
    """Лист «Операции» — именно то имя, на котором импорт падал."""
    rows = _parse_cf_best_rows(_sheet([EXPORT_HEADER, EXPORT_ROW], title='Операции'))
    assert len(rows) == 1
    r = rows[0]
    assert r['expense'] == 1879.53
    assert r['vat_rate'] == 22
    assert r['counterparty'] == 'ДИАЛОГ СТОЛИЦА ООО'
    assert r['inn'] == '7722319430'
    assert r['document_link'] == 'https://diadoc.kontur.ru/x'   # «Документ» → ссылка


def test_any_sheet_name_works_if_the_header_is_there():
    rows = _parse_cf_best_rows(_sheet([EXPORT_HEADER, EXPORT_ROW], title='Лист1'))
    assert len(rows) == 1


def test_instruction_sheet_is_not_mistaken_for_data():
    """Ловушка шаблона: на листе «Инструкция» есть строка со словом «Статус».

    Прежний поиск («первая строка, где встречается ячейка статус») выбрал бы
    именно её, и импорт разобрал бы описание колонок как операции.

    Лист с данными назван «Операции», а не «CF BEST», намеренно: имя «CF BEST»
    проверяется первым и увело бы выбор мимо ловушки — тест остался бы зелёным
    при любой наивной эвристике. Проверено мутацией.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = 'Инструкция'
    ws.append(['Колонка', 'Что писать'])
    ws.append(['Дата', 'дата операции'])
    ws.append(['Статус', 'выберите из списка'])
    data = wb.create_sheet('Операции')
    data.append(EXPORT_HEADER)
    data.append(EXPORT_ROW)
    buf = io.BytesIO()
    wb.save(buf)

    raw, header_idx, col_map = _find_import_sheet(buf.getvalue())
    assert set(col_map.values()) >= _CF_BEST_REQUIRED_FIELDS
    assert str(raw.iloc[header_idx].tolist()[2]).strip() == 'Статус'
    rows = _parse_cf_best_rows(buf.getvalue())
    assert len(rows) == 1 and rows[0]['status'] == 'ПЛАН ОПЛАТ'


def test_missing_column_names_the_sheet_and_the_row():
    """Отказ должен говорить, где искать, а не «нет листа CF BEST»."""
    header = [c for c in EXPORT_HEADER if c not in ('НДС %',)]
    row = [v for c, v in zip(EXPORT_HEADER, EXPORT_ROW) if c not in ('НДС %',)]
    with pytest.raises(ValueError) as e:
        _parse_cf_best_rows(_sheet([header, row], title='Операции'))
    msg = str(e.value)
    assert 'Операции' in msg and 'vat_rate' in msg


def test_no_recognisable_header_at_all():
    with pytest.raises(ValueError):
        _parse_cf_best_rows(_sheet([['раз', 'два'], [1, 2]], title='Лист1'))


# ---- два заголовка на одно поле ----

def test_duplicate_headers_for_one_field_do_not_break_parsing():
    """Файл может нести и «НДС», и «НДС %» — оба ведут в vat_rate.

    Без правила «первое вхождение выигрывает» в выборку попали бы две колонки
    с одним именем, и pandas отдал бы DataFrame вместо колонки.
    """
    header = ['Дата', 'Статус', 'Поступления', 'Списания', 'Банк', 'Период',
              'НДС', 'НДС %', 'Статья', 'Контрагент', '№ ДС', 'Счет',
              'Счет от дата', 'Назначение']
    row = [None, 'ОПЛАЧЕНО', None, 1000, 'АльфаБанк', '2026-08', 20, 22,
           'IT', 'ООО РОМАШКА', None, '5', '2026-08-01', None]
    mapped = _map_header_row(pd.Series(header))
    assert list(mapped.values()).count('vat_rate') == 1
    rows = _parse_cf_best_rows(_sheet([header, row], title='Операции'))
    assert rows[0]['vat_rate'] == 20        # выигрывает первая колонка, «НДС»
