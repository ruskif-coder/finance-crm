"""«План продаж» (владелец 24.10.2026, docs/ПЛАН_план_продаж.md).

Выполнение = факт + в работе (сделки со «Сборки»), прогноз = + планируемые: зелёный 100 %, оранжевый 50 %, серый 0.
Разрезы независимые; сделки без услуги / продажника — отдельными строками «без …». Цели правятся пачкой, каждая правка —
в журнал «было → стало»; «свои» видят общую цель и свою строку; ночной снимок даёт «▲▼ за 7 дней».
"""
from datetime import date, timedelta
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from sqlalchemy import text

from app.database import SessionLocal
from app.sales import plan

YEAR = 2091          # год, в котором на стенде нет ни сделок, ни целей — тесты владеют им целиком


def _d(month, layer, light=None, amount=100.0, service=1, rep=10, stage_key="booking"):
    return SimpleNamespace(month=month, layer=layer, light=light, amount=amount, service_id=service,
                           sales_rep_id=rep, stage_key=stage_key)


# ── расчёт ──────────────────────────────────────────────────────────────────
def test_layers_and_lights_go_to_their_buckets():
    deals = [_d(1, "фактические", amount=300), _d(1, "реализуемые", amount=200),
             _d(1, "планируемые", "green", 100), _d(1, "планируемые", "orange", 80), _d(1, "планируемые", "grey", 50)]
    cells = plan.aggregate(deals)
    c = cells[("общий", 0)][0]
    assert (c["fact"], c["work"], c["green"], c["orange"], c["grey"]) == (300, 200, 100, 80, 50)
    assert plan.done(c) == 500
    assert plan.forecast(c) == 300 + 200 + 100 + 40, "оранжевый — 50 %, серый в прогноз не идёт"


def test_deals_without_service_or_rep_get_their_own_rows():
    cells = plan.aggregate([_d(3, "фактические", service=None, rep=None, amount=70), _d(3, "фактические", amount=30)])
    assert cells[("без услуги", 0)][2]["fact"] == 70
    assert cells[("без продажника", 0)][2]["fact"] == 70
    assert cells[("услуга", 1)][2]["fact"] == 30 and cells[("продажник", 10)][2]["fact"] == 30
    assert cells[("общий", 0)][2]["fact"] == 100


def test_planned_without_light_takes_it_from_the_stage():
    """Пустой светофор — по стадии, одним правилом с дашбордом (`probability.effective`): Бронь — оранжевый."""
    cells = plan.aggregate([_d(5, "планируемые", None, 100, stage_key="booking"),
                            _d(5, "планируемые", None, 40, stage_key="media_plan")])
    c = cells[("общий", 0)][4]
    assert c["orange"] == 100 and c["grey"] == 40


# ── цели: сохранение и журнал ───────────────────────────────────────────────
@pytest.fixture
def db():
    s = SessionLocal()
    _purge(s)
    yield s
    _purge(s)
    s.close()


def _purge(s):
    s.rollback()
    s.execute(text("DELETE FROM sales_targets WHERE year = :y"), {"y": YEAR})
    s.execute(text("DELETE FROM sales_plan_snapshots WHERE year = :y"), {"y": YEAR})
    s.execute(text("DELETE FROM audit_log WHERE action = 'sales_target_set' AND details LIKE :p"), {"p": f"{YEAR}-%"})
    s.commit()


_ADMIN = SimpleNamespace(role=SimpleNamespace(key="admin"), id=None, name="тест", role_id=None)


def _rep_and_service(s):
    rep = s.execute(text("SELECT id FROM sales_reps ORDER BY id LIMIT 1")).scalar()
    svc = s.execute(text("SELECT id FROM sales_services WHERE is_active ORDER BY id LIMIT 1")).scalar()
    if not (rep and svc):
        pytest.skip("на стенде нет продажника или услуги")
    return rep, svc


def test_targets_are_saved_logged_and_cleared(db):
    rep, svc = _rep_and_service(db)
    items = [plan.TargetItem(month=1, scope="общий", amount=1_000_000),
             plan.TargetItem(month=1, scope="услуга", key_id=svc, amount=400_000),
             plan.TargetItem(month=1, scope="продажник", key_id=rep, amount=300_000)]
    assert plan.save_targets(db, YEAR, items, _ADMIN) == 3
    assert plan.save_targets(db, YEAR, [plan.TargetItem(month=1, scope="общий", amount=1_200_000)], _ADMIN) == 1
    assert plan.save_targets(db, YEAR, [plan.TargetItem(month=1, scope="общий", amount=1_200_000)], _ADMIN) == 0, \
        "без изменений — ни записи, ни строки журнала"
    hist = plan.history(db, YEAR, 1, 1, "общий", 0)
    assert [h["was"] for h in hist] == [1_000_000, None] and [h["now"] for h in hist] == [1_200_000, 1_000_000]
    plan.save_targets(db, YEAR, [plan.TargetItem(month=1, scope="продажник", key_id=rep, amount=None)], _ADMIN)
    goals = {(t.scope, t.sales_rep_id) for t in db.execute(text(
        "SELECT scope, sales_rep_id FROM sales_targets WHERE year = :y"), {"y": YEAR}).all()}
    assert ("продажник", rep) not in goals, "пустое значение снимает цель"
    assert plan.history(db, YEAR, 1, 1, "продажник", rep)[0]["now"] is None


def test_bad_targets_are_refused(db):
    rep, svc = _rep_and_service(db)
    for bad in (plan.TargetItem(month=13, scope="общий", amount=1),
                plan.TargetItem(month=1, scope="общий", amount=-5),
                plan.TargetItem(month=1, scope="услуга", amount=5),                  # без услуги
                plan.TargetItem(month=1, scope="продажник", key_id=10 ** 8, amount=5),
                plan.TargetItem(month=1, scope="бренд", amount=5)):
        with pytest.raises(HTTPException) as e:
            plan.save_targets(db, YEAR, [bad], _ADMIN)
        assert e.value.status_code in (400, 404)
    assert not db.execute(text("SELECT count(*) FROM sales_targets WHERE year = :y"), {"y": YEAR}).scalar()


# ── экран: видимость и строки ───────────────────────────────────────────────
def test_own_scope_sees_total_and_own_rep_only(db):
    rep, svc = _rep_and_service(db)
    plan.save_targets(db, YEAR, [plan.TargetItem(month=2, scope="общий", amount=900),
                                 plan.TargetItem(month=2, scope="услуга", key_id=svc, amount=500),
                                 plan.TargetItem(month=2, scope="продажник", key_id=rep, amount=400)], _ADMIN)
    full = plan.screen(db, YEAR, own=None)
    assert {r["scope"] for r in full["rows"]} >= {"общий", "услуга", "продажник"}
    mine = plan.screen(db, YEAR, own=[rep])
    assert {r["scope"] for r in mine["rows"]} == {"общий", "продажник"}
    assert {r["key_id"] for r in mine["rows"] if r["scope"] == "продажник"} == {rep}
    tot = next(r for r in full["rows"] if r["scope"] == "общий")
    assert tot["months"][1]["goal"] == 900 and tot["months"][0]["goal"] is None, "нет цели — None, не ноль"


# ── ночной снимок ───────────────────────────────────────────────────────────
def test_snapshot_is_idempotent_prunes_and_gives_week_delta(db):
    today = date(YEAR, 6, 15)
    db.execute(text("INSERT INTO sales_plan_snapshots (day, year, month, scope, key_id, done, forecast) "
                    "VALUES (:d, :y, 6, 'общий', 0, 10, 100), (:old, :y, 6, 'общий', 0, 1, 1)"),
               {"d": today - timedelta(days=7), "old": today - timedelta(days=120), "y": YEAR})
    db.commit()
    cells = {("общий", 0): [plan.empty() for _ in range(12)]}
    cells[("общий", 0)][5].update(fact=50, work=20, green=30)
    assert plan.write_snapshot(db, YEAR, cells, today) == 12
    assert plan.write_snapshot(db, YEAR, cells, today) == 12, "повторный прогон за тот же день — перезапись, не дубль"
    n_old = db.execute(text("SELECT count(*) FROM sales_plan_snapshots WHERE year = :y AND day = :d"),
                       {"y": YEAR, "d": today - timedelta(days=120)}).scalar()
    assert n_old == 0, "старше 90 дней — убирается"
    week = plan.week_ago(db, YEAR, today)
    assert week[("общий", 0)][5] == 100


# ── ручки и права ───────────────────────────────────────────────────────────
def test_routes_need_the_right_and_answer(db):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.routers.auth import get_current_user

    class _U:
        def __init__(self, key):
            self.id, self.name, self.email, self.role_id = None, "т", "t@t", -1
            self.role, self.consent_accepted_at = SimpleNamespace(key=key), True
    with TestClient(app) as c:
        try:
            app.dependency_overrides[get_current_user] = lambda: _U("no_such_role")
            assert c.get(f"/api/sales-plan?year={YEAR}").status_code == 403
            assert c.put("/api/sales-plan/targets", json={"year": YEAR, "items": []}).status_code == 403
            app.dependency_overrides[get_current_user] = lambda: _U("admin")
            r = c.get(f"/api/sales-plan?year={YEAR}")
            assert r.status_code == 200 and r.json()["can_edit"] is True
            assert c.get(f"/api/sales-plan/cell?year={YEAR}&m_from=5&m_to=2&scope=общий").status_code == 400
            assert c.get("/api/sales-plan?year=1990").status_code == 400
        finally:
            app.dependency_overrides.clear()


# ── Excel: выгрузка и загрузка целей ────────────────────────────────────────
def test_excel_roundtrip_previews_only_real_changes(db):
    import io
    import openpyxl
    from app.sales import plan_xlsx
    rep, svc = _rep_and_service(db)
    plan.save_targets(db, YEAR, [plan.TargetItem(month=1, scope="общий", amount=1_000_000),
                                 plan.TargetItem(month=2, scope="продажник", key_id=rep, amount=500_000)], _ADMIN)
    body = plan_xlsx.export(db, YEAR, own=None)
    wb = openpyxl.load_workbook(io.BytesIO(body))
    assert wb.sheetnames[:2] == ["Выполнение", "Цели"], "отчёт первым, шаблон правки вторым"
    run = wb["Выполнение"]
    assert str(run["A1"].value).startswith(f"План продаж {YEAR}"), "у отчёта есть шапка, это не голая таблица"
    assert any(c.fill.fgColor.rgb not in (None, "00000000") for row in run.iter_rows() for c in row), "проценты подсвечены"
    ws = wb["Цели"]
    hr = plan_xlsx.header_row(ws)
    head = [c.value for c in ws[hr]]
    assert head[:4] == ["Код", "Разрез", "Название", "янв"] and head[15] == "Год"
    rows = {r[0].value: r for r in ws.iter_rows(min_row=hr + 1) if r[0].value}
    assert rows["общий:0"][3].value == 1_000_000
    rows["общий:0"][3].value = 1_100_000             # было 1 000 000 → станет 1 100 000
    rows[f"продажник:{rep}"][4].value = None         # пусто — не менять
    rows[f"услуга:{svc}"][5].value = 250_000         # новая цель
    ws.append(["продажник:999999999", "продажник", "нет такого", 5])
    buf = io.BytesIO()
    wb.save(buf)
    out = plan_xlsx.preview(db, YEAR, buf.getvalue())
    ch = {(c["scope"], c["key_id"], c["month"]): (c["was"], c["now"]) for c in out["changes"]}
    assert ch == {("общий", 0, 1): (1_000_000, 1_100_000), ("услуга", svc, 3): (None, 250_000)}
    assert out["errors"] and "999999999" in out["errors"][0]


def test_excel_import_refuses_garbage(db):
    from app.sales import plan_xlsx
    with pytest.raises(HTTPException) as e:
        plan_xlsx.preview(db, YEAR, b"not an excel file")
    assert e.value.status_code == 400


def test_inactive_reps_go_last_or_hide_and_active_sellers_always_show(db):
    """Владелец 24.10.2026: неактивные — последними, если за год есть выполненное; без него — скрыты (даже с целью).
    Активные продавцы — строкой всегда, чтобы им можно было поставить цель."""
    inactive = db.execute(text("""SELECT r.id FROM sales_reps r LEFT JOIN users u ON u.id = r.user_id
                                   WHERE NOT r.is_active OR u.is_active = 0 ORDER BY r.id LIMIT 1""")).scalar()
    seller = db.execute(text("""SELECT r.id FROM sales_reps r JOIN users u ON u.id = r.user_id JOIN roles ro ON ro.id = u.role_id
                                 WHERE r.is_active AND u.is_active = 1 AND ro.staff_group = 'seller' ORDER BY r.id LIMIT 1""")).scalar()
    if not (inactive and seller):
        pytest.skip("на стенде нет неактивного продажника или активного продавца")
    plan.save_targets(db, YEAR, [plan.TargetItem(month=1, scope="продажник", key_id=inactive, amount=100)], _ADMIN)
    rows = plan.screen(db, YEAR)["rows"]
    reps = [r for r in rows if r["scope"] == "продажник"]
    assert seller in {r["key_id"] for r in reps}, "активный продавец без сделок — строкой"
    assert inactive not in {r["key_id"] for r in reps}, "неактивный без выполненного за год — скрыт, даже с целью"
    # с выполненным — последним среди продажников и с пометкой
    cells = {("продажник", inactive): [dict(plan.empty(), fact=50.0)] + [plan.empty() for _ in range(11)]}
    rows2 = plan.rows_for(db, YEAR, cells, plan.targets_of(db, YEAR), {}, own=None)
    reps2 = [r for r in rows2 if r["scope"] == "продажник"]
    assert reps2[-1]["key_id"] == inactive and reps2[-1]["inactive"] is True
