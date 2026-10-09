# -*- coding: utf-8 -*-
"""«Паблишеры → Аудитория» (владелец 09.10.2026): заявленные площадкой цифры (MAU/DAU/соцдем,
с источником и датой) рядом с тем, что намерили мы (показы нашего счётчика, верификатор,
уники Adfox, fill DSP) за окно в N дней.

Два правила, которые держат эти тесты:
  * заявленное — строки с датой и источником, а не колонки: новый замер не затирает старый;
  * «подтверждено нами» складывает ТОЛЬКО наш счётчик (stat_sources.OWN); верификатор —
    отдельное число, никогда не в сумме.
"""
from datetime import date, timedelta

import pytest
from sqlalchemy import text

from app.database import SessionLocal
from app.sales import publisher_audience as A


@pytest.fixture
def db():
    s = SessionLocal()
    try:
        yield s
    finally:
        s.rollback()
        s.close()


def test_permission_section_exists():
    from app.permissions import SECTIONS
    sec = next(s for s in SECTIONS if s["key"] == "dir_publishers_audience")
    assert sec["group"] == "Паблишеры"
    assert set(sec["actions"]) == {"view", "edit"}


def test_metric_catalog_is_closed():
    assert "mau" in A.METRICS and "dau" in A.METRICS
    with pytest.raises(ValueError):
        A.check_metric("visits_per_second")


def test_measured_separates_own_counter_from_verifier(db):
    """Берём любую площадку, у которой за окно есть и наши показы, и показы Weborama,
    и проверяем, что в `shows` вошли только наши."""
    day_to = date.today()
    day_from = day_to - timedelta(days=30)
    rows = db.execute(text("""
        SELECT p.publisher_id,
               sum(s.shows) FILTER (WHERE s.source = ANY(:own)) AS own,
               sum(s.shows) FILTER (WHERE s.source = ANY(:ver)) AS ver
          FROM ad_campaign_stat s JOIN ad_campaign_placement p ON p.id = s.placement_id
         WHERE s.date BETWEEN :a AND :b
         GROUP BY p.publisher_id
        HAVING sum(s.shows) FILTER (WHERE s.source = ANY(:ver)) > 0
    """), {"own": list(A.OWN), "ver": list(A.VERIFIER), "a": day_from, "b": day_to}).all()
    if not rows:
        pytest.skip("на стенде нет площадки с показами верификатора за окно")
    pid, own, ver = rows[0]
    got = A.measured(db, day_from, day_to)[pid]
    assert got["shows"] == int(own or 0)
    assert got["verifier_shows"] == int(ver)
    assert got["shows"] != got["shows"] + got["verifier_shows"]


def test_overview_lists_every_active_publisher_even_without_data(db):
    out = A.overview(db, days=30)
    active = db.execute(text("SELECT count(*) FROM sales_publishers WHERE is_active")).scalar()
    assert len(out["publishers"]) == active
    empty = [p for p in out["publishers"] if p["measured"]["shows"] is None and not p["declared"]]
    for p in empty:
        assert p["gap_pct"] is None


def test_declared_takes_latest_measurement_per_metric(db):
    pid = db.execute(text("SELECT id FROM sales_publishers WHERE is_active LIMIT 1")).scalar()
    db.execute(text("DELETE FROM sales_publisher_audience WHERE publisher_id = :p AND note = 'test'"), {"p": pid})
    for v, d in ((1_000_000, date.today() - timedelta(days=2)), (2_000_000, date.today() - timedelta(days=1))):
        A.add_measure(db, pid, metric="mau", value=v, source="тест", measured_at=d, note="test",
                      surface_kind=None, segment=None, user_id=None)
    db.flush()
    dec = A.declared(db, [pid])[pid]
    assert dec["mau"]["value"] == 2_000_000
    assert dec["mau"]["measured_at"] == date.today() - timedelta(days=1)
    assert dec["mau"]["source"] == "тест"
    db.rollback()


# ── ручки ─────────────────────────────────────────────────────────────────────
from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402
from app.routers.auth import get_current_user  # noqa: E402


class _Role:
    def __init__(self, key):
        self.key, self.id = key, None


class _User:
    def __init__(self, key):
        self.id, self.name, self.email = None, "test", "t@t"
        self.role, self.role_id, self.consent_accepted_at = _Role(key), None, True


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _as(key):
    app.dependency_overrides[get_current_user] = lambda: _User(key)


def test_admin_sees_overview_and_sources(client):
    _as("admin")
    r = client.get("/api/publisher-audience?days=30")
    assert r.status_code == 200
    body = r.json()
    assert {"window", "publishers", "totals", "metrics"} <= set(body)
    assert body["window"]["days"] == 30
    assert client.get("/api/publisher-audience/sources").json()[:1] != []


def test_role_without_the_right_is_refused(client):
    _as("sales")
    assert client.get("/api/publisher-audience").status_code == 403


def test_window_is_bounded(client):
    _as("admin")
    assert client.get("/api/publisher-audience?days=0").status_code == 400
    assert client.get("/api/publisher-audience?days=366").status_code == 400


def test_measure_validation_answers_4xx_with_a_reason(client, db):
    _as("admin")
    pid = db.execute(text("SELECT id FROM sales_publishers WHERE is_active LIMIT 1")).scalar()
    base = {"metric": "mau", "value": 1000, "source": "тест", "measured_at": str(date.today())}
    for bad in ({"metric": "visits_per_second"}, {"value": -1}, {"source": "  "},
                {"measured_at": str(date.today() + timedelta(days=1))}, {"surface_kind": "tv"}):
        r = client.post(f"/api/publisher-audience/{pid}/measures", json={**base, **bad})
        assert r.status_code == 400, bad
        assert r.json()["detail"]
    assert client.post("/api/publisher-audience/999999/measures", json=base).status_code == 404


def test_measure_roundtrip_is_logged_and_removable(client, db):
    _as("admin")
    pid = db.execute(text("SELECT id FROM sales_publishers WHERE is_active LIMIT 1")).scalar()
    r = client.post(f"/api/publisher-audience/{pid}/measures",
                    json={"metric": "dau", "value": 4321, "source": "тест-ручка",
                          "measured_at": str(date.today()), "note": "test"})
    assert r.status_code == 200
    mid = r.json()["id"]
    try:
        card = client.get(f"/api/publisher-audience/{pid}").json()
        assert any(h["id"] == mid for h in card["history"])
        assert card["declared"]["dau"]["source"] == "тест-ручка"
        n = db.execute(text("SELECT count(*) FROM audit_log WHERE action = 'publisher_audience_add' "
                            "AND entity_id = :p AND details LIKE '%тест-ручка%'"), {"p": pid}).scalar()
        assert n >= 1
    finally:
        assert client.delete(f"/api/publisher-audience/measures/{mid}").status_code == 200
    assert client.delete(f"/api/publisher-audience/measures/{mid}").status_code == 404
    left = db.execute(text("SELECT count(*) FROM sales_publisher_audience WHERE id = :m"), {"m": mid}).scalar()
    assert left == 0


def test_gap_and_stickiness_follow_the_formulas():
    dec = {"mau": {"value": 1_000_000}, "dau": {"value": 120_000}}
    assert A._gap(dec, {"uniques": 250_000}) == 25.0
    assert A._gap(dec, {"uniques": None}) is None
    assert A._gap({}, {"uniques": 5}) is None


# ── сетка, пачка, Excel ───────────────────────────────────────────────────────
import io  # noqa: E402

SRC = "тест-пачка"


def _cleanup(db):
    db.execute(text("DELETE FROM sales_publisher_audience WHERE source = :s"), {"s": SRC})
    db.execute(text("DELETE FROM audit_log WHERE details LIKE :d"), {"d": f"%{SRC}%"})
    db.commit()


def _one_active(db):
    pid = db.execute(text("SELECT id FROM sales_publishers WHERE is_active AND status = 'СОТРУДНИЧАЕМ' ORDER BY id LIMIT 1")).scalar()
    if not pid:
        pytest.skip("нет активной площадки")
    return pid


def test_field_catalog_is_consistent():
    keys = [f[0] for f in A.FIELDS]
    assert len(set(keys)) == len(keys) and len({f[1] for f in A.FIELDS}) == len(A.FIELDS)
    for key, _label, metric, surface, segment in A.FIELDS:
        assert metric in A.METRICS
        assert key == (f"{metric}:{segment}" if segment else metric) + (f"@{surface}" if surface else "")


def test_grid_lists_active_publishers_with_all_fields(client, db):
    _as("admin")
    body = client.get("/api/publisher-audience/grid").json()
    assert [f["key"] for f in body["fields"]] == [f[0] for f in A.FIELDS] + [f[0] for f in A.SYS_FIELDS]
    assert {f["group"] for f in body["fields"]} == {"audience", "system"}
    active = db.execute(text("SELECT count(*) FROM sales_publishers WHERE is_active AND status = 'СОТРУДНИЧАЕМ'")).scalar()
    assert len(body["publishers"]) == active


def test_bulk_writes_once_and_skips_the_same_value(client, db):
    _as("admin")
    pid = _one_active(db)
    try:
        payload = {"source": SRC, "measured_at": str(date.today()),
                   "items": [{"publisher_id": pid, "key": "wau", "value": 777}, {"publisher_id": pid, "key": "share:Ж", "value": 61.5}]}
        r1 = client.post("/api/publisher-audience/bulk", json=payload).json()
        assert r1["written"] == 2 and r1["skipped_same"] == 0
        r2 = client.post("/api/publisher-audience/bulk", json=payload).json()
        assert r2["written"] == 0 and r2["skipped_same"] == 2          # повтор не плодит дубли
        grid = client.get("/api/publisher-audience/grid").json()
        row = next(p for p in grid["publishers"] if p["id"] == pid)
        assert row["values"]["wau"]["value"] == 777 and row["values"]["wau"]["source"] == SRC
    finally:
        _cleanup(db)


def test_bulk_refuses_bad_input_with_a_reason(client, db):
    _as("admin")
    pid = _one_active(db)
    base = {"source": SRC, "measured_at": str(date.today())}
    cases = [({"items": [{"publisher_id": pid, "key": "wau", "value": -1}]}, 400),
             ({"items": [{"publisher_id": pid, "key": "share:Ж", "value": 101}]}, 400),
             ({"items": [{"publisher_id": pid, "key": "нет_такого", "value": 1}]}, 400),
             ({"items": [{"publisher_id": 999999, "key": "wau", "value": 1}]}, 404),
             ({"items": []}, 400), ({"items": [{"publisher_id": pid, "key": "wau", "value": 1}], "source": "  "}, 400)]
    for extra, code in cases:
        r = client.post("/api/publisher-audience/bulk", json={**base, **extra})
        assert r.status_code == code, extra
        assert r.json()["detail"]
    left = db.execute(text("SELECT count(*) FROM sales_publisher_audience WHERE source = :s"), {"s": SRC}).scalar()
    assert left == 0


def _book_with(db, pid, cells, src=SRC):
    """Выгрузка, в которую вписаны значения для площадки pid."""
    from openpyxl import load_workbook
    ws_book = load_workbook(io.BytesIO(A.export_xlsx(db)))
    ws = ws_book["Аудитория"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == pid:
            for label, v in cells.items():
                row[head.index(label)].value = v
            row[head.index("Источник")].value = src
            break
    out = io.BytesIO()
    ws_book.save(out)
    return out.getvalue()


def test_excel_roundtrip_preview_then_apply(client, db):
    _as("admin")
    pid = _one_active(db)
    try:
        book = _book_with(db, pid, {"DAU": 4242, "Доля 25–54, %": "51,5"})
        files = {"file": ("a.xlsx", book, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        prev = client.post("/api/publisher-audience/import?apply=false", files=files).json()
        assert prev["applied"] is False and prev["written"] == 2 and prev["errors_total"] == 0
        assert db.execute(text("SELECT count(*) FROM sales_publisher_audience WHERE source = :s"), {"s": SRC}).scalar() == 0
        done = client.post("/api/publisher-audience/import?apply=true", files=files).json()
        assert done["written"] == 2
        assert db.execute(text("SELECT count(*) FROM sales_publisher_audience WHERE source = :s"), {"s": SRC}).scalar() == 2
        again = client.post("/api/publisher-audience/import?apply=true", files=files).json()
        assert again["written"] == 0 and again["errors_total"] == 0     # повтор того же файла ничего не пишет
    finally:
        _cleanup(db)


def test_excel_rejects_rows_without_source_and_bad_numbers(client, db):
    _as("admin")
    pid = _one_active(db)
    book = _book_with(db, pid, {"MAU": "много"}, src="")
    files = {"file": ("a.xlsx", book, "application/octet-stream")}
    out = client.post("/api/publisher-audience/import?apply=true", files=files).json()
    assert out["written"] == 0 and out["errors_total"] >= 1
    assert any("не число" in e or "источник" in e for e in out["errors"])
    assert client.post("/api/publisher-audience/import", files={"file": ("a.txt", b"x", "text/plain")}).status_code == 400
    assert client.post("/api/publisher-audience/import", files={"file": ("a.xlsx", b"not a zip", "application/octet-stream")}).status_code == 400


def test_excel_export_is_a_workbook_with_the_catalog_columns(client):
    _as("admin")
    r = client.get("/api/publisher-audience/export.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(r.content))["Аудитория"]
    head = [c.value for c in ws[1]]
    assert head[:3] == ["ID", "Площадка", "Вид"] and head[-2:] == ["Источник", "Дата замера"]
    assert [f[1] for f in A.FIELDS] + [f[1] for f in A.SYS_FIELDS] == head[3:-2]


def test_bulk_and_import_need_the_edit_right(client):
    _as("sales")
    assert client.post("/api/publisher-audience/bulk", json={"source": "x", "measured_at": str(date.today()), "items": []}).status_code == 403
    assert client.get("/api/publisher-audience/export.xlsx").status_code == 403


# ── поля, что в системе уже есть: трафик, SimilarWeb, запросы кода, покрытие ──
def _traffic_snapshot(db, pid, scope):
    return [tuple(r) for r in db.execute(text(
        "SELECT id, value, depth, measured_at, source FROM sales_publisher_traffic WHERE publisher_id = :p AND scope = :s ORDER BY id"),
        {"p": pid, "s": scope}).all()]


def _restore_traffic(db, pid, scope, snap):
    db.execute(text("DELETE FROM sales_publisher_traffic WHERE publisher_id = :p AND scope = :s"), {"p": pid, "s": scope})
    for _id, v, d, m, src in snap:
        db.execute(text("INSERT INTO sales_publisher_traffic (id, publisher_id, scope, value, depth, measured_at, source) "
                        "VALUES (:i, :p, :s, :v, :d, :m, :src)"), {"i": _id, "p": pid, "s": scope, "v": v, "d": d, "m": m, "src": src})
    db.commit()


def test_system_catalog_points_at_real_scopes_and_surfaces():
    from app.ad.balance import REQ_SCOPE, SW_SCOPES
    from app.sales.models import TRAFFIC_SCOPES
    # карточка пишет TRAFFIC_SCOPES, балансировщик — SimilarWeb и запросы по поверхностям (тот же склад)
    known = set(TRAFFIC_SCOPES) | set(SW_SCOPES.values()) | set(REQ_SCOPE.values())
    for key, _label, kind, scope, col in A.SYS_FIELDS:
        assert key in A.SYS_BY_KEY and key not in A.FIELD_BY_KEY
        if kind == "traffic":
            assert scope in known, f"{scope} никто не читает — запись молча потеряется"
            assert col in ("value", "depth")
        else:
            assert kind == "coverage" and scope in ("web", "app")


def test_system_field_is_written_like_the_card_does_and_read_back(client, db):
    """Поле из карточки площадки: замер текущего месяца ПЕРЕЗАПИСЫВАЕТСЯ (как в карточке), прошлые
    замеры не трогаются, вторая колонка (глубина) не теряется."""
    _as("admin")
    pid = db.execute(text("SELECT publisher_id FROM sales_publisher_traffic WHERE scope = 'web' AND depth IS NOT NULL "
                          "GROUP BY publisher_id LIMIT 1")).scalar()
    if not pid:
        pytest.skip("нет площадки с замером web")
    snap = _traffic_snapshot(db, pid, "web")
    try:
        r = client.post("/api/publisher-audience/bulk", json={"source": SRC, "measured_at": str(date.today()),
                        "items": [{"publisher_id": pid, "key": "tr:web", "value": 123456}]})
        assert r.status_code == 200 and r.json()["written"] == 1
        row = next(p for p in client.get("/api/publisher-audience/grid").json()["publishers"] if p["id"] == pid)
        assert row["values"]["tr:web"]["value"] == 123456 and row["values"]["tr:web"]["source"] == A.SOURCE_SYSTEM
        last = db.execute(text("SELECT value, depth FROM sales_publisher_traffic WHERE publisher_id = :p AND scope = 'web' "
                               "ORDER BY measured_at DESC LIMIT 1"), {"p": pid}).first()
        assert last[0] == 123456 and last[1] == snap[-1][2], "глубина последнего замера потерялась"
        again = client.post("/api/publisher-audience/bulk", json={"source": SRC, "measured_at": str(date.today()),
                            "items": [{"publisher_id": pid, "key": "tr:web", "value": 123456}]}).json()
        assert again["written"] == 0 and again["skipped_same"] == 1
        assert db.execute(text("SELECT count(*) FROM sales_publisher_audience WHERE source = :s"), {"s": SRC}).scalar() == 0,             "системное поле не должно попадать в таблицу замеров аудитории"
    finally:
        _restore_traffic(db, pid, "web", snap)
        _cleanup(db)
    assert _traffic_snapshot(db, pid, "web") == snap


def test_coverage_needs_an_existing_surface_and_stays_within_100(client, db):
    _as("admin")
    row = db.execute(text("SELECT p.id FROM sales_publishers p WHERE p.is_active AND NOT EXISTS "
                          "(SELECT 1 FROM sales_publisher_surfaces s WHERE s.publisher_id = p.id AND s.kind = 'app') LIMIT 1")).first()
    if row:
        r = client.post("/api/publisher-audience/bulk", json={"source": SRC, "measured_at": str(date.today()),
                        "items": [{"publisher_id": row[0], "key": "cov:app", "value": 50}]})
        assert r.status_code == 400 and "поверхност" in r.json()["detail"]
    pid = db.execute(text("SELECT publisher_id FROM sales_publisher_surfaces WHERE kind = 'web' LIMIT 1")).scalar()
    r = client.post("/api/publisher-audience/bulk", json={"source": SRC, "measured_at": str(date.today()),
                    "items": [{"publisher_id": pid, "key": "cov:web", "value": 101}]})
    assert r.status_code == 400


def test_overview_carries_the_system_block_and_request_load(db):
    out = A.overview(db, days=30)
    assert all("system" in p and "load_pct" in p for p in out["publishers"])
    with_req = [p for p in out["publishers"] if p["load_pct"] is not None]
    for p in with_req:
        assert p["measured"]["shows"] and any(k.startswith("tr:ad_requests") for k in p["system"])
