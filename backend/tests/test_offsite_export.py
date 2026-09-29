# -*- coding: utf-8 -*-
"""Архив по РК для площадок без нашего кода (владелец 29.09.2026)."""
import io
import zipfile

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.database import SessionLocal
from app.sales.models import SalesDeal
from app.traffic import offsite_export as OX


def _deal_with_offsite(db):
    return db.execute(text("""
        SELECT s.deal_id FROM launch_prep_pair p
          JOIN launch_prep_creative_set s ON s.id = p.set_id
          JOIN launch_prep_target t ON t.id = p.target_id
          JOIN sales_publishers pb ON pb.id = t.publisher_id
          JOIN launch_prep_creative_file f ON f.set_id = s.id
         WHERE p.sent_at IS NOT NULL AND pb.our_code IS NOT TRUE
         ORDER BY s.deal_id DESC LIMIT 1""")).scalar()


def test_archive_has_passport_and_named_banners():
    db = SessionLocal()
    try:
        did = _deal_with_offsite(db)
        if not did:
            # На стенде все отправки ушли площадкам с нашим кодом — снимаем признак у одной
            # внутри транзакции теста (откат в finally, в базу не пишется).
            row = db.execute(text("""
                SELECT s.deal_id, t.publisher_id FROM launch_prep_pair p
                  JOIN launch_prep_creative_set s ON s.id = p.set_id
                  JOIN launch_prep_target t ON t.id = p.target_id
                  JOIN launch_prep_creative_file f ON f.set_id = s.id
                 WHERE p.sent_at IS NOT NULL ORDER BY s.deal_id DESC LIMIT 1""")).first()
            if not row:
                pytest.skip("на стенде нет отправленных креативов")
            db.execute(text("UPDATE sales_publishers SET our_code = false WHERE id = :p"), {"p": row[1]})
            did = row[0]
        deal = db.get(SalesDeal, did)
        body, name = OX.build(db, deal)
        assert name.endswith(".zip")
        z = zipfile.ZipFile(io.BytesIO(body))
        passport = [n for n in z.namelist() if n.endswith(".xlsx")]
        assert len(passport) == 1
        ws = load_workbook(io.BytesIO(z.read(passport[0]))).active
        head = [c.value for c in ws[1]]
        assert head == OX.HEAD and head[0] == "Статус у площадки"
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert rows
        names = set(z.namelist())
        for r in rows:
            path = r[head.index("Путь в архиве")]
            if not str(path).startswith("(файл"):
                assert path in names, "путь из паспорта не находится в архиве"
                assert path.endswith(r[head.index("Имя архива в выгрузке")])
                assert r[head.index("Название РК в Adfox")].startswith(
                    r[head.index("Имя архива в выгрузке")].rsplit(".", 1)[0])
        # Ни один путь не повторяется: web и app одной площадки — разные папки (ревью 29.09).
        assert len(z.namelist()) == len(set(z.namelist()))
        # Имена внутри архива — латиница (кириллица ломает распаковщики Windows).
        assert all(n.isascii() for n in names)
    finally:
        db.rollback()
        db.close()


def test_our_code_publishers_are_not_exported():
    import inspect
    assert "our_code.isnot(True)" in inspect.getsource(OX.build)


def test_adfox_name_is_archive_plus_traffic_initials():
    assert OX.traffic_initials("Жанна Смирнова") == "ZS"
    assert OX.traffic_initials("Гресева Дарья") == "DG", "имя первым, даже если в справочнике фамилия первой"
    assert OX.traffic_initials("Кисляков Алексей") == "AK"
    assert OX.traffic_initials(None) == ""
    assert OX.adfox_name("54ZYCH-SMK1-cr2-01.zip", "ZS") == "54ZYCH-SMK1-cr2-01_ZS"
    assert OX.adfox_name("54ZYCH-SMK1-cr2-01.zip", "") == "54ZYCH-SMK1-cr2-01"
