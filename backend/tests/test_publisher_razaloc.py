# -*- coding: utf-8 -*-
"""Разаллокация площадке по согласованной сверке (владелец 29.10.2026). Числа — руками по образцу владельца за 08.2026:
5 064 показа × CPM 100 = 506,40 без НДС, 617,81 с НДС 22 %."""
from datetime import date
from io import BytesIO

import openpyxl

from app.publisher_recon import razaloc as R

INITIAL = {"kind": "initial", "ord_id": "CT-A", "advertiser_inn": "7700000001", "advertiser_name": "Рекламодатель А",
           "contractor_inn": "7700000002", "contractor_name": "Агентство Б", "number": "Д-1",
           "date": date(2024, 12, 23), "type": "ServiceAgreement", "subject_type": "OrgDistribution",
           "action_type": None, "is_agent_acting_for_publisher": False}
MEDIATION = {**INITIAL, "ord_id": "CT-B", "advertiser_name": "Рекламодатель В", "number": "ДОГ-1", "type": "MediationContract",
             "subject_type": "Mediation", "action_type": "Contracting"}
DIRECT = {"kind": "final", "ord_id": "CT-C", "advertiser_inn": "7700000003", "advertiser_name": "Рекламодатель Г",
          "contractor_inn": "7700000004", "contractor_name": "НАШЕ ЮРЛИЦО ООО", "number": "РМ-1",
          "date": date(2026, 6, 1), "type": "ServiceAgreement", "subject_type": None, "action_type": None,
          "is_agent_acting_for_publisher": False}


def _line(code, volume, contract):
    return {"deal_code": code, "agreed_volume": volume, "contract": contract}


def test_two_deals_of_one_contract_are_one_row_with_budget_by_publisher_cpm():
    out = R.compose([_line("AAA111", 3000, INITIAL), _line("BBB222", 2064, INITIAL)], "site.test", 100, 22)
    assert out["missing"] == []
    (row,) = out["rows"]
    assert row["ID ИД медиаскаут"] == "CT-A"
    assert row["Показы"] == 5064 and row["Расчетные показы"] == 5064
    assert row["Бюджет до НДС"] == 506.40
    assert row["Бюджет с НДС"] == 617.81
    assert row["Налоговая форма"] == "С НДС" and row["НДС,%"] == 0.22
    assert row["Тип договора"] == "Договор оказания услуг"
    assert row["Вид договора/Предмет договора"] == "Договор на организацию распространения рекламы"
    assert row["РА1 действует в интересах"] == ""


def test_mediation_contract_says_in_whose_interests():
    row = R.compose([_line("AAA111", 22659, MEDIATION)], "site.test", 100, 22)["rows"][0]
    assert row["Тип договора"] == "Посреднический договор"
    assert row["Вид деятельности"] == "Заключение договоров"
    assert row["РА1 действует в интересах"] == "В интересах РД"


def test_direct_advertiser_goes_with_our_final_contract_and_us_as_ra1():
    row = R.compose([_line("GQNK5U", 5057, DIRECT)], "site.test", 100, 22)["rows"][0]
    assert row["ID ИД медиаскаут"] == "CT-C"
    assert row["Наименование РА1"] == "НАШЕ ЮРЛИЦО ООО"
    assert row["Вид договора/Предмет договора"] == "Договор на распространение рекламы"


def test_publisher_without_vat_gets_no_vat_and_equal_budgets():
    row = R.compose([_line("AAA111", 1000, INITIAL)], "x.ru", 125, 0)["rows"][0]
    assert row["Налоговая форма"] == "Без НДС" and row["НДС,%"] == 0
    assert row["Бюджет до НДС"] == row["Бюджет с НДС"] == 125.0


def test_deal_without_ord_contract_is_reported_not_silently_dropped():
    out = R.compose([_line("AAA111", 1000, INITIAL), _line("ZZZ999", 500, None)], "x.ru", 100, 22)
    assert out["missing"] == [{"deal_code": "ZZZ999", "volume": 500}]
    assert [r["Показы"] for r in out["rows"]] == [1000]


def test_zero_volume_lines_do_not_make_rows():
    assert R.compose([_line("AAA111", 0, INITIAL)], "x.ru", 100, 22)["rows"] == []


def test_xlsx_has_the_owner_sample_header_and_rows():
    out = R.compose([_line("AAA111", 5064, INITIAL)], "site.test", 100, 22)
    ws = openpyxl.load_workbook(BytesIO(R.xlsx([("site.test", out["rows"])]))).active
    head = [c.value for c in ws[1]]
    assert head == ["Сайт", "ID ИД медиаскаут", "ИНН РД", "Наименование РД", "ИНН РА1", "Наименование РА1",
                    "Номер договора", "Дата договора", "Тип договора", "Вид договора/Предмет договора",
                    "Вид деятельности", "РА1 действует в интересах", "Налоговая форма", "Показы", "CPM", "НДС,%",
                    "Расчетные показы", "Бюджет до НДС", "Бюджет с НДС"]
    assert ws.cell(2, 1).value == "site.test" and ws.cell(2, 18).value == 506.4


def test_several_publishers_are_one_file_with_a_sheet_each():
    """Владелец 29.10.2026: у кабинета нескольких площадок — один файл, площадка на своей вкладке."""
    a = R.compose([_line("AAA111", 1000, INITIAL)], "a.ru", 100, 22)["rows"]
    b = R.compose([_line("BBB222", 2000, MEDIATION)], "b.ru", 100, 22)["rows"]
    wb = openpyxl.load_workbook(BytesIO(R.xlsx([("a.ru", a), ("b.ru", b)])))
    assert wb.sheetnames == ["a.ru", "b.ru"]
    assert wb["b.ru"].cell(2, 2).value == "CT-B" and wb["b.ru"].cell(2, 14).value == 2000


def test_sheet_names_survive_excel_limits():
    """Имя вкладки в Excel — до 31 знака, без []:*?/\ и без повторов; иначе openpyxl падает или Excel чинит файл."""
    long = "очень-длинное-имя-площадки-на-сорок-знаков.ru"
    wb = openpyxl.load_workbook(BytesIO(R.xlsx([(long, []), (long, []), ("a/b:c", [])])))
    assert [len(n) <= 31 for n in wb.sheetnames] == [True] * 3
    assert len(set(wb.sheetnames)) == 3 and wb.sheetnames[2] == "a_b_c"


# ── ручки: до согласования нет, без права нет, чужую через шлюз не открыть ──────────────────────────

import pytest  # noqa: E402
from fastapi import HTTPException  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import app.model_registry  # noqa: E402,F401
from app.database import SessionLocal  # noqa: E402
from app.main import app  # noqa: E402
from app.publisher_requests import PublisherRequest  # noqa: E402
from app.routers.auth import get_current_user  # noqa: E402


class _Role:
    def __init__(self, key):
        self.key, self.id = key, None


class _User:
    def __init__(self, key):
        self.id, self.name, self.email = None, "test", "t@t"
        self.role, self.role_id, self.consent_accepted_at = _Role(key), None, True


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def test_screen_needs_the_interaction_right(client):
    app.dependency_overrides[get_current_user] = lambda: _User("role_13")      # трафик — права нет
    assert client.get("/api/publisher-recon/1/razaloc").status_code == 403
    assert client.get("/api/publisher-recon/1/razaloc.xlsx").status_code == 403


def test_not_agreed_recon_has_no_razaloc():
    db = SessionLocal()
    try:
        req = PublisherRequest(kind="сверка", publisher_id=1, period="2026-10")
        db.add(req)
        db.flush()
        with pytest.raises(HTTPException) as e:
            R.response(db, [req])
        assert e.value.status_code == 409
    finally:
        db.rollback()
        db.close()


def test_gateway_razaloc_does_not_open_someone_elses_recon():
    from app.routers.cabinet_gateway_recon import razaloc
    db = SessionLocal()
    try:
        req = PublisherRequest(kind="сверка", publisher_id=1, period="2026-10", verdict="согласовано")
        db.add(req)
        db.flush()
        with pytest.raises(HTTPException) as e:
            razaloc(ids=[req.id], account_id=0, db=db)
        assert e.value.status_code == 404
    finally:
        db.rollback()
        db.close()
