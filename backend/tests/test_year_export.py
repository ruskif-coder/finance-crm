"""
Лист «Годовой МП»: подытоги месяца и бренда не должны попадать в сумму дважды.

Ради чего тест: на листе три уровня сумм (месяц → бренд → ИТОГО), и все они лежат
ВНУТРИ одного блока строк. Наивное `=СУММ(первая:последняя)`, которым считается ИТОГО
на обычном месячном листе, здесь сложило бы и данные, и подытоги — годовой бюджет
удвоился бы, а в Excel это выглядит как правдоподобное число, а не как ошибка.
"""
import re

from openpyxl import Workbook

from app.routers.media_plans import subtotal_formulas
from app.year_mp_export import (
    DISCOUNT_FILL,
    DISCOUNT_HEADER,
    METRICS_SLOTS,
    brand_groups,
    highlight_discounts,
    metrics_extra_rows,
    year_row_hook,
)

# Раскладка блока: полоса бренда, полоса месяца, данные, подытог месяца, …, подытог бренда.
# Ровно то, что собирает year_rows: два бренда, у первого два месяца, у второго один.
LAYOUT = [
    (10, {"_band": "БРЕНД А"}),
    (11, {"_band": "Бренд А · Янв"}),
    # 12, 13 — данные
    (14, {"_subtotal": "Итого · Янв", "_kind": "month"}),
    (15, {"_band": "Бренд А · Мар"}),
    # 16 — данные
    (17, {"_subtotal": "Итого · Мар", "_kind": "month"}),
    (18, {"_subtotal": "Итого · Бренд А за год", "_kind": "brand"}),
    (19, {"_band": "БРЕНД Б"}),
    (20, {"_band": "Бренд Б · Янв"}),
    # 21 — данные
    (22, {"_subtotal": "Итого · Янв", "_kind": "month"}),
    (23, {"_subtotal": "Итого · Бренд Б за год", "_kind": "brand"}),
]
# В шаблоне блок начинается с несуммируемой колонки («Место размещения», B) — в неё
# и уходит подпись подытога. Держим её в раскладке теста: без неё подпись села бы на
# первую же денежную колонку и затёрла формулу.
COLS = {"place": "B", "net": "Q", "gross": "S"}
TCOL = {"net": "Q", "gross": "S"}


def _run():
    ws = Workbook().active
    sum_rows = year_row_hook(ws, COLS, TCOL, 10, 23, list(LAYOUT))
    return ws, sum_rows


def _refs(formula):
    """Номера строк, на которые ссылается формула, с раскрытием диапазонов."""
    out = set()
    for a, b in re.findall(r"Q(\d+)(?::Q(\d+))?", formula or ""):
        out.update(range(int(a), int(b or a) + 1))
    return out


def test_month_subtotal_sums_only_its_own_data_rows():
    ws, _ = _run()
    assert _refs(ws["Q14"].value) == {12, 13}
    assert _refs(ws["Q17"].value) == {16}
    assert _refs(ws["Q22"].value) == {21}


def test_brand_subtotal_sums_month_subtotals_not_data():
    """Иначе месяц вошёл бы в бренд дважды — и как данные, и как свой подытог."""
    ws, _ = _run()
    assert _refs(ws["Q18"].value) == {14, 17}
    assert _refs(ws["Q23"].value) == {22}


def test_grand_total_sums_brand_subtotals():
    _, sum_rows = _run()
    assert sum_rows == [18, 23]


def test_no_row_is_counted_twice_down_the_cascade():
    """Раскрыв каскад до листьев, каждая строка данных обязана встретиться один раз."""
    ws, sum_rows = _run()

    def leaves(row):
        formula = ws[f"Q{row}"].value
        if not formula:                       # строка данных — сама себе лист
            return [row]
        return [x for r in sorted(_refs(formula)) for x in leaves(r)]

    got = [x for r in sum_rows for x in leaves(r)]
    assert sorted(got) == [12, 13, 16, 21]
    assert len(got) == len(set(got))


def test_label_does_not_overwrite_a_summed_column():
    """Подпись садится в несуммируемую колонку — иначе подытог стал бы текстом."""
    ws, _ = _run()
    assert ws["B14"].value == "Итого · Янв"
    assert ws["Q14"].value.startswith("=SUM(")
    assert ws["Q10"].value is None            # полоса остаётся пустой


# Текущая раскладка листа: месяц виден в колонке «Период», поперечных рядов нет,
# подытог только по бренду. Каскад из тестов выше при этом сохранён — месячные
# подытоги убрали по решению владельца, а не потому, что механизм их не тянет.
FLAT = [
    (10, {"_band": "БРЕНД А"}),
    # 11, 12, 13 — данные всех месяцев подряд
    (14, {"_subtotal": "Итого · Бренд А за год", "_kind": "brand"}),
    (15, {"_band": "БРЕНД Б"}),
    # 16 — данные
    (17, {"_subtotal": "Итого · Бренд Б за год", "_kind": "brand"}),
]


def test_brand_subtotal_falls_back_to_its_own_rows_without_month_level():
    ws = Workbook().active
    sum_rows = year_row_hook(ws, COLS, TCOL, 10, 17, list(FLAT))
    assert _refs(ws["Q14"].value) == {11, 12, 13}
    assert _refs(ws["Q17"].value) == {16}
    assert sum_rows == [14, 17]


def test_brand_groups_reads_ranges_from_the_bands():
    """Блок показателей считает по этим диапазонам — ошибка здесь тихо испортит цифры."""
    assert brand_groups(FLAT, 10, 17) == [("БРЕНД А", 11, 13), ("БРЕНД Б", 16, 16)]


def test_group_ends_at_the_next_service_row_not_at_the_subtotal():
    """Подытог бренда в диапазон входить не должен: он сложился бы со своими же строками."""
    (_, _, last), = brand_groups(FLAT[:2], 10, 17)
    assert last == 13                          # 14 — строка подытога


def test_band_without_rows_is_not_a_group():
    """Пустая полоса дала бы диапазон наизнанку и =СУММ по чужим строкам."""
    assert brand_groups([(10, {"_band": "ПУСТОЙ"}), (11, {"_band": "БРЕНД"})], 10, 12) == \
        [("БРЕНД", 12, 12)]


def test_block_grows_only_past_five_brands():
    """До пяти брендов медиаплан стоит на месте; дальше — по строке на бренд."""
    assert metrics_extra_rows(3) == 0
    assert metrics_extra_rows(METRICS_SLOTS) == 0
    assert metrics_extra_rows(7) == 2


def _discount_sheet():
    """Лист с колонкой скидки: ставка, ноль, формула и итог — как в реальной книге."""
    ws = Workbook().active
    ws["O5"] = DISCOUNT_HEADER
    ws["O6"], ws["O7"], ws["O8"], ws["O9"] = 50, 0, "=SUM(O6:O7)", 100
    return ws


def _green(ws, coord):
    c = ws[coord]
    return bool(c.fill.patternType) and str(c.fill.fgColor.rgb) == DISCOUNT_FILL


def test_only_cells_with_a_real_discount_turn_green():
    ws = _discount_sheet()
    assert highlight_discounts(ws) == 2
    assert _green(ws, "O6") and _green(ws, "O9")
    assert not _green(ws, "O7")          # ноль — это отсутствие скидки, а не скидка
    assert not _green(ws, "O8")          # формула не число, итог не подсвечиваем
    assert not _green(ws, "O5")          # заголовок колонки


def test_discount_columns_are_found_by_header_in_both_tables():
    """Размещения и доп. услуги — две таблицы со своим заголовком скидки на листе."""
    ws = _discount_sheet()
    ws["O20"], ws["O21"] = DISCOUNT_HEADER, 30
    assert highlight_discounts(ws) == 3
    assert _green(ws, "O21")


def test_single_rows_are_not_written_as_degenerate_ranges():
    """`Q14:Q14` Excel съест, но в формуле бренда это читается как диапазон данных."""
    out = subtotal_formulas(COLS, TCOL, [14, 17])
    assert out["Q"] == "=SUM(Q14,Q17)"
    assert out["S"] == "=SUM(S14,S17)"


# ── шапка листа «Годовой МП» вырезана, показатели переехали в «Сводную» ──
# Ради чего: удаление строк в openpyxl переносит только значения — объединения и высоты
# остаются на прежних номерах, и шапка садится на чужие строки. Проверяем оба.

def test_drop_rows_moves_merges_and_heights():
    from app.year_mp_export import _drop_rows

    wb = Workbook()
    ws = wb.active
    ws["B9"] = "шапка"
    ws["B18"] = "БАЗОВЫЕ УСЛУГИ:"
    ws.merge_cells("B9:C10")        # внутри вырезаемого — исчезает
    ws.merge_cells("B19:B20")       # ниже — поднимается на 9
    ws.merge_cells("B3:D5")         # выше — не двигается
    ws.row_dimensions[19].height = 40

    _drop_rows(ws, 9, 9)

    merged = {str(r) for r in ws.merged_cells.ranges}
    assert ws["B9"].value == "БАЗОВЫЕ УСЛУГИ:"
    assert merged == {"B3:D5", "B10:B11"}
    assert ws.row_dimensions[10].height == 40
    assert ws.row_dimensions[19].height is None


def test_kpi_table_counts_by_the_purchase_sheet():
    """Сводная считается формулами по строкам бренда, а не своими числами: правка
    объёма прямо в книге обязана пересчитать показатели."""
    from app.year_mp_export import _kpi_table

    wb = Workbook()
    ws = wb.active
    brands = [{"name": "Бренд",
               "services": [{"cost": [100] + [0] * 11, "vol": [1000] + [0] * 11, "is_cpm": True}]}]
    # Полоса бренда на листе закупки пишется капсом — сводка обязана его найти.
    year_ref = ("Годовой МП", {"net_nodisc": "N", "disc_rub": "P", "net": "Q",
                               "gross": "S", "clicks": "X", "checks": "AC",
                               "revenue": "AF"}, [("БРЕНД", 13, 21)])

    last = _kpi_table(ws, 10, brands, year_ref, 0.22)

    assert ws["F11"].value == "=SUM('Годовой МП'!N13:N21)"
    assert ws["I11"].value == "=SUM('Годовой МП'!Q13:Q21)"
    # CPM — от цены БЕЗ НДС (колонка I), как и построчно на листе закупки.
    assert ws["K11"].value == '=IF(E11>0,I11/E11*1000,"")'
    # Удельные показатели в ИТОГО — взвешенные, а не сумма по брендам.
    assert ws[f"M{last}"].value.startswith("=IF(SUM(")
    assert ws[f"I{last}"].value == "=SUM(I11:I11)"


def test_kpi_table_survives_a_plan_without_a_purchase_sheet():
    """План из одних доп. услуг: листа «Годовой МП» нет, ссылаться не на что —
    остаются суммы плана, а не битые формулы."""
    from app.year_mp_export import _kpi_table

    wb = Workbook()
    ws = wb.active
    brands = [{"name": "Б", "services": [{"cost": [100] + [0] * 11, "vol": [0] * 12, "is_cpm": False}]}]

    _kpi_table(ws, 10, brands, None, 0.22)

    assert ws["F11"].value == 100
    assert ws["J11"].value == 122
    assert ws["N11"].value is None      # дохода взять неоткуда — пусто, а не ноль


# ── источник строк: проверенный МП вместо плана (вариант A, 2026-08-22) ──

class _Line:
    """Строка плана: месяц 4 разбит кнопкой «+ сделка» на две группы."""
    id = 5
    year = 2026
    brief = {}
    service_forecast = {}
    products = {"4": [{"ref_id": 28, "type": "service", "amount": 500000, "deal_idx": 0},
                      {"ref_id": 17, "type": "service", "amount": 200000, "deal_idx": 1}]}


class _Svc:
    def __init__(self, name, price):
        self.name, self.unit_price = name, price
        self.separate_price = False
        self.calc_form = "CPM"
        self.placement_type = "Banners"


SVC = {28: _Svc("еФарм", 450), 17: _Svc("Polza", 300)}


def test_parts_falls_back_to_the_plan_without_a_verified_mp():
    from app.year_mp_export import _parts

    rows, _ = _parts(_Line(), 4, SVC, {}, _Line.products["4"], {})
    assert [r["position"] for r in rows] == ["еФарм", "Polza"]
    assert all(r["discount"] == 0 for r in rows)      # в плане скидки нет вовсе


def test_parts_replaces_only_the_checked_deal_of_a_split_month():
    """Месяц, проверенный наполовину: подменяется сделка с проверенным МП, вторая
    считается из плана. Подмена ячейки целиком выкинула бы её услуги из книги —
    молча и правдоподобно."""
    from app.year_mp_export import _parts

    mp_row = {"position": "еФарм", "format": None, "model": "CPM", "inventory": "web",
              "volume": 1111111, "unit_price": 450, "discount": 0.15, "forecast": {}}
    verified = {(5, 4): {0: ([mp_row], [])}}

    rows, _ = _parts(_Line(), 4, SVC, {}, _Line.products["4"], verified)

    assert [r["position"] for r in rows] == ["еФарм", "Polza"]
    assert rows[0]["discount"] == 0.15      # из проверенного МП
    assert rows[1]["discount"] == 0         # из плана
