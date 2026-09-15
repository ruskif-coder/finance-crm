# -*- coding: utf-8 -*-
"""Выгрузка медиаплана в Excel: клон строки-образца должен быть ПОЛНОЙ копией.

Жалоба владельца 15.09.2026 по сделке WGUPN5: «пропало наименование доп. услуг».

Разбор. Шаблон (`app/templates/mp_template.xlsx`) держит по одной строке-образцу на
блок: `{{r.*}}` для размещений и `{{e.*}}` для доп. услуг. Под факт они клонируются
`_clone_below`, и клон получал значение, шрифт, заливку, рамку, выравнивание, высоту —
всё, кроме ОБЪЕДИНЕНИЙ ВНУТРИ строки. А название доп. услуги в шаблоне растянуто на
`C..H`: без объединения оно оставалось в узкой колонке C с переносом по словам и
клонированной фиксированной высотой, то есть обрезалось до невидимого.

Отсюда и форма жалобы: ПЕРВАЯ доп. услуга выглядела правильно — она садится в саму
строку-образец, объединение которой никуда не делось, — а вторая и последующие
приезжали с пустой колонкой названия. На плане с одной доп. услугой дефект не виден
вовсе, поэтому он и дожил до жалобы.

Прибор рендерит настоящий шаблон настоящим кодом рендера: подделка книги проверяла бы
подделку. На стенде без шаблона тест пропускается — это не повод краснеть, но и не
повод считать проверку пройденной.
"""
import os

import pytest

from app.routers.media_plans import (TEMPLATE_PATH, render_mp_sheet,
                                     _find_token_row, _token_cols)

pytestmark = pytest.mark.filterwarnings("ignore::DeprecationWarning")


def _sheet():
    from openpyxl import load_workbook
    path = os.path.abspath(TEMPLATE_PATH)
    if not os.path.exists(path):
        pytest.skip("шаблон mp_template.xlsx недоступен")
    wb = load_workbook(path)
    return wb, (wb["МП"] if "МП" in wb.sheetnames else wb.active)


def _full(rows, extras):
    return {"rows": rows, "extras": extras, "advertiser": "Рекламодатель",
            "brand": "Бренд", "agency": "Агентство", "title": "[тест] выгрузка",
            "period": "2026-09", "geo": "РФ", "targeting": {}, "created_at": None,
            "date_from": None, "date_to": None}


ROW = {"position": "еФарм", "format": "Banners", "model": "CPM", "inventory": "web",
       "volume": 1000000, "unit_price": 330, "discount": 0, "forecast": {}}
NAMES = ["Sales-Lift", "Brand-Lift", "Охватный отчёт"]


def test_every_extra_keeps_its_name_cell_wide():
    """Каждая строка доп. услуги несёт имя и объединение под него, а не только первая."""
    wb, ws = _sheet()
    src_e = _find_token_row(ws, "{{e.")
    if not src_e:
        pytest.skip("в шаблоне нет блока доп. услуг")
    name_col = _token_cols(ws, src_e, "e").get("name")
    tpl_merge = [(mr.min_col, mr.max_col) for mr in ws.merged_cells.ranges
                 if mr.min_row == mr.max_row == src_e]

    render_mp_sheet(ws, _full([dict(ROW)], [{"name": n, "price": 100000, "total": 90000}
                                            for n in NAMES]))

    # Строки доп. услуг ищем по именам: блок сдвинулся вниз вместе со вставками.
    found = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value in NAMES:
                found[cell.value] = cell
    assert set(found) == set(NAMES), f"в книге нет наименований: {set(NAMES) - set(found)}"

    if not tpl_merge:
        return                      # в шаблоне имя не растянуто — проверять нечего
    for name, cell in found.items():
        assert cell.column_letter == name_col, f"«{name}» не в своей колонке"
        widths = [(mr.min_col, mr.max_col) for mr in ws.merged_cells.ranges
                  if mr.min_row == mr.max_row == cell.row]
        assert widths == tpl_merge, (
            f"у строки «{name}» потеряно объединение под название: {widths} вместо {tpl_merge}")


def test_cloned_rows_keep_the_alignment_of_the_template():
    """Выравнивание клонов повторяет образец — его задаёт аккаунт в шаблоне, а не код.

    Вторая половина той же жалобы («всё выравнивание по центру») кодом не
    воспроизводится: прибор фиксирует это утверждение, чтобы будущая правка рендера
    не сломала то, что сегодня цело.
    """
    wb, ws = _sheet()
    src_r = _find_token_row(ws, "{{r.")
    if not src_r:
        pytest.skip("в шаблоне нет блока размещений")
    want = {c.column_letter: (c.alignment.horizontal, c.alignment.vertical)
            for c in ws[src_r] if isinstance(c.value, str) and "{{r." in c.value}

    render_mp_sheet(ws, _full([dict(ROW) for _ in range(4)], []))

    for i in range(4):
        r = src_r + i
        for col, pair in want.items():
            got = ws[f"{col}{r}"].alignment
            assert (got.horizontal, got.vertical) == pair, (
                f"строка {r}, колонка {col}: выравнивание {got.horizontal}/{got.vertical} "
                f"вместо {pair[0]}/{pair[1]}")
