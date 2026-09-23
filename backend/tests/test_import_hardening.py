# -*- coding: utf-8 -*-
"""Прибор: импорт-сверка не портит деньги и периоды (аудит 23.09.2026, этап 1 плана).

  · 2.M1 — подтверждение конфликта по ОДНОМУ полю переписывало квартальный период
    «Q3 2026» в месяц даты: разбор файла сводит квартал к месяцу, а применение писало
    это значение, хотя период никто не менял. Отчёты переставали делить сумму на три месяца;
  · 2.M2 — запрет отрицательных сумм и ставки НДС вне 0–100 был только у формы, файл шёл
    мимо: «Списания = −5000» переворачивал знак в ДДС и P&L;
  · 2.M3 — период, который Excel превратил в дату, читался строкой «2026-01-01 00:00:00»
    и не распознавался: операция получала месяц оплаты или теряла период вовсе;
  · `vat_fact` из файла побеждал серверный расчёт (аудит 11.09, №9) — НДС всегда считает сервер.

Файл строится настоящей выгрузкой и читается настоящим разбором.
"""
import asyncio
import io
from datetime import date, datetime

import openpyxl
import pytest
from fastapi import HTTPException, UploadFile

import app.ad.models           # noqa: F401
import app.launch_prep.models  # noqa: F401
import app.notify.models       # noqa: F401
import app.ord.models          # noqa: F401
from app.database import SessionLocal
from app.models import Article, Counterparty, Operation, User
from app.routers import operations as ops

DS = "прибор импорта"


@pytest.fixture
def db():
    s = SessionLocal()
    s.commit = s.flush
    try:
        yield s
    finally:
        s.rollback()
        s.close()


@pytest.fixture
def admin(db):
    u = (db.query(User).filter(User.is_active == 1, User.role.has(key="admin"))
         .order_by(User.id).first())
    if u is None:
        pytest.skip("нужна учётка админа")
    return u


@pytest.fixture
def quarter_op(db):
    art = db.query(Article).order_by(Article.id).first()
    cp = db.query(Counterparty).order_by(Counterparty.id).first()
    op = Operation(date=date(2099, 8, 15), status="ОПЛАЧЕНО", income=900.0, expense=0.0,
                   bank="АльфаБанк", period="Q3 2099", vat_rate=22,
                   vat_fact=ops.compute_vat_fact(900.0, 0, 22), article_id=art.id,
                   counterparty_id=cp.id, ds_num=DS, invoice="1")
    db.add(op)
    db.flush()
    return op


def _export(db, admin, counterparty_id) -> bytes:
    resp = ops.export_operations(status=None, bank=None, date_from=None, date_to=None,
                                 article_id=None, counterparty_id=[counterparty_id],
                                 period=None, gaps=None, db=db, current_user=admin)

    async def read():
        return b"".join([c async for c in resp.body_iterator])
    return asyncio.run(read())


def _edit(data: bytes, fn) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    fn(ws, [c.value for c in ws[1]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _mine(ws, hdr):
    i_ds = hdr.index("№ ДС")
    return [row for row in ws.iter_rows(min_row=2) if row[i_ds].value == DS]


def _preview(db, admin, data):
    res = asyncio.run(ops.import_preview(file=UploadFile(filename="s.xlsx", file=io.BytesIO(data)),
                                         db=db, current_user=admin))
    return res, ops.IMPORT_SYNC_CACHE[res["import_id"]]["rows"]


def _apply(db, admin, res, keys):
    return asyncio.run(ops.import_apply(ops.ImportApplyRequest(import_id=res["import_id"],
                                                               confirmed_keys=keys),
                                        db=db, current_user=admin))


def test_confirming_another_field_keeps_the_quarter(db, admin, quarter_op):
    def edit(ws, hdr):
        for row in _mine(ws, hdr):
            row[hdr.index("Назначение")].value = "правка назначения"
    res, cache = _preview(db, admin, _edit(_export(db, admin, quarter_op.counterparty_id), edit))
    conflict = next(c for c in cache if c["data"].get("ds_num") == DS)
    assert conflict["status"] == "conflict"
    _apply(db, admin, res, [conflict["key"]])
    db.expire_all()
    op = db.query(Operation).filter(Operation.id == quarter_op.id).first()
    assert op.description == "правка назначения"
    assert op.period == "Q3 2099", "квартал переписан в месяц даты"


def test_moving_the_date_to_another_month_keeps_an_untouched_quarter(db, admin, quarter_op):
    """Дату сдвинули с августа на сентябрь, ячейку «Период» не трогали. Разбор сводит
    «Q3 2099» к месяцу НОВОЙ даты (2099-09), сравнение со старым периодом по СТАРОЙ
    дате даёт 2099-08 — и квартал переписывался месяцем (ревью 23.09.2026)."""
    def edit(ws, hdr):
        for row in _mine(ws, hdr):
            row[hdr.index("Дата")].value = date(2099, 9, 20)
    res, cache = _preview(db, admin, _edit(_export(db, admin, quarter_op.counterparty_id), edit))
    conflict = next(c for c in cache if c["data"].get("ds_num") == DS)
    _apply(db, admin, res, [conflict["key"]])
    db.expire_all()
    op = db.query(Operation).filter(Operation.id == quarter_op.id).first()
    assert op.date == date(2099, 9, 20)
    assert op.period == "Q3 2099", "квартал переписан месяцем новой даты"


def test_a_really_changed_period_is_still_written(db, admin, quarter_op):
    def edit(ws, hdr):
        for row in _mine(ws, hdr):
            row[hdr.index("Период")].value = "2099-09"
    res, cache = _preview(db, admin, _edit(_export(db, admin, quarter_op.counterparty_id), edit))
    conflict = next(c for c in cache if c["data"].get("ds_num") == DS)
    _apply(db, admin, res, [conflict["key"]])
    db.expire_all()
    assert db.query(Operation).filter(Operation.id == quarter_op.id).first().period == "2099-09"


@pytest.mark.parametrize("field,value,says", [("income", -5000.0, "отрицательн"),
                                              ("expense", -1.0, "отрицательн"),
                                              ("vat_rate", 150.0, "вне диапазона"),
                                              ("vat_rate", -22.0, "вне диапазона")])
def test_signs_and_vat_range_are_checked_for_every_path(field, value, says):
    # Банк задан и текст отказа проверяется: иначе строку отклоняло бы ДРУГОЕ правило
    # («оплаченная без банка»), и тест оставался зелёным без проверки знака (ревью 23.09.2026).
    op = Operation(status="ОПЛАЧЕНО", income=0.0, expense=100.0, vat_rate=0, article_id=1,
                   period="2099-01", bank="АльфаБанк")
    assert ops._operation_problem(op) is None, "исходная строка должна быть годной"
    if field == "income":
        op.expense = 0.0
    setattr(op, field, value)
    problem = ops._operation_problem(op)
    assert problem and says in problem, f"{field}={value}: {problem!r}"


def test_negative_amount_in_a_file_is_refused(db, admin, quarter_op):
    def edit(ws, hdr):
        for row in _mine(ws, hdr):
            row[hdr.index("ID")].value = None
            row[hdr.index("Счет")].value = "2"          # другая операция — новая строка
            row[hdr.index("Поступления")].value = -5000
    res, _ = _preview(db, admin, _edit(_export(db, admin, quarter_op.counterparty_id), edit))
    with pytest.raises(HTTPException) as e:
        _apply(db, admin, res, [])
    assert e.value.status_code == 400
    assert "отрицательн" in str(e.value.detail), e.value.detail


def test_period_turned_into_a_date_by_excel_is_read_as_its_month(db, admin, quarter_op):
    def edit(ws, hdr):
        for row in _mine(ws, hdr):
            row[hdr.index("Период")].value = datetime(2099, 1, 1)
            row[hdr.index("Дата")].value = datetime(2099, 3, 20)   # месяц оплаты — другой
    rows = ops._parse_cf_best_rows(_edit(_export(db, admin, quarter_op.counterparty_id), edit))
    mine = [r for r in rows if r["ds_num"] == DS]
    assert [r["period"] for r in mine] == ["2099-01"]


def test_vat_fact_from_the_file_is_ignored(db, admin, quarter_op):
    def edit(ws, hdr):
        for row in _mine(ws, hdr):
            row[hdr.index("ID")].value = None
            row[hdr.index("Счет")].value = "3"
            col = next(i for i, h in enumerate(hdr) if h and "сумма" in str(h).lower() and "ндс" in str(h).lower())
            row[col].value = 1
    rows = ops._parse_cf_best_rows(_edit(_export(db, admin, quarter_op.counterparty_id), edit))
    mine = [r for r in rows if r["ds_num"] == DS]
    assert mine[0]["vat_fact"] == ops.compute_vat_fact(900.0, 0, 22)


def test_paid_operation_without_a_bank_is_refused():
    """Оплаченная без банка выпадала из остатков и сводки ДДС (аудит 23.09.2026, 2.L8)."""
    op = Operation(status="ОПЛАЧЕНО", income=100.0, expense=0.0, vat_rate=0, article_id=1,
                   period="2099-01", bank=None)
    assert ops._operation_problem(op)
    op.bank = "АльфаБанк"
    assert not ops._operation_problem(op)
    op.status, op.bank = "ПЛАН ПОСТУПЛЕНИЙ", None       # у плана банк не обязателен
    assert not ops._operation_problem(op)
