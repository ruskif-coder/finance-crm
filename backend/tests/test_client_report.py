# -*- coding: utf-8 -*-
"""Отчёт клиенту по РК (владелец 05.10.2026): суммы блоков сходятся, Weborama в отчёт не
попадает, показы в модели — реальные; ручка закрыта правами и отдаёт Excel."""
import io
from datetime import date

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.ad import client_report as CR
from app.database import SessionLocal
from app.main import app
from app.routers.auth import get_current_user


def _camp_with_fact(db):
    row = db.execute(text("""
        SELECT c.id, c.deal_id FROM ad_campaign c
         WHERE EXISTS (SELECT 1 FROM ad_campaign_stat s WHERE s.campaign_id = c.id
                         AND s.source IN ('dsp', 'ms', 'adfox') AND s.shows > 0)
         ORDER BY c.id DESC LIMIT 1""")).first()
    if row is None:
        pytest.skip("на стенде нет РК с боевым фактом")
    return row


def test_blocks_add_up_and_weborama_is_out():
    db = SessionLocal()
    try:
        cid, _ = _camp_with_fact(db)
        rep = CR.build(db, cid)
        own = db.execute(text("""SELECT coalesce(sum(shows), 0) FROM ad_campaign_stat
            WHERE campaign_id = :c AND source IN ('dsp', 'ms', 'adfox')
              AND date BETWEEN :a AND :b"""),
            {"c": cid, "a": rep.head["period_from"], "b": rep.head["period_to"]}).scalar()
        assert rep.totals["shows"] == own, "в отчёте не только наш боевой факт"
        assert sum(r["shows"] for r in rep.by_placement) == rep.totals["shows"]
        assert sum(r["shows"] for r in rep.by_day) == rep.totals["shows"]
        assert sum(v[0] for d in rep.site_days.values() for v in d.values()) == rep.totals["shows"]
    finally:
        db.close()


@pytest.fixture(autouse=True)
def _keep_coefs():
    """Коэффициенты, зафиксированные тестом, убираются: иначе тест занял бы дни стенда
    значениями по умолчанию раньше настроек владельца (05.10.2026)."""
    db = SessionLocal()
    start = db.execute(text("SELECT coalesce(max(id), 0) FROM report_daily_coef")).scalar()
    db.close()
    yield
    db = SessionLocal()
    db.execute(text("DELETE FROM report_daily_coef WHERE id > :s"), {"s": start})
    db.commit()
    db.close()


def test_model_keeps_real_shows():
    db = SessionLocal()
    try:
        cid, _ = _camp_with_fact(db)
        f, m = CR.build(db, cid), CR.build(db, cid, model=True)
        assert m.totals["shows"] == f.totals["shows"]
        assert all(r["uniques"] is not None for r in m.by_placement)
        assert all(r["uniques"] is None for r in f.by_placement)
    finally:
        db.close()


class _Role:
    def __init__(self, key):
        self.key = key


class _User:
    def __init__(self, key):
        self.id, self.name, self.email = None, "t", "t@t"
        self.role = _Role(key)
        self.role_id = -1                # роли нет — ни одного права
        self.consent_accepted_at = True


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def test_admin_downloads_excel_with_both_sets(client):
    db = SessionLocal()
    _, deal_id = _camp_with_fact(db)
    db.close()
    app.dependency_overrides[get_current_user] = lambda: _User("admin")
    r = client.get(f"/api/client-report/deal/{deal_id}.xlsx")
    assert r.status_code == 200, r.text[:200]
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Отчёт", "По неделям", "Отчёт (модель)", "По неделям (модель)"]


def test_no_campaign_and_bad_period(client):
    app.dependency_overrides[get_current_user] = lambda: _User("admin")
    assert client.get("/api/client-report/deal/999999999.xlsx").status_code == 404
    db = SessionLocal()
    _, deal_id = _camp_with_fact(db)
    db.close()
    r = client.get(f"/api/client-report/deal/{deal_id}.xlsx?date_from=2026-10-05&date_to=2026-10-01")
    assert r.status_code == 400
    assert client.get(f"/api/client-report/deal/{deal_id}.xlsx?date_from=05.10").status_code == 400


def test_role_without_dashboards_is_refused(client):
    app.dependency_overrides[get_current_user] = lambda: _User("no_such_role")
    assert client.get("/api/client-report/deal/1.xlsx").status_code == 403


def test_period_respected():
    db = SessionLocal()
    try:
        cid, _ = _camp_with_fact(db)
        full = CR.build(db, cid)
        d = full.by_day[0]["label"]
        one = CR.build(db, cid, d, d)
        assert [r["label"] for r in one.by_day] == [d]
        assert one.totals["shows"] == full.by_day[0]["shows"]
        assert isinstance(d, date)
    finally:
        db.close()


def test_no_stats_no_report(client):
    """Отчёт — только при статистике хоть за один день (владелец 05.10.2026)."""
    db = SessionLocal()
    row = db.execute(text("""
        SELECT c.deal_id FROM ad_campaign c WHERE NOT EXISTS (
            SELECT 1 FROM ad_campaign_stat s WHERE s.campaign_id = c.id
               AND s.source IN ('dsp', 'ms', 'adfox') AND s.shows > 0) LIMIT 1""")).first()
    db.close()
    if row is None:
        pytest.skip("на стенде нет РК без статистики")
    app.dependency_overrides[get_current_user] = lambda: _User("admin")
    assert client.get(f"/api/client-report/deal/{row[0]}.xlsx").status_code == 409
    db = SessionLocal()
    assert CR.deal_report_ready(db, row[0]) is False
    db.close()


# ── ревью 05.10.2026: границы периода ───────────────────────────────────────
D = date


def test_period_without_flight_dates_uses_data():
    """Нет даты старта/конца РК — период и колонки дней от первого до последнего дня данных."""
    p = CR.period_of(start=None, end=None, first=D(2026, 10, 2), last=D(2026, 10, 9),
                     date_from=None, date_to=None)
    assert (p["from"], p["to"], p["days_to"]) == (D(2026, 10, 2), D(2026, 10, 9), D(2026, 10, 9))


def test_days_cover_data_after_flight_end():
    """Докрутка после конца флайта не выпадает из колонок дней."""
    p = CR.period_of(start=D(2026, 10, 1), end=D(2026, 10, 31), first=D(2026, 10, 1),
                     last=D(2026, 11, 2), date_from=None, date_to=None)
    assert p["days_to"] == D(2026, 11, 2) and p["to"] == D(2026, 11, 2)


def test_default_days_run_to_flight_end():
    p = CR.period_of(start=D(2026, 10, 1), end=D(2026, 10, 31), first=D(2026, 10, 1),
                     last=D(2026, 10, 4), date_from=None, date_to=None)
    assert (p["to"], p["days_to"]) == (D(2026, 10, 4), D(2026, 10, 31))


def test_custom_period_is_kept():
    p = CR.period_of(start=D(2026, 10, 1), end=D(2026, 10, 31), first=D(2026, 10, 1),
                     last=D(2026, 10, 20), date_from=D(2026, 10, 5), date_to=D(2026, 10, 11))
    assert (p["from"], p["to"], p["days_to"]) == (D(2026, 10, 5), D(2026, 10, 11), D(2026, 10, 11))


def test_period_after_data_is_refused(client):
    db = SessionLocal()
    _, deal_id = _camp_with_fact(db)
    db.close()
    app.dependency_overrides[get_current_user] = lambda: _User("admin")
    r = client.get(f"/api/client-report/deal/{deal_id}.xlsx?date_from=2030-01-01")
    assert r.status_code == 409 and "период" in r.json()["detail"]


def test_days_are_full_calendar_month_for_short_campaign():
    """Владелец 05.10.2026: РК на 15 дней «поехала» — шаблон рассчитан на полный месяц.
    Колонки дней — весь календарный месяц; дни вне РК пустые, а не нули."""
    from app.ad import client_report_xlsx as X
    head = {"period_from": date(2026, 10, 15), "period_to": date(2026, 10, 18),
            "days_to": date(2026, 10, 29), "last_data": date(2026, 10, 18),
            "date_start": date(2026, 10, 15), "date_end": date(2026, 10, 29),
            "advertiser": "А", "code": "TEST01"}
    tot = {"shows": 40, "clicks": 1, "ctr": 2.5, "plan": None, "done_pct": None}
    pl = [{"label": "site.ru", "shows": 40, "clicks": 1, "uniques": None}]
    days = [{"label": date(2026, 10, d), "shows": 10, "clicks": 0} for d in (15, 16, 17, 18)]
    rep = CR.Report(head=head, totals=tot, by_placement=pl, by_day=days)
    got = X._days_of(rep)
    assert got[0] == date(2026, 10, 1) and got[-1] == date(2026, 10, 31) and len(got) == 31
    wb = openpyxl.load_workbook(io.BytesIO(X.to_xlsx(rep, rep)))
    ws = wb.worksheets[0]
    vals = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value == "01.10":
                head_row, first_col = c.row, c.column
                vals = {ws.cell(head_row, first_col + i).value: ws.cell(head_row + 1, first_col + i).value
                        for i in range(31)}
                break
        if vals:
            break
    assert vals, "колонки 01.10 нет"
    assert vals["01.10"] is None and vals["14.10"] is None, "до старта — пусто, не ноль"
    assert vals["15.10"] == 10 and vals["31.10"] is None
