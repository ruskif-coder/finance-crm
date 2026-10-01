# -*- coding: utf-8 -*-
"""Паспорт по всей РК (владелец 30.09.2026): .xlsx без креативов; внешний тег пикселя
попадает в колонку «Пиксель Weborama» (раньше колонка брала только пиксель размещения)."""
import io
from types import SimpleNamespace as NS

from openpyxl import load_workbook
from sqlalchemy import text

import app.main  # noqa: F401
from app.database import SessionLocal
from app.traffic import offsite_export as OX


def test_pixel_for_external_tag_comes_from_deal():
    deal = NS(weborama_pixel=True, weborama_pixel_mode="external", weborama_pixel_tag="<img src=x>")
    assert OX.pixel_for(deal, NS(weborama_pixel="own")) == "<img src=x>"


def test_pixel_for_own_and_off():
    assert OX.pixel_for(NS(weborama_pixel=True, weborama_pixel_mode="own", weborama_pixel_tag=None),
                        NS(weborama_pixel="P")) == "P"
    assert OX.pixel_for(NS(weborama_pixel=False, weborama_pixel_mode="own", weborama_pixel_tag=None),
                        None) == "не нужен"


def test_whole_passport_is_xlsx_over_all_publishers():
    db = SessionLocal()
    try:
        did = db.execute(text(
            "SELECT s.deal_id FROM launch_prep_pair p JOIN launch_prep_creative_set s ON s.id=p.set_id "
            "WHERE p.sent_at IS NOT NULL GROUP BY s.deal_id ORDER BY count(*) DESC LIMIT 1")).scalar()
        if not did:
            import pytest
            pytest.skip("нет отправленных пар")
        from app.sales.models import SalesDeal
        body, name = OX.build(db, db.get(SalesDeal, did), whole=True)
        assert name.endswith(".xlsx")
        ws = load_workbook(io.BytesIO(body)).active
        assert ws.max_row > 1 and ws.cell(1, 20).value == "Пиксель Weborama"
    finally:
        db.close()


def test_external_tag_marks_every_placement_pixel_ready():
    """Внешний тег сделки — у всех площадок «пиксель получен», вставки не нужны."""
    from app.ad import external as E
    p = NS(is_direct=False, id=1, weborama_pixel=None, status="ждёт запуска")
    st = E._state_of(p, {}, set(), [], ext_tag=True)
    assert st["weborama"]["state"] == E.READY
    assert E._state_of(p, {}, set(), [])["weborama"]["state"] == E.MISSING


def test_rk_holds_only_deal_recipients_not_archived():
    """Состав РК = получатели сделки без архивных (владелец 30.09.2026)."""
    from app.ad import build as B
    db = SessionLocal()
    db.commit = db.flush
    try:
        cid, did = db.execute(text(
            "SELECT a.id, a.deal_id FROM ad_campaign a WHERE EXISTS (SELECT 1 FROM launch_prep_target t "
            "WHERE t.deal_id = a.deal_id) ORDER BY a.id LIMIT 1")).first()
        B.sync_placements(db, db.get(B.AdCampaign, cid), commit=False)
        have = {r[0] for r in db.execute(text(
            "SELECT publisher_id FROM ad_campaign_placement WHERE campaign_id = :c AND status = :w "
            "AND NOT EXISTS (SELECT 1 FROM ad_campaign_creative c WHERE c.placement_id = ad_campaign_placement.id)"),
            {"c": cid, "w": B.PLACEMENT_WAIT})}
        assert have <= B.deal_publishers(db, did)
        arch = {r[0] for r in db.execute(text("SELECT id FROM sales_publishers WHERE status = 'АРХИВ'"))}
        assert not (B.deal_publishers(db, did) & arch)
    finally:
        db.rollback()
        db.close()


def test_passport_pixel_macro_follows_channel():
    """Итоговый тег — под канал площадки: DSP — {RND}, Adfox — %system.random% (30.09.2026)."""
    raw = "https://wcm.weborama-tech.ru/fcgi-bin/dispatch.fcgi?a.A=im&a.wi=~WIDTH~&a.he=~HEIGHT~&g.lu=[RANDOM]"
    deal = NS(weborama_pixel=True, weborama_pixel_mode="own", weborama_pixel_tag=None)
    pl = NS(weborama_pixel=raw)
    dsp = OX.pixel_for(deal, pl, NS(domain="a.ru", our_code=True), "dsp", "240x400")
    adf = OX.pixel_for(deal, pl, NS(domain="b.ru", our_code=False), "adfox", "240x400")
    assert "{RND}&a.ycp=https://a.ru" in dsp and "a.wi=240" in dsp
    assert "%system.random%&a.ycp=https://b.ru" in adf
    assert "[RANDOM]" not in dsp + adf
    assert dsp != adf



def test_passport_volume_is_per_pair_not_whole_placement():
    """«План показов» — объём пары «креатив × площадка» (владелец 01.10.2026), тот же, что
    уходит лимитом креатива в DSP (`build.creative_plans`), а не весь план площадки.
    Замер 01.10: у Maksavit.ru в DLMBGB два креатива, а в строке стоял весь план — 214 836."""
    from app.ad.build import creative_plans
    from app.ad.models import AdCampaign
    from app.sales.models import SalesDeal
    db = SessionLocal()
    try:
        did = db.execute(text("""
            SELECT a.deal_id FROM ad_campaign_creative c JOIN ad_campaign a ON a.id = c.campaign_id
              JOIN ad_campaign_placement pl ON pl.id = c.placement_id
             WHERE c.pair_id IS NOT NULL AND pl.plan_show > 0
             GROUP BY a.deal_id, c.placement_id HAVING count(*) > 1 LIMIT 1""")).scalar()
        if not did:
            import pytest
            pytest.skip("нет площадки с несколькими креативами и планом")
        want = {}
        for camp in db.query(AdCampaign).filter(AdCampaign.deal_id == did):
            plans = creative_plans(db, camp)
            for cid, pair_id in db.execute(text(
                    "SELECT id, pair_id FROM ad_campaign_creative WHERE campaign_id = :c"), {"c": camp.id}):
                if pair_id and plans.get(cid) is not None:
                    want[pair_id] = round(plans[cid])
        got = OX.pair_volumes(db, db.get(SalesDeal, did))
        assert want and all(got.get(k) == v for k, v in want.items()), (want, got)
        # И паспорт берёт именно их: строк с целым планом площадки при двух креативах нет.
        whole = {r[0] for r in db.execute(text("""
            SELECT pl.plan_show FROM ad_campaign_placement pl JOIN ad_campaign a ON a.id = pl.campaign_id
             WHERE a.deal_id = :d AND pl.plan_show > 0
               AND (SELECT count(*) FROM ad_campaign_creative c WHERE c.placement_id = pl.id) > 1
        """), {"d": did})}
        body, _ = OX.build(db, db.get(SalesDeal, did), whole=True)
        ws = load_workbook(io.BytesIO(body)).active
        pi = [c.value for c in ws[1]].index("План показов")
        vals = [r[pi] for r in ws.iter_rows(min_row=2, values_only=True)]
        assert not ({round(w) for w in whole} & set(vals) - set(want.values()))
    finally:
        db.close()
