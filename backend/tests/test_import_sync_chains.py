# -*- coding: utf-8 -*-
"""Прибор: сверка импорта не склеивает цепочку частичных оплат (аудит 23.09.2026, 2.H1).

ЧТО БЫЛО. Для строк с № ДС и Счётом существующие операции клались в словарь по паре
номеров — оставалась одна, последняя по id. Проверено на стенде: выгрузка операций
одного контрагента, загруженная обратно БЕЗ ЕДИНОЙ ПРАВКИ, давала «Конфликтов: 5». Каждый транш
цепочки «прилож. 9 / 424» сравнивался с одной и той же операцией на 100 000, у всех был
ОДИН ключ подтверждения, и одна галка переписала бы эту операцию три раза подряд.

ЧТО ДЕРЖИМ:
  · выгрузка, загруженная обратно как есть, — ноль конфликтов в цепочке;
  · правка одного транша — ровно один конфликт, с ТОЙ операцией, и применение меняет её одну;
  · без колонки ID строка находит свою операцию по сумме и статусу;
  · одинаковые номера у ДРУГОГО контрагента — другая операция, а не конфликт;
  · сопоставить однозначно нечем — строка помечается «неоднозначной» и не применяется.

Файл строится настоящей выгрузкой (`/operations/export`) и читается настоящим разбором —
проверяется весь круг, а не функция сопоставления в отрыве от формата.
"""
import asyncio
import io
from datetime import date

import openpyxl
import pytest
from fastapi import UploadFile

import app.ad.models           # noqa: F401
import app.launch_prep.models  # noqa: F401
import app.notify.models       # noqa: F401
import app.ord.models          # noqa: F401
from app.database import SessionLocal
from app.models import Article, Counterparty, Operation, User
from app.routers import operations as ops

# Своя цепочка в транзакции теста — по образцу «прилож. 9 / 424» стенда (у одного контрагента: четыре
# оплаты под одной парой номеров). До ревью 23.09.2026 тесты брали ту цепочку со стенда, и на
# проде и свежей базе все они молча пропускались.
CHAIN, INV = "прибор прилож. 9", "424"


@pytest.fixture
def db():
    s = SessionLocal()
    s.commit = s.flush
    try:
        yield s
    finally:
        s.rollback()
        s.close()


@pytest.fixture
def admin(db):
    u = (db.query(User).filter(User.is_active == 1, User.role.has(key="admin"))
         .order_by(User.id).first())
    if u is None:
        pytest.skip("нужна учётка админа")
    return u


@pytest.fixture
def chain(db):
    art = db.query(Article).order_by(Article.id).first()
    if art is None:
        pytest.skip("нужна хотя бы одна статья")
    cp = Counterparty(name="прибор цепочки оплат", inn="7799000001", status="действующий")
    db.add(cp)
    db.flush()
    rows = []
    for i, (status, amount) in enumerate((("ОПЛАЧЕНО", 100_000.0), ("ОПЛАЧЕНО", 50_000.0),
                                          ("ОПЛАЧЕНО", 30_000.0),
                                          ("ПЛАН ПОСТУПЛЕНИЙ", 20_000.0))):
        op = Operation(date=date(2099, 1, 10 + i), status=status, income=amount, expense=0.0,
                       bank="АльфаБанк" if status == "ОПЛАЧЕНО" else None, period="2099-01",
                       vat_rate=0, vat_fact=0, article_id=art.id, counterparty_id=cp.id,
                       ds_num=CHAIN, invoice=INV, description=f"транш {i + 1}",
                       parent_operation_id=rows[0].id if rows else None)
        db.add(op)
        db.flush()
        rows.append(op)
    return rows


def _export(db, admin, counterparty_id) -> bytes:
    resp = ops.export_operations(status=None, bank=None, date_from=None, date_to=None,
                                 article_id=None, counterparty_id=[counterparty_id],
                                 period=None, gaps=None, db=db, current_user=admin)

    async def read():
        return b"".join([c async for c in resp.body_iterator])
    return asyncio.run(read())


def _edit(data: bytes, fn) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    fn(ws, hdr)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _preview(db, admin, data: bytes):
    up = UploadFile(filename="sync.xlsx", file=io.BytesIO(data))
    res = asyncio.run(ops.import_preview(file=up, db=db, current_user=admin))
    cache = ops.IMPORT_SYNC_CACHE[res["import_id"]]["rows"]
    return res, cache


def _chain_rows(cache, chain_ids):
    return [c for c in cache if c["data"].get("op_id") in chain_ids
            or (_ours(c) and c["data"].get("op_id") is None)]


def _ours(c):
    return c["data"].get("ds_num") == CHAIN and str(c["data"].get("invoice")) == INV


def test_export_reimported_as_is_has_no_conflicts_in_the_chain(db, admin, chain):
    ids = {o.id for o in chain}
    _, cache = _preview(db, admin, _export(db, admin, chain[0].counterparty_id))
    mine = _chain_rows(cache, ids)
    assert len(mine) == len(chain)
    assert [c["status"] for c in mine] == ["unchanged"] * len(chain), mine
    # И каждая строка нашла СВОЮ операцию, а не одну на всех.
    assert {c["existing_id"] for c in mine} == ids


def test_editing_one_tranche_gives_one_conflict_with_that_operation(db, admin, chain):
    target = chain[1]

    def edit(ws, hdr):
        i_id, i_desc = hdr.index("ID"), hdr.index("Назначение")
        for row in ws.iter_rows(min_row=2):
            if row[i_id].value == target.id:
                row[i_desc].value = "правка одного транша"
    res, cache = _preview(db, admin, _edit(_export(db, admin, target.counterparty_id), edit))
    conflicts = [c for c in _chain_rows(cache, {o.id for o in chain}) if c["status"] == "conflict"]
    assert len(conflicts) == 1 and conflicts[0]["existing_id"] == target.id
    # Ключ подтверждения у строки свой: одна галка — одна операция.
    keys = [c["key"] for c in cache if c["status"] == "conflict"]
    assert len(keys) == len(set(keys))

    before = {o.id: (o.income, o.description) for o in chain}
    ops.IMPORT_SYNC_CACHE[res["import_id"]]["rows"] = cache
    asyncio.run(ops.import_apply(ops.ImportApplyRequest(import_id=res["import_id"],
                                                        confirmed_keys=[conflicts[0]["key"]]),
                                 db=db, current_user=admin))
    db.expire_all()
    for o in db.query(Operation).filter(Operation.id.in_(list(before))).all():
        if o.id == target.id:
            assert o.description == "правка одного транша"
            assert o.income == before[o.id][0]
        else:
            assert (o.income, o.description) == before[o.id], f"задета чужая операция #{o.id}"


def test_rows_without_id_find_their_tranche_by_amount(db, admin, chain):
    def drop_ids(ws, hdr):
        i_id = hdr.index("ID")
        for row in ws.iter_rows(min_row=2):
            row[i_id].value = None
    _, cache = _preview(db, admin, _edit(_export(db, admin, chain[0].counterparty_id), drop_ids))
    mine = [c for c in cache if _ours(c)]
    assert [c["status"] for c in mine] == ["unchanged"] * len(chain)
    assert {c["existing_id"] for c in mine} == {o.id for o in chain}


def test_foreign_id_does_not_steal_an_operation(db, admin, chain):
    """ID из файла — подсказка, не приговор: файл мог прийти с другого стенда, где этот
    номер — чужая операция. Строка с номерами цепочки и ID чужой операции сопоставляется
    по содержанию, а чужая операция не трогается."""
    stranger = db.query(Operation).filter(Operation.ds_num != CHAIN).order_by(Operation.id).first()  # noqa: E501
    target = chain[0]

    def swap(ws, hdr):
        i_id = hdr.index("ID")
        for row in ws.iter_rows(min_row=2):
            if row[i_id].value == target.id:
                row[i_id].value = stranger.id
    _, cache = _preview(db, admin, _edit(_export(db, admin, target.counterparty_id), swap))
    assert all(c["existing_id"] != stranger.id for c in cache)
    mine = [c for c in cache if _ours(c)]
    assert {c["existing_id"] for c in mine} == {o.id for o in chain}


def test_same_numbers_at_another_counterparty_are_a_new_operation(db, admin, chain):
    # «Другая фирма» доказывается только ИНН: название в файле свободный текст, и его
    # несовпадение даёт «неоднозначную» строку (test_import_match_counterparty).
    own_inn = (chain[0].counterparty.inn or "").strip() if chain[0].counterparty else ""
    if not own_inn:
        pytest.skip("у контрагента цепочки нет ИНН — другую фирму нечем доказать")
    other = (db.query(Counterparty).filter(Counterparty.id != chain[0].counterparty_id,
                                           Counterparty.inn.isnot(None), Counterparty.inn != "",
                                           Counterparty.inn != own_inn)
             .order_by(Counterparty.id).first())

    def move(ws, hdr):
        i_id, i_cp, i_inn = hdr.index("ID"), hdr.index("Контрагент"), hdr.index("ИНН")
        for row in ws.iter_rows(min_row=2):
            if row[i_id].value == chain[0].id:
                row[i_id].value, row[i_cp].value, row[i_inn].value = None, other.name, other.inn
    _, cache = _preview(db, admin, _edit(_export(db, admin, chain[0].counterparty_id), move))
    moved = [c for c in cache if c["data"].get("counterparty") == other.name and _ours(c)]
    assert len(moved) == 1 and moved[0]["status"] == "new"


def test_undecidable_rows_are_marked_and_not_applied(db, admin):
    art = db.query(Article).order_by(Article.id).first()
    cp = db.query(Counterparty).order_by(Counterparty.id).first()
    for amount in (100.0, 200.0):
        db.add(Operation(date=date(2099, 2, 1), status="ОПЛАЧЕНО", income=amount, expense=0.0,
                         bank="АльфаБанк", period="2099-02", vat_rate=0, vat_fact=0,
                         article_id=art.id, counterparty_id=cp.id, ds_num="прибор неясн.",
                         invoice="1"))
    db.flush()
    data = _export(db, admin, cp.id)

    def blur(ws, hdr):
        i_id, i_ds, i_inc = hdr.index("ID"), hdr.index("№ ДС"), hdr.index("Поступления")
        for row in ws.iter_rows(min_row=2):
            if row[i_ds].value == "прибор неясн.":
                row[i_id].value = None
                row[i_inc].value = (row[i_inc].value or 0) + 50     # 150 и 250: ни одна не совпадает
    res, cache = _preview(db, admin, _edit(data, blur))
    mine = [c for c in cache if c["data"].get("ds_num") == "прибор неясн."]
    assert [c["status"] for c in mine] == ["ambiguous", "ambiguous"]
    assert res["summary"]["ambiguous"] >= 2
    assert all(len(a["candidates"]) == 2 for a in res["ambiguous"]
               if a["incoming"].get("ds_num") == "прибор неясн.")


def test_new_tranche_from_a_file_joins_its_chain(db, admin, chain):
    """В файле на транш больше, чем в базе: все операции цепочки заняты своими строками,
    и лишняя строка — новая. Она заводится ЧАСТЬЮ той же цепочки, а не одиночкой рядом
    (ревью 23.09.2026): иначе разметка цепочек её потом не подхватит — в группе «кто-то уже
    связан»."""
    def add_tranche(ws, hdr):
        i_id, i_inc = hdr.index("ID"), hdr.index("Поступления")
        src = next(r for r in ws.iter_rows(min_row=2) if r[i_id].value == chain[1].id)
        vals = [c.value for c in src]
        vals[i_id], vals[i_inc] = None, 7_777.0
        ws.append(vals)
    res, cache = _preview(db, admin, _edit(_export(db, admin, chain[0].counterparty_id), add_tranche))
    new = [c for c in cache if _ours(c) and c["status"] == "new"]
    assert len(new) == 1, [c["status"] for c in cache if _ours(c)]
    ops.IMPORT_SYNC_CACHE[res["import_id"]]["rows"] = cache
    asyncio.run(ops.import_apply(ops.ImportApplyRequest(import_id=res["import_id"],
                                                        confirmed_keys=[]),
                                 db=db, current_user=admin))
    made = (db.query(Operation).filter(Operation.ds_num == CHAIN, Operation.income == 7_777.0)
            .one())
    assert made.parent_operation_id == chain[0].id
