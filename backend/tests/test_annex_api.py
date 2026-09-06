# -*- coding: utf-8 -*-
"""Приложения к договору: черновик, подтверждение, занятый номер.

Смысл разделения на черновик и подтверждение: брошенная сборка не должна сжигать номер,
а занятый номер не должен меняться задним числом. Оба края тихие — ошибка вылезает у
клиента, который получил второе приложение с тем же номером.
"""
from datetime import date
from types import SimpleNamespace

import pytest
from fastapi import HTTPException

import app.launch_prep.models  # noqa: F401
import app.ord.models          # noqa: F401
from app.database import SessionLocal
from app.models import Contract, User
from app.routers import annexes as A
from sqlalchemy import func

from app.sales.models import SalesAnnex, SalesDealAnnexAllocation as Alloc


@pytest.fixture(autouse=True)
def no_leftovers():
    """Ни один тест не оставляет приложений на стенде — и это проверяется, а не имеется
    в виду.

    Прогон 05.09.2026 дважды оставил черновики, и обнаружилось это счётом строк вручную,
    а не тестом. Такие следы опасны не сами по себе: приложение занимает номер в договоре,
    и следующий выпуск по нему поедет. Уборка внутри теста через `finally` перестаёт
    работать ровно там, где нужна, — при падении до присвоения переменной.
    """
    db = SessionLocal()
    high = db.query(func.max(SalesAnnex.id)).scalar() or 0
    db.close()
    yield
    db = SessionLocal()
    left = db.query(SalesAnnex).filter(SalesAnnex.id > high).all()
    ids = [a.id for a in left]
    if ids:
        db.query(Alloc).filter(Alloc.annex_id.in_(ids)).delete(synchronize_session=False)
        db.query(SalesAnnex).filter(SalesAnnex.id.in_(ids)).delete(synchronize_session=False)
        db.commit()
    db.close()
    assert not ids, f'тест оставил приложения на стенде: {ids}'


@pytest.fixture
def env():
    db = SessionLocal()
    c = db.query(Contract).order_by(Contract.id).first()
    u = db.query(User).filter(User.is_active == 1).order_by(User.id).first()
    if not c or not u:
        db.close()
        pytest.skip('нужны договор и активная учётка')
    was_start = c.annex_start_no
    actor = SimpleNamespace(id=u.id, name=u.name,
                            role=SimpleNamespace(key='admin', is_master=True))
    yield SimpleNamespace(db=db, c=c, user=actor)
    db.query(SalesAnnex).filter(SalesAnnex.contract_id == c.id).delete()
    c.annex_start_no = was_start
    db.commit()
    db.close()


def _draft(env, amount=732000, pf=date(2026, 6, 1), pt=date(2026, 6, 30)):
    return A.create_annex(A.AnnexIn(contract_id=env.c.id, period_from=pf, period_to=pt,
                                    total_amount=amount), env.db, env.user)


def test_preview_writes_nothing_and_burns_no_number(env):
    """Предпросмотр — не действие. Если бы он заводил строку, каждая примерка съедала бы
    номер, и нумерация поехала бы у клиента."""
    before = env.db.query(SalesAnnex).count()
    p = A.preview(contract_id=env.c.id, period_from=date(2026, 6, 1),
                  period_to=date(2026, 6, 30), amount=732000, db=env.db, user=env.user)
    assert p["no"] and p["date"] == date(2026, 5, 31)
    assert env.db.query(SalesAnnex).count() == before


def test_draft_has_no_number_until_confirmed(env):
    """Черновиков может быть сколько угодно, и номера у них нет: иначе брошенная сборка
    навсегда занимает место в нумерации договора."""
    a = _draft(env)
    b = _draft(env)
    assert a["no"] is None and b["no"] is None
    assert a["is_draft"] and b["is_draft"]


def test_confirm_takes_the_next_number_within_the_contract(env):
    """Стартовый номер — последний, выданный вне системы. Первое наше идёт следом."""
    A.set_start_no(env.c.id, A.StartNoIn(annex_start_no=73), env.db, env.user)
    a = A.confirm_annex(_draft(env)["id"], A.ConfirmIn(), env.db, env.user)
    b = A.confirm_annex(_draft(env)["id"], A.ConfirmIn(), env.db, env.user)
    assert (a["no"], b["no"]) == (74, 75)
    assert a["number"] == "Приложение № 74"


def test_confirmed_annex_is_frozen(env):
    """Документ ушёл клиенту. Правка задним числом означала бы два разных приложения под
    одним номером — и второе клиент не видел."""
    a = A.confirm_annex(_draft(env)["id"], A.ConfirmIn(), env.db, env.user)
    with pytest.raises(HTTPException) as e1:
        A.edit_annex(a["id"], A.AnnexIn(contract_id=env.c.id, period_from=date(2026, 7, 1),
                                        period_to=date(2026, 7, 31), total_amount=1),
                     env.db, env.user)
    assert e1.value.status_code == 400
    with pytest.raises(HTTPException) as e2:
        A.confirm_annex(a["id"], A.ConfirmIn(), env.db, env.user)
    assert e2.value.status_code == 400


def test_taken_number_is_refused_with_the_number_in_the_message(env):
    """Уникальность держит база, а не проверка в коде: между сборкой и подтверждением
    номер мог занять кто-то другой. В тексте отказа обязано стоять само число — после
    отката объект сброшен, и «Номер None уже занят» не сказал бы ничего.
    """
    A.confirm_annex(_draft(env)["id"], A.ConfirmIn(no=5), env.db, env.user)
    with pytest.raises(HTTPException) as e:
        A.confirm_annex(_draft(env)["id"], A.ConfirmIn(no=5), env.db, env.user)
    assert e.value.status_code == 409
    assert "5" in str(e.value.detail) and "None" not in str(e.value.detail)


def test_editing_a_draft_recomputes_the_date(env):
    """Дата считается от периода, а не вводится: сдвинули размещение на июль — дата
    обязана стать 30 июня, иначе документ датирован не тем месяцем."""
    a = _draft(env)
    out = A.edit_annex(a["id"], A.AnnexIn(contract_id=env.c.id,
                                          period_from=date(2026, 7, 1),
                                          period_to=date(2026, 7, 31),
                                          total_amount=800000), env.db, env.user)
    assert out["date"] == date(2026, 6, 30)


def test_backwards_period_is_refused_everywhere(env):
    """Конец раньше начала — опечатка, и принять её значит выпустить документ на период
    отрицательной длины."""
    with pytest.raises(HTTPException):
        A.preview(contract_id=env.c.id, period_from=date(2026, 6, 30),
                  period_to=date(2026, 6, 1), amount=1, db=env.db, user=env.user)
    with pytest.raises(HTTPException):
        _draft(env, pf=date(2026, 6, 30), pt=date(2026, 6, 1))


def test_number_hint_never_becomes_the_start_number_by_itself(env):
    """Подсказка из старых операций справочная: там встречается мусор вроде 590425.
    Стартовый номер остаётся пустым, пока его не проставит человек."""
    A.set_start_no(env.c.id, A.StartNoIn(annex_start_no=None), env.db, env.user)
    p = A.preview(contract_id=env.c.id, period_from=date(2026, 6, 1),
                  period_to=date(2026, 6, 30), amount=1, db=env.db, user=env.user)
    assert p["annex_start_no"] is None
    assert p["no"] == 1, 'без стартового номера нумерация начинается с единицы'


# ── имя файла ────────────────────────────────────────────────────────────────

def test_file_name_carries_all_four_parts():
    """Правило владельца 05.09.2026: номер, дата, рекламодатель с брендом, юрлицо по
    договору. По имени документ находят в папке, не открывая его."""
    n = A.file_name(no=68, doc_date=date(2026, 5, 31),
                    brand="Dr. Reddy / Хелинорм", payer="ДИДЖИТАЛ ИНСПИРЕЙШН ООО")
    assert n.endswith(".pdf")
    for part in ("ДС №68", "31.05.2026", "Хелинорм", "ДИДЖИТАЛ ИНСПИРЕЙШН ООО"):
        assert part in n


def test_slash_from_the_brand_never_reaches_the_file_name():
    """Бренд собирается как «Рекламодатель / Бренд». Слэш в имени файла — это путь:
    документ не сохранился бы вовсе, и выглядело бы это как поломка выгрузки."""
    n = A.file_name(no=1, doc_date=date(2026, 5, 31),
                    brand="Dr. Reddy / Хелинорм", payer='ООО "Ромашка"')
    assert not set(n) & {"/", chr(92), ":", "*", "?", '"', "<", ">", "|"}


def test_empty_parts_drop_out_instead_of_leaving_a_gap():
    """У черновика номера ещё нет, у части контрагентов не заполнен бренд. Пустая часть
    выпадает целиком — иначе в имени повисает разделитель без содержимого."""
    n = A.file_name(no=None, doc_date=date(2026, 5, 31), brand=None, payer="ООО Ромашка")
    assert n == "ДС без номера от 31.05.2026 — ООО Ромашка.pdf"
    assert A.file_name(no=7, doc_date=None, brand=None, payer=None) == "ДС №7.pdf"
    assert A.file_name(no=None, doc_date=None, brand=None, payer=None) == "ДС без номера.pdf"


def test_file_name_stays_within_the_filesystem_limit():
    """255 байт — предел почти везде, а кириллица в UTF-8 занимает по два."""
    n = A.file_name(no=68, doc_date=date(2026, 5, 31), brand="Бренд " * 40,
                    payer="Контрагент " * 40)
    assert len(n.encode("utf-8")) < 255


# ── проверка реквизитов ──────────────────────────────────────────────────────

def test_export_is_refused_while_requisites_are_missing(env):
    """Бумага с прочерком там, где стоит подпись, уходит клиенту и подписывается не
    глядя. Проверка стоит на ВЫГРУЗКЕ, а не на подтверждении: номер занять заранее можно,
    напечатать документ с пустым местом — нет."""
    a = _draft(env)
    with pytest.raises(HTTPException) as e:
        A.annex_pdf(a["id"], env.db, env.user)
    assert e.value.status_code == 400
    # Отказ обязан называть недостающее поимённо, иначе он означает «сходи поищи».
    assert "должность подписанта" in str(e.value.detail) \
        or "ФИО подписанта" in str(e.value.detail) \
        or "основание полномочий" in str(e.value.detail) \
        or "ИНН" in str(e.value.detail) or "бренд" in str(e.value.detail)


def test_signer_is_saved_from_the_annex_screen_and_lifts_the_block(env):
    """Реквизиты подписанта правятся на экране ДС и ложатся в карточку контрагента
    (решение владельца 05.09.2026): к моменту сборки выясняется, что должности нет, и
    уходить за ней в справочник — терять место."""
    from app.models import Counterparty
    cp = (env.db.query(Counterparty)
          .filter(Counterparty.id == env.c.counterparty_id).first())
    if not cp:
        pytest.skip('нужен договор с контрагентом')
    was = (cp.director_name, cp.signer_position, cp.signer_basis)
    try:
        a = _draft(env)
        before = A.get_annex(a["id"], env.db, env.user)["doc"]["missing"]
        A.set_signer(cp.id, A.SignerIn(director_name="Иванов Иван Иванович",
                                       signer_position="Генеральный директор",
                                       signer_basis="Устав"), env.db, env.user)
        after = A.get_annex(a["id"], env.db, env.user)["doc"]["missing"]
        assert len(after) < len(before)
        for gone in ("заказчик: должность подписанта", "заказчик: основание полномочий",
                     "заказчик: ФИО подписанта"):
            assert gone not in after
    finally:
        cp.director_name, cp.signer_position, cp.signer_basis = was
        env.db.commit()


def test_signer_patch_touches_only_what_was_sent(env):
    """Общая правка контрагента перезаписывает карточку целиком. Если бы экран ДС ходил
    туда, частичный запрос стёр бы адрес, банк и телефон."""
    from app.models import Counterparty
    cp = (env.db.query(Counterparty)
          .filter(Counterparty.id == env.c.counterparty_id).first())
    if not cp:
        pytest.skip('нужен договор с контрагентом')
    was_addr, was_pos = cp.address, cp.signer_position
    try:
        A.set_signer(cp.id, A.SignerIn(signer_position="Директор"), env.db, env.user)
        env.db.refresh(cp)
        assert cp.signer_position == "Директор"
        assert cp.address == was_addr, 'непереданное поле обязано остаться нетронутым'
    finally:
        cp.signer_position = was_pos
        env.db.commit()


def test_the_document_is_built_once_for_the_screen_and_for_the_paper(env):
    """Экран сборки и печать берут ОДИН расчёт. Считай экран сам — аккаунт проверял бы
    одни числа, а в бумагу уходили другие, и разошлись бы они молча."""
    a = _draft(env)
    from_screen = A.get_annex(a["id"], env.db, env.user)["doc"]
    stored = A._annex(env.db, a["id"])
    for_paper = A._doc(env.db, stored)
    assert from_screen["amount"] == for_paper["amount"]
    assert from_screen["vat"] == for_paper["vat"]
    assert from_screen["missing"] == for_paper["missing"]


# ── создание из сделки ───────────────────────────────────────────────────────

def _deal_ready_for_annex(db):
    """Сделка, из которой ДС собирается целиком: плательщик с договором и период."""
    from sqlalchemy import text as sa_text
    return db.execute(sa_text("""
        SELECT d.id FROM sales_deals d
          JOIN contracts c ON c.counterparty_id = d.payer_counterparty_id
         WHERE d.period_from IS NOT NULL
           AND COALESCE(d.amount_with_vat, d.amount) > 0
         ORDER BY d.id LIMIT 1
    """)).scalar()


def test_annex_is_assembled_from_the_deal_without_asking_again(env):
    """Кнопка «Создать» в документах сделки. Плательщик даёт договор, период берётся у
    сделки, сумма — её собственная: переспрашивать это формой значит просить человека
    ввести то, что система уже знает."""
    from app.sales.models import SalesAnnex, SalesDealAnnexAllocation as Alloc
    did = _deal_ready_for_annex(env.db)
    if not did:
        pytest.skip('на стенде нет сделки с плательщиком-договором и периодом')
    out = A.create_from_deal(did, env.db, env.user)
    try:
        assert out["is_draft"] and out["period_from"] and out["total_amount"] > 0
        # Дата документа — последний день месяца ПЕРЕД периодом, как и везде.
        assert out["date"].month != out["period_from"].month
    finally:
        env.db.query(Alloc).filter(Alloc.annex_id == out["id"]).delete()
        env.db.query(SalesAnnex).filter(SalesAnnex.id == out["id"]).delete()
        env.db.commit()


def test_the_link_to_the_deal_is_actually_written(env):
    """До 05.09.2026 `deal_ids` участвовал только в сборке предпросмотра и не сохранялся:
    приложение выпускалось, а связь со сделкой не возникала — вернуть выручку авторам
    было нечем, и таблица размещения в документе бралась из воздуха."""
    from app.sales.models import SalesAnnex, SalesDealAnnexAllocation as Alloc
    did = _deal_ready_for_annex(env.db)
    if not did:
        pytest.skip('на стенде нет подходящей сделки')
    out = A.create_from_deal(did, env.db, env.user)
    try:
        rows = env.db.query(Alloc).filter(Alloc.annex_id == out["id"]).all()
        assert [(r.deal_id, r.amount) for r in rows] == [(did, out["total_amount"])]
    finally:
        env.db.query(Alloc).filter(Alloc.annex_id == out["id"]).delete()
        env.db.query(SalesAnnex).filter(SalesAnnex.id == out["id"]).delete()
        env.db.commit()


def test_missing_source_data_is_named_not_guessed(env):
    """Приложение не на тот договор исправляется только выпуском нового номера, поэтому
    нехватка исходных данных — отказ с объяснением, а не умолчание."""
    from app.sales.models import SalesDeal
    d = (env.db.query(SalesDeal)
         .filter(SalesDeal.payer_counterparty_id.is_(None),
                 SalesDeal.counterparty_id.is_(None)).first())
    if d:
        with pytest.raises(HTTPException) as e:
            A.create_from_deal(d.id, env.db, env.user)
        assert e.value.status_code == 400 and "плательщик" in str(e.value.detail)
    with pytest.raises(HTTPException) as e2:
        A.create_from_deal(10**9, env.db, env.user)
    assert e2.value.status_code == 404


# ── номер и дата руками ──────────────────────────────────────────────────────

def test_number_below_the_start_is_refused_with_both_numbers_named(env):
    """Стартовый номер — последний, выданный ВНЕ системы. Занять его или меньший значит
    выпустить второй документ под числом, которое у клиента уже есть, и заметит это
    клиент. В отказе стоят оба числа: человек в этот момент выбирает номер руками."""
    A.set_start_no(env.c.id, A.StartNoIn(annex_start_no=73), env.db, env.user)
    for bad in (73, 1, 0, -5):
        with pytest.raises(HTTPException) as e:
            A.confirm_annex(_draft(env)["id"], A.ConfirmIn(no=bad), env.db, env.user)
        assert e.value.status_code == 400
        if bad > 0:
            assert "73" in str(e.value.detail) and str(bad) in str(e.value.detail)
    # А следующий за стартовым проходит.
    out = A.confirm_annex(_draft(env)["id"], A.ConfirmIn(no=74), env.db, env.user)
    assert out["no"] == 74


def test_taken_number_names_the_document_that_holds_it(env):
    """«Конфликт» человеку ничего не говорит. Сказать, что 74 занят приложением от
    31.05.2026, — говорит: он видит, куда смотреть."""
    A.set_start_no(env.c.id, A.StartNoIn(annex_start_no=None), env.db, env.user)
    first = A.confirm_annex(_draft(env)["id"], A.ConfirmIn(no=74), env.db, env.user)
    with pytest.raises(HTTPException) as e:
        A.confirm_annex(_draft(env)["id"], A.ConfirmIn(no=74), env.db, env.user)
    assert e.value.status_code == 409
    assert "74" in str(e.value.detail)
    assert first["date"].strftime("%d.%m.%Y") in str(e.value.detail)


def test_date_can_be_set_by_hand_and_otherwise_stays_computed(env):
    """Правило «последний день предыдущего месяца» покрывает обычный случай, но документ
    могут датировать числом, о котором договорились. Пусто — остаётся посчитанная."""
    auto = A.confirm_annex(_draft(env)["id"], A.ConfirmIn(), env.db, env.user)
    assert auto["date"] == date(2026, 5, 31)
    manual = A.confirm_annex(_draft(env)["id"],
                             A.ConfirmIn(doc_date=date(2026, 5, 29)), env.db, env.user)
    assert manual["date"] == date(2026, 5, 29)


def test_the_screen_is_told_where_numbering_starts(env):
    """Без границы на экране отказ «номер не больше стартового» выглядит придиркой
    непонятно к чему."""
    A.set_start_no(env.c.id, A.StartNoIn(annex_start_no=73), env.db, env.user)
    out = A.get_annex(_draft(env)["id"], env.db, env.user)
    assert out["annex_start_no"] == 73 and out["next_no"] == 74


# ── экран сборки: предпросмотр и сделки ──────────────────────────────────────

def test_declension_preview_works_on_unsaved_input(env):
    """Строка «в лице …» под полями обязана меняться по мере ввода, ещё до сохранения.
    Считает её сервер: правила склонения живут в одном месте, и вторая реализация на
    клиенте разошлась бы — в предпросмотре одна фамилия, в документе другая."""
    out = A.party_phrase(A.PhraseIn(director_name="Кинчиков Павел Сергеевич",
                                    signer_position="Управляющий",
                                    signer_basis="Доверенность № 5 от 01.01.2026"), env.user)
    assert out["position_gen"] == "Управляющего"
    assert out["director_gen"] == "Кинчикова Павла Сергеевича"
    assert out["basis_gen"] == "Доверенности № 5 от 01.01.2026"
    assert out["acting"] == "действующего"
    # Пустой ввод не выдумывается: экран покажет прочерки, а не чужие формы.
    empty = A.party_phrase(A.PhraseIn(), env.user)
    assert not empty["position_gen"] and not empty["director_gen"]


def test_screen_gets_the_deals_the_annex_covers(env):
    """Из документа надо уметь вернуться к тому, из чего он собран. В адресах сделок
    используется `code`, а не id — так они выглядят везде в интерфейсе."""
    from app.sales.models import SalesAnnex, SalesDealAnnexAllocation as Alloc
    did = _deal_ready_for_annex(env.db)
    if not did:
        pytest.skip('на стенде нет подходящей сделки')
    out = A.create_from_deal(did, env.db, env.user)
    try:
        deals = A.get_annex(out["id"], env.db, env.user)["deals"]
        assert [x["id"] for x in deals] == [did]
        assert deals[0]["amount"] == out["total_amount"]
        assert "code" in deals[0] and "title" in deals[0]
    finally:
        env.db.query(Alloc).filter(Alloc.annex_id == out["id"]).delete()
        env.db.query(SalesAnnex).filter(SalesAnnex.id == out["id"]).delete()
        env.db.commit()


# ── .docx ────────────────────────────────────────────────────────────────────

def test_docx_and_pdf_are_built_from_one_and_the_same_document(env):
    """Два формата одного документа с разными числами — самая дорогая из возможных здесь
    ошибок: подписывают один, а сверяют другой. Защита в том, что считать в сборщике
    .docx нечего — он получает готовый словарь."""
    from io import BytesIO
    from docx import Document
    from app.sales import annex_docx
    a = _draft(env)
    doc = A._doc(env.db, A._annex(env.db, a["id"]))
    d = Document(BytesIO(annex_docx.build_docx(doc)))
    text = "\n".join(p.text for p in d.paragraphs)
    assert doc["amount_words"] in text
    assert doc["vat_words"] in text
    # Номер сравнивается в НЕРАЗРЫВНОЙ форме: в документе он с неразрывными пробелами и
    # дефисами, иначе Word рвёт его между строк (см. тест ниже).
    assert annex_docx._nb(doc["contract"]["number"] or "") in text


def test_docx_is_landscape_like_the_paper(env):
    """Таблица из тринадцати колонок в книжной A4 сжимается до нечитаемого."""
    from io import BytesIO
    from docx import Document
    from app.sales import annex_docx
    a = _draft(env)
    d = Document(BytesIO(annex_docx.build_docx(A._doc(env.db, A._annex(env.db, a["id"])))))
    sec = d.sections[0]
    assert sec.page_width > sec.page_height


def test_docx_is_refused_on_the_same_grounds_as_pdf(env):
    """Документ с прочерком на месте подписанта одинаково опасен в любом формате —
    значит и проверка полноты обязана быть одна, а не «для PDF строгая, для ворда нет»."""
    a = _draft(env)
    doc = A._doc(env.db, A._annex(env.db, a["id"]))
    if not doc["missing"]:
        pytest.skip('реквизиты стенда полны — отказу неоткуда взяться')
    with pytest.raises(HTTPException) as e:
        A.annex_docx(a["id"], env.db, env.user)
    assert e.value.status_code == 400


def test_docx_table_fits_inside_the_page(env):
    """Таблица шире полосы набора выезжает за правое поле — в Word это видно сразу, а в
    коде не видно вовсе. Так и случилось 05.09.2026: python-docx создаёт документ
    размером Letter, а колонки были посчитаны под A4.

    Проверяется не «ширина равна 26,5», а само отношение: сумма колонок не больше полосы.
    Тогда тест переживёт смену полей и формата."""
    from io import BytesIO
    from docx import Document
    from docx.shared import Emu
    from app.sales import annex_docx
    a = _draft(env)
    d = Document(BytesIO(annex_docx.build_docx(A._doc(env.db, A._annex(env.db, a["id"])))))
    sec = d.sections[0]
    usable = Emu(sec.page_width - sec.left_margin - sec.right_margin).cm
    assert round(sec.page_width.cm, 1) == 29.7, 'страница обязана быть A4, а не Letter'
    for t in d.tables:
        assert sum(c.width.cm for c in t.columns) <= usable + 0.01


def test_money_columns_hold_values_over_a_million(env):
    """Проверка владельца 05.09.2026: «копейки едут». Причина была не в формате числа —
    разряды и так неразрывные, — а в ширине: колонка НДС в 1,44 см ломала «90 163,97»
    после запятой, и на бумаге это читается как другое число.

    Проверяются не сегодняшние суммы, а самый длинный правдоподобный случай: миллиарды
    показов и десятки миллионов рублей. Ширина считается от кегля 7,5 pt Times, где знак
    занимает ~0,132 см, плюс поля ячейки.
    """
    from datetime import date as _d
    from io import BytesIO
    from docx import Document
    from app.sales import annex_docx
    doc = A._doc(env.db, A._annex(env.db, _draft(env, amount=12345678.90)["id"]))
    doc["rows"] = [{"network": "SIMB-AD", "position": "Альфарм-Таргет", "geo": "РФ",
                    "format": "Баннеры", "device": "Кросс-девайс", "model": "CPM",
                    "volume": 1234567890, "unit_price": 1234.56,
                    "date_from": _d(2026, 6, 1), "date_to": _d(2026, 6, 30),
                    "amount": 10119409.02, "vat": 2226269.88, "gross": 12345678.90}]
    t = Document(BytesIO(annex_docx.build_docx(doc))).tables[0]
    for i in (7, 9, 10, 11, 12):
        text = t.rows[1].cells[i].text
        need = len(text) * 0.132 + 0.38
        assert t.columns[i].width.cm >= need, \
            f'колонка {i} узка для «{text}»: нужно {need:.2f} см, есть {t.columns[i].width.cm:.2f}'


def test_contract_number_never_splits_across_lines(env):
    """«РМ 30-11-2023» Word ломает и по пробелу, и по дефису — выходит «№ РМ 30-11-» на
    одной строке и «2023» на следующей, то есть на вид два разных номера. В печатной форме
    это закрыто `white-space: nowrap`, здесь — самими символами."""
    from io import BytesIO
    from docx import Document
    from app.sales import annex_docx
    a = _draft(env)
    doc = A._doc(env.db, A._annex(env.db, a["id"]))
    num = (doc.get("contract") or {}).get("number")
    if not num or not any(ch in num for ch in " -"):
        pytest.skip('в номере договора нечего разрывать')
    d = Document(BytesIO(annex_docx.build_docx(doc)))
    text = "\n".join(p.text for p in d.paragraphs)
    unbreakable = annex_docx._nb("№ " + num)
    assert unbreakable in text
    assert f"№ {num} об оказании" not in text, 'остался разрываемый вариант'


# ── реестр и выгрузка ────────────────────────────────────────────────────────

def test_registry_row_carries_the_columns_the_screen_shows(env):
    """Реестр показывает ЧЬЁ и по какому договору раньше, чем сам документ. Если эти поля
    не приходят, колонки молча заполняются прочерками — экран выглядит рабочим, а искать
    по нему нечего."""
    _draft(env)
    row = [i for i in A.list_annexes(env.db, env.user)["items"]
           if i["contract_id"] == env.c.id][0]
    for key in ("counterparty_name", "counterparty_inn", "contract_number",
                "contract_date", "period_from", "total_amount", "deals"):
        assert key in row, f'реестру нужна колонка {key}'


def test_export_keeps_the_order_the_screen_sent(env):
    """Сортировка живёт на экране. Повторять её на сервере значило бы завести вторую
    реализацию того же правила — они расходятся молча, и выгрузка перестаёт совпадать
    с тем, что человек видел перед нажатием."""
    import asyncio
    from io import BytesIO
    from openpyxl import load_workbook
    a1, a2 = _draft(env), _draft(env, amount=111)
    r = A.export_annexes(A.ExportIn(ids=[a2["id"], a1["id"]]), env.db, env.user)

    async def grab():
        return b"".join([c async for c in r.body_iterator])
    ws = load_workbook(BytesIO(asyncio.run(grab()))).active
    amounts = [ws.cell(row=i, column=9).value for i in (2, 3)]
    assert amounts == [a2["total_amount"], a1["total_amount"]], 'порядок экрана не соблюдён'


def test_export_head_matches_the_screen_columns(env):
    """Шапка выгрузки повторяет колонки реестра: файл читают вместо экрана, и другой
    набор столбцов в нём означает другой документ."""
    import asyncio
    from io import BytesIO
    from openpyxl import load_workbook
    _draft(env)
    r = A.export_annexes(A.ExportIn(), env.db, env.user)

    async def grab():
        return b"".join([c async for c in r.body_iterator])
    ws = load_workbook(BytesIO(asyncio.run(grab()))).active
    head = [c.value for c in ws[1]]
    assert head[:5] == ["Юрлицо", "ИНН", "Договор", "Дата договора", "Приложение"]
    assert "Сделки" in head and "Статус" in head


# ── договор с экрана сборки ──────────────────────────────────────────────────

def test_document_carries_the_contract_link_and_file_flag(env):
    """Приложение проверяют рядом с договором, а не поиском по реестру. Экрану нужны две
    вещи: куда вести ссылку и есть ли что скачивать.

    Замерено 06.09.2026: ссылка в ЭДО заполнена у 45 договоров из 173, прикреплённых
    файлов нет ни одного — то есть кнопка скачивания на стенде не появится, и проверять
    её надо на обоих состояниях, а не на текущих данных.
    """
    from app.models import Contract
    a = _draft(env)
    c = env.db.query(Contract).filter(Contract.id == env.c.id).first()
    was = (c.document_link, c.attached_filename)
    try:
        c.document_link, c.attached_filename = None, None
        env.db.flush()
        out = A._doc(env.db, A._annex(env.db, a["id"]))["contract"]
        assert out["link"] is None and out["has_file"] is False

        c.document_link = "https://diadoc.kontur.ru/143bcbfb-a011-4791-9316-981158fa555"
        c.attached_filename = f"{c.id}_dogovor.pdf"
        env.db.flush()
        out = A._doc(env.db, A._annex(env.db, a["id"]))["contract"]
        assert out["link"].startswith("https://diadoc.kontur.ru/")
        assert out["has_file"] is True
        assert out["id"] == c.id, 'по id экран зовёт ручку договоров за файлом'
    finally:
        c.document_link, c.attached_filename = was
        env.db.commit()


# ── один рекламодатель на приложение ─────────────────────────────────────────

def _two_deals_of_different_advertisers(db):
    from sqlalchemy import text as sa_text
    return db.execute(sa_text("""
        SELECT (SELECT id FROM sales_deals WHERE advertiser_id = a.id ORDER BY id LIMIT 1),
               (SELECT id FROM sales_deals WHERE advertiser_id = b.id ORDER BY id LIMIT 1)
          FROM sales_advertisers a, sales_advertisers b
         WHERE a.id < b.id
           AND EXISTS (SELECT 1 FROM sales_deals WHERE advertiser_id = a.id)
           AND EXISTS (SELECT 1 FROM sales_deals WHERE advertiser_id = b.id)
         LIMIT 1
    """)).first()


def test_one_annex_is_one_advertiser(env):
    """Правило владельца 06.09.2026. Формулировка в документе говорит про материалы ОДНОГО
    бренда: собрав сделки разных рекламодателей, мы выпустили бы бумагу, где под одной
    услугой перечислены два чужих друг другу бренда, и заметил бы это юрист клиента.

    Проверено на живом примере 06.09.2026: черновик на сделки S35ZA8 (Promomed) и ZCBPLS
    (Dr. Reddy's) собирался молча, а в тексте стояло «материалов бренда: «Dr. Reddy's /
    Хелинорм, Promomed / Салвисар»».
    """
    pair = _two_deals_of_different_advertisers(env.db)
    if not pair or not all(pair):
        pytest.skip('нужны две сделки разных рекламодателей')
    with pytest.raises(HTTPException) as e:
        A.create_annex(A.AnnexIn(contract_id=env.c.id, deal_ids=list(pair),
                                 period_from=date(2026, 6, 1), period_to=date(2026, 6, 30),
                                 total_amount=100), env.db, env.user)
    assert e.value.status_code == 400
    # В отказе стоят и рекламодатели, и метки сделок: иначе непонятно, что разделять.
    detail = str(e.value.detail)
    assert "рекламодател" in detail.lower()
    for did in pair:
        code = env.db.execute(
            __import__("sqlalchemy").text("SELECT code FROM sales_deals WHERE id = :i"),
            {"i": did}).scalar()
        assert (code or str(did)) in detail


def test_deals_of_one_advertiser_still_assemble(env):
    """Проверка не должна запрещать то, ради чего она заведена: несколько сделок ОДНОГО
    рекламодателя в одном приложении — обычный случай."""
    from sqlalchemy import text as sa_text
    from app.sales.models import SalesAnnex, SalesDealAnnexAllocation as Alloc
    ids = [r[0] for r in env.db.execute(sa_text("""
        SELECT id FROM sales_deals
         WHERE advertiser_id = (SELECT advertiser_id FROM sales_deals
                                 WHERE advertiser_id IS NOT NULL
                                 GROUP BY advertiser_id HAVING count(*) > 1
                                 ORDER BY advertiser_id LIMIT 1)
         ORDER BY id LIMIT 2
    """)).all()]
    if len(ids) < 2:
        pytest.skip('нужны две сделки одного рекламодателя')
    out = A.create_annex(A.AnnexIn(contract_id=env.c.id, deal_ids=ids,
                                   period_from=date(2026, 6, 1), period_to=date(2026, 6, 30),
                                   total_amount=1000), env.db, env.user)
    rows = env.db.query(Alloc).filter(Alloc.annex_id == out["id"]).all()
    assert sorted(r.deal_id for r in rows) == sorted(ids)
    # Разнесённое сходится с приложением до копейки.
    assert round(sum(r.amount for r in rows), 2) == 1000.0
    env.db.query(Alloc).filter(Alloc.annex_id == out["id"]).delete()
    env.db.query(SalesAnnex).filter(SalesAnnex.id == out["id"]).delete()
    env.db.commit()


def test_empty_selection_exports_nothing_not_everything(env):
    """Ревью 06.09.2026: `ids or list(items)` превращал пустой отбор в «выгрузить всё».
    Человек ищет строку, под которую ничего не подошло, жмёт выгрузку — и получает весь
    реестр вместо пустого файла. Пустой СПИСОК и ОТСУТСТВИЕ списка — разные состояния."""
    import asyncio
    from io import BytesIO
    from openpyxl import load_workbook
    _draft(env)

    def rows_of(payload):
        r = A.export_annexes(payload, env.db, env.user)

        async def grab():
            return b"".join([c async for c in r.body_iterator])
        return load_workbook(BytesIO(asyncio.run(grab()))).active.max_row

    assert rows_of(A.ExportIn(ids=[])) == 1, 'в файле должна остаться только шапка'
    assert rows_of(A.ExportIn()) > 1, 'без отбора выгружается весь реестр'


def test_registry_says_how_many_there_are_in_total(env):
    """Молча обрезанный список выглядит как «документ пропал», а не как «показано не всё»:
    искать будут документ, а не страницу. Поэтому потолок назван числом."""
    _draft(env)
    out = A.list_annexes(env.db, env.user)
    assert out["limit"] == A.LIST_LIMIT
    assert out["total"] >= len(out["items"])
    assert out["truncated"] is (out["total"] > len(out["items"]))
