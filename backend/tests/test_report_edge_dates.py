# -*- coding: utf-8 -*-
"""Две ошибки сервера на краевых датах (аудит 23.09.2026, этап 8.4: 2.L2, 2.L3).

2.L2. Выгрузка финотчёта за диапазон без операций падала 500: формула итога брала
первую месячную колонку, а их не было. Теперь — файл с шапкой и пометкой.

2.L3. Дашборд руководителя сравнивает с тем же днём прошлого года через
`today.replace(year=…-1)` — 29 февраля такого дня нет, и экран падал раз в четыре года.
"""
import asyncio
import io
from datetime import date

import pytest
from openpyxl import load_workbook

import app.main  # noqa: F401 — все модели
from app.database import SessionLocal
from app.models import User
from app.routers import exec_dashboard, finreport


@pytest.fixture
def ctx():
    db = SessionLocal()
    user = db.query(User).filter(User.email == "d.makarov@simb-ad.com").first()
    yield db, user
    db.rollback()
    db.close()


def test_an_empty_range_exports_a_file_not_a_500(ctx):
    db, user = ctx
    resp = finreport.export_finreport(basis="accrual", vat="net", granularity="month",
                                      date_from="2099-01", date_to="2099-02",
                                      db=db, current_user=user)

    async def read():
        return b"".join([c async for c in resp.body_iterator])
    body = asyncio.run(read())
    ws = load_workbook(io.BytesIO(body)).active
    text = " ".join(str(c.value) for r in ws.iter_rows() for c in r if c.value)
    assert "операций нет" in text


def test_the_29th_of_february_does_not_break_the_dashboard(ctx):
    db, _ = ctx
    out = exec_dashboard._booking(db, date(2028, 2, 29))
    assert out["months"], "блок брони пуст"
