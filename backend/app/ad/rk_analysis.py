# -*- coding: utf-8 -*-
"""Страница разбора одной РК для трафика (владелец 11.10.2026, макет «Разбор РК»).

Сервер отдаёт данные в МЕЛКОМ зерне одним ответом, а суммы, проценты, прогноз и выводы считает
страница (`frontend/lib/rkAnalysis.mjs`): так выключение площадки пересчитывает всё мгновенно.
Здесь нет своих правил — только сбор из того, что уже считается в системе:

* факт площадка × день × источник — `ad_campaign_stat` (DSP, Adfox, Weborama; `offered` с 06.10);
* креатив × день с расходом — сырьё DSP `dsp_stat_raw` через `stat_store.hash_map`;
* уники — модель SIMB ID, та же, что в «Отчёте клиенту с поправкой» (`client_report.build(model=True)`):
  коэффициенты пары «площадка × креатив» за день фиксируются разово, поэтому страница и Excel клиенту
  показывают одно число;
* события — журнал действий (запуск/пауза площадок, статус РК, ссылки нацеливания, «Обновить в DSP»,
  «Темп размещения») и прогоны биддера; у событий площадки — её `placement_id`;
* порог расхождения с Weborama — `stat_sources.goal_limit` из целей сделки;
* проверки — записанный ночной прогон `rk_check`;
* прошлый флайт — предыдущие РК той же связки «рекламодатель + бренд» (по умолчанию последняя).
"""
from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from app.ad.boost import boost_window  # noqa: F401 — единое определение окна буста, тест читает его здесь
from typing import Dict, List, Optional

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.ad.stat_sources import COMBAT, VERIFIER, goal_limit

TECH = "tech"
# действие журнала → вид события на оси
_CAMPAIGN_ACTIONS = {"ad_campaign_status": "campaign", "targeting_link": TECH, "dsp_refresh": "campaign"}
_PLACEMENT_KIND = re.compile(r"→\s*(\S+)")


def _iso(v) -> Optional[str]:
    return v.isoformat() if isinstance(v, (date, datetime)) else v


def head_of(db: Session, code: str) -> Optional[dict]:
    r = db.execute(text("""
        SELECT c.id, c.date_start, c.date_end, c.plan_show, c.plan_budget, c.status, c.ms_campaign_xxhash,
               d.id AS deal_id, d.code, d.title, d.advertiser_id, d.brand_id,
               coalesce(a.short_name, a.name) AS adv, b.name AS brand, st.name AS stage
          FROM ad_campaign c JOIN sales_deals d ON d.id = c.deal_id
          LEFT JOIN sales_advertisers a ON a.id = d.advertiser_id
          LEFT JOIN sales_brands b ON b.id = d.brand_id
          LEFT JOIN sales_stages st ON st.id = d.our_stage_id
         WHERE upper(d.code) = upper(:c)"""), {"c": code}).mappings().first()
    return dict(r) if r else None


def placements_of(db: Session, campaign_id: int) -> List[dict]:
    rows = db.execute(text("""
        SELECT p.id, p.publisher_id, coalesce(pb.domain, pb.name) AS domain, pb.name, p.status, p.share, pb.cpm_contract AS cpm,
               p.plan_show, p.first_started_at, p.weborama_pixel IS NOT NULL AS pixel,
               (SELECT count(*) FROM ad_campaign_creative cr WHERE cr.placement_id = p.id
                  AND coalesce(cr.ms_creative_xxhash, '') <> '') AS cr_dsp,
               (SELECT count(*) FROM ad_campaign_creative cr WHERE cr.placement_id = p.id) AS cr_all
          FROM ad_campaign_placement p JOIN sales_publishers pb ON pb.id = p.publisher_id
         WHERE p.campaign_id = :c ORDER BY p.plan_show DESC NULLS LAST, p.id"""), {"c": campaign_id}).mappings().all()
    base = _plans_without_boost(db, campaign_id)
    return [{**dict(r), "share": float(r["share"] or 0),
             "plan_show": float(base.get(r["id"], r["plan_show"]) or 0),
             "cpm": float(r["cpm"]) if r["cpm"] is not None else None,
             "first_started_at": _iso(r["first_started_at"]), "external": not r["cr_dsp"]} for r in rows]


def _plans_without_boost(db: Session, campaign_id: int) -> Dict[int, float]:
    """Планы площадок без «темпа размещения», пока он действует; без буста — пусто (берётся записанный план).

    Записанный план при бусте поднят — так он уходит в DSP. Страница же добавляет буст сама, множителем на
    его дни (`boostK` в `lib/rkAnalysis.mjs`): с поднятым планом буст считался дважды, а «План» шапки
    расходился с планом РК в карточке сделки (9E8Z4U, 26.10.2026: 2 184 257 против 1 826 485)."""
    from app.ad import boost
    from app.ad import build as ad_build
    if not boost.is_live(boost.active(db, campaign_id)):
        return {}
    _pls, out = ad_build.campaign_layout(db, campaign_id, with_boost=False)
    return {r["id"]: r["plan_show"] for r in out["rows"]}


def stat_of(db: Session, campaign_id: int, until: date) -> List[dict]:
    """Факт площадка × день × источник (наш боевой факт и верификатор)."""
    rows = db.execute(text("""
        SELECT placement_id, date, source, shows, clicks, uniques, offered FROM ad_campaign_stat
         WHERE campaign_id = :c AND placement_id IS NOT NULL AND date <= :u AND source = ANY(:s)
         ORDER BY date, placement_id"""),
        {"c": campaign_id, "u": until, "s": list(COMBAT) + list(VERIFIER)}).mappings().all()
    out = []
    for r in rows:
        src = "weborama" if r["source"] in VERIFIER else ("dsp" if r["source"] in ("dsp", "ms") else "adfox")
        out.append({"placement_id": r["placement_id"], "date": _iso(r["date"]), "source": src,
                    "shows": int(r["shows"] or 0), "clicks": int(r["clicks"] or 0),
                    "uniques": r["uniques"], "offered": r["offered"]})
    return out


def creative_days_of(db: Session, dsp_db, head: dict, until: date) -> List[dict]:
    """Креатив × день из сырья DSP: показы, клики, расход, предложено. Нет сырья — пусто."""
    if dsp_db is None or not head.get("ms_campaign_xxhash"):
        return []
    from app.dsp import stat_store
    from app.dsp.client import PROD, MsClient
    cl = MsClient(contour=PROD)
    hmap = stat_store.hash_map(db, dsp_db, [head["id"]], contour=cl.contour, partner=cl.partner_xxhash)
    labels = {r[0]: (r[1], r[2], r[3], r[4]) for r in db.execute(text("""
        SELECT cr.id, cr.placement_id, coalesce(cs.no, cr.creative_no), cs.title, cs.report_group
          FROM ad_campaign_creative cr LEFT JOIN launch_prep_creative_set cs ON cs.id = cr.root_set_id
         WHERE cr.campaign_id = :c"""), {"c": head["id"]}).all()}
    out = []
    for day, cr, shows, clicks, spend, offered in dsp_db.execute(text("""
        SELECT (ts AT TIME ZONE 'Europe/Moscow')::date, ms_creative_xxhash, shows, clicks, spend, offered
          FROM dsp_stat_raw
         WHERE ms_campaign_xxhash = :h AND ms_source_key IS NULL AND ms_creative_xxhash IS NOT NULL
           AND (ts AT TIME ZONE 'Europe/Moscow')::date <= :u"""),
            {"h": head["ms_campaign_xxhash"].upper(), "u": until}).all():
        rid = hmap.rows.get((cr or "").upper())
        if rid is None or rid not in labels:
            continue
        pid, no, title, grp = labels[rid]
        out.append({"date": _iso(day), "placement_id": pid, "set_no": no, "set_title": title or "", "group": grp,
                    "shows": int(shows or 0), "clicks": int(clicks or 0),
                    "spend": float(spend) if spend is not None else 0.0, "offered": offered})
    return out


def model_report(db: Session, dsp_db, head: dict, until: date):
    """Отчёт клиенту «с поправкой» за флайт по дату факта — источник уников страницы."""
    from app.ad import client_report
    start = head["date_start"]
    if not start or until < start:
        return None
    return client_report.build(db, head["id"], date_from=start, date_to=until, dsp_db=dsp_db, model=True)


def model_uniques_of(db: Session, dsp_db, head: dict, placements: List[dict], until: date, rep=None) -> List[dict]:
    """Уники модели SIMB ID по площадке за день — та же функция, что «Отчёт клиенту с поправкой»."""
    rep = rep or model_report(db, dsp_db, head, until)
    if rep is None:
        return []
    pid_of = {p["domain"]: p["id"] for p in placements}
    out = []
    for site, days in rep.site_days.items():
        pid = pid_of.get(site)
        if pid is None:
            continue
        for d, v in days.items():
            if len(v) > 2 and v[2]:
                out.append({"placement_id": pid, "date": _iso(d), "uniques": round(v[2])})
    return out


def creative_uniques_of(db: Session, head: dict, rep) -> List[dict]:
    """Уники модели по креативу × площадке × дню (владелец 24.10.2026: разрез «по креативам»)."""
    if rep is None or not rep.pair_days:
        return []
    of = {r[0]: (r[1], r[2], r[3]) for r in db.execute(text("""
        SELECT cr.id, cr.placement_id, coalesce(cs.no, cr.creative_no), cs.report_group
          FROM ad_campaign_creative cr LEFT JOIN launch_prep_creative_set cs ON cs.id = cr.root_set_id
         WHERE cr.campaign_id = :c"""), {"c": head["id"]}).all()}
    return [{"date": _iso(day), "placement_id": of[cid][0], "set_no": of[cid][1], "group": of[cid][2],
             "uniques": round(v[2])}
            for (cid, day), v in rep.pair_days.items() if cid in of and len(v) > 2 and v[2]]


def events_of(db: Session, head: dict, placements: List[dict]) -> List[dict]:
    """События на ось времени: журнал действий по сделке и площадкам этой РК, «Темп размещения», биддер."""
    code, deal_id = head["code"], head["deal_id"]
    pid_of_pub = {p["publisher_id"]: p["id"] for p in placements}
    out = []
    for at, action, details in db.execute(text("""
        SELECT created_at, action, details FROM audit_log
         WHERE entity_type = 'sales_deal' AND entity_id = :d AND action = ANY(:a) ORDER BY created_at"""),
            {"d": deal_id, "a": list(_CAMPAIGN_ACTIONS)}).all():
        out.append({"at": _iso(at), "kind": _CAMPAIGN_ACTIONS[action], "placement_id": None, "text": details or action})
    for at, pub, details in db.execute(text("""
        SELECT created_at, entity_id, details FROM audit_log
         WHERE action = 'ad_placement_status' AND details LIKE :p ORDER BY created_at"""),
            {"p": f"РК {code}:%"}).all():
        m = _PLACEMENT_KIND.search(details or "")
        to = m.group(1) if m else ""
        kind = "start" if to.startswith("запущен") else "pause"
        out.append({"at": _iso(at), "kind": kind, "placement_id": pid_of_pub.get(pub), "text": details})
    for r in db.execute(text("""SELECT pct, days, starts_on, until, created_at, ended_at, ended_reason
                                  FROM ad_campaign_boost WHERE campaign_id = :c ORDER BY id"""), {"c": head["id"]}).mappings().all():
        out.append({"at": _iso(r["created_at"]), "kind": "boost", "placement_id": None,
                    "text": f"Темп размещения +{r['pct']} % на {r['days']} дн. ({r['starts_on']:%d.%m}–{r['until']:%d.%m})"})
    for at, n in db.execute(text("""
        SELECT b.started_at, count(ch.id) FROM bidder_run b JOIN bidder_run_change ch ON ch.run_id = b.id
         WHERE ch.campaign_id = :c GROUP BY b.id, b.started_at ORDER BY 1"""), {"c": head["id"]}).all():
        out.append({"at": _iso(at), "kind": "bidder", "placement_id": None,
                    "text": f"Ночной пересчёт долей: изменено планов площадок {n}"})
    out.sort(key=lambda e: e["at"] or "")
    return out


def checks_of(db: Session, campaign_id: int) -> Optional[dict]:
    """Сводка последнего ночного прогона проверок (та же, что строка «Сводка» на дашборде)."""
    from app.ad import rk_check as R
    day = db.execute(text("SELECT max(run_date) FROM rk_check_result")).scalar()
    if day is None:
        return None
    rows = db.execute(text("""SELECT campaign_id, check_key, level, message FROM rk_check_result
                               WHERE run_date = :d AND placement_id IS NULL"""), {"d": day}).all()
    by: Dict[int, list] = {}
    for cid, key, lv, msg in rows:
        by.setdefault(cid, []).append({"key": key, "level": lv, "message": msg})
    if campaign_id not in by:
        return None
    common = R.common_keys({cid: {x["key"]: x["level"] for x in v} for cid, v in by.items()})
    return {"run_date": str(day), **R.build_summary(by[campaign_id], common)}


def prev_options(db: Session, head: dict) -> List[dict]:
    """Прошлые РК той же связки «рекламодатель + бренд», с фактом, раньше этой — от новой к старой."""
    if not head.get("advertiser_id"):
        return []
    rows = db.execute(text("""
        SELECT c.id, d.code, d.title, c.date_start, c.date_end
          FROM ad_campaign c JOIN sales_deals d ON d.id = c.deal_id
         WHERE d.advertiser_id = :a AND d.brand_id IS NOT DISTINCT FROM :b AND c.id <> :c
           AND c.date_start < :s
           AND EXISTS (SELECT 1 FROM ad_campaign_stat s WHERE s.campaign_id = c.id AND s.source = ANY(:own) AND s.shows > 0)
         ORDER BY c.date_start DESC, c.id DESC LIMIT 12"""),
        {"a": head["advertiser_id"], "b": head["brand_id"], "c": head["id"], "s": head["date_start"],
         "own": list(COMBAT)}).mappings().all()
    return [{"campaign_id": r["id"], "code": r["code"], "title": r["title"],
             "date_start": _iso(r["date_start"]), "date_end": _iso(r["date_end"])} for r in rows]


def prev_flight(db: Session, campaign_id: int) -> Optional[dict]:
    """Прошлый флайт для сравнения ПО НОМЕРУ ДНЯ флайта (не по дате): факт и WR по домену и дню."""
    h = db.execute(text("SELECT date_start, date_end, plan_show FROM ad_campaign WHERE id = :c"),
                   {"c": campaign_id}).mappings().first()
    if not h or not h["date_start"]:
        return None
    rows = db.execute(text("""
        SELECT coalesce(pb.domain, pb.name) AS domain, s.date, s.source, sum(s.shows) AS shows, sum(s.clicks) AS clicks
          FROM ad_campaign_stat s JOIN ad_campaign_placement p ON p.id = s.placement_id
          JOIN sales_publishers pb ON pb.id = p.publisher_id
         WHERE s.campaign_id = :c AND s.source = ANY(:s) GROUP BY 1, 2, 3"""),
        {"c": campaign_id, "s": list(COMBAT) + list(VERIFIER)}).mappings().all()
    length = (h["date_end"] - h["date_start"]).days + 1 if h["date_end"] else None
    return {"campaign_id": campaign_id, "plan_show": float(h["plan_show"] or 0), "length": length,
            "stat": [{"domain": r["domain"], "day_no": (r["date"] - h["date_start"]).days + 1,
                      "source": "weborama" if r["source"] in VERIFIER else "own",
                      "shows": int(r["shows"] or 0), "clicks": int(r["clicks"] or 0)} for r in rows]}


def build(db: Session, head: dict, dsp_db=None, prev_id: Optional[int] = None,
          today: Optional[date] = None) -> dict:
    from app.ad import build as ad_build
    from app.ad import stat_sources as S
    today = today or date.today()
    as_of = S.fact_as_of(db, today)
    placements = placements_of(db, head["id"])
    options = prev_options(db, head)
    chosen = prev_id if prev_id in {o["campaign_id"] for o in options} else (options[0]["campaign_id"] if options else None)
    goals = ad_build.deal_goals(db, head["deal_id"]) or {}
    model = model_report(db, dsp_db, head, as_of)
    return {
        "head": {**{k: _iso(v) for k, v in head.items() if k not in ("advertiser_id", "brand_id")},
                 "plan_show": float(head["plan_show"] or 0), "plan_budget": float(head["plan_budget"] or 0),
                 "goals": goals, "wr_limit": goal_limit(goals), "as_of": _iso(as_of),
                 "offer_from": _iso(_first_offer_day(db, head["id"]))},
        "placements": placements,
        "stat": stat_of(db, head["id"], as_of),
        "model_uniques": model_uniques_of(db, dsp_db, head, placements, as_of, rep=model),
        "creative_uniques": creative_uniques_of(db, head, model),
        "creative_days": creative_days_of(db, dsp_db, head, as_of),
        "events": events_of(db, head, placements),
        # Окна «темпа размещения»: в эти дни план дня поднят, строки таблицы подкрашены (владелец 13.10.2026).
        "boosts": boosts_of(db, head["id"]),
        "checks": checks_of(db, head["id"]),
        "prev": {"options": options, "selected": chosen,
                 "flight": prev_flight(db, chosen) if chosen else None},
    }


def boosts_of(db: Session, campaign_id: int) -> List[dict]:
    from app.ad.boost import windows_of
    return windows_of(db, [campaign_id]).get(campaign_id, [])


def _first_offer_day(db: Session, campaign_id: int) -> Optional[date]:
    return db.execute(text("SELECT min(date) FROM ad_campaign_stat WHERE campaign_id = :c AND source = 'dsp' AND offered IS NOT NULL"),
                      {"c": campaign_id}).scalar()


# ── Excel текущего среза ──────────────────────────────────────────────────────
MAX_ROWS, MAX_COLS = 10000, 40


def export_xlsx(title: str, columns: List[dict], rows: List[dict]) -> bytes:
    """Срез страницы как есть: заголовок, шапка колонок, строки, итог (если передан строкой)."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from app.xlsx_safe import save_workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Срез"
    ws.append([title[:250]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append([c["label"] for c in columns])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(c["key"]) for c in columns])
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = max(10, min(40, len(c["label"]) + 4))
    buf = io.BytesIO()
    save_workbook(wb, buf)
    return buf.getvalue()


def flight_days(head: dict) -> List[str]:
    s, e = head.get("date_start"), head.get("date_end")
    if not s or not e:
        return []
    s, e = date.fromisoformat(str(s)[:10]), date.fromisoformat(str(e)[:10])
    return [(s + timedelta(days=i)).isoformat() for i in range((e - s).days + 1)]
