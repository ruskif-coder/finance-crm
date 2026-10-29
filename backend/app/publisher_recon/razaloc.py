# -*- coding: utf-8 -*-
"""Разаллокация площадке (владелец 29.10.2026): после согласованной сверки мы сами раскладываем итог площадки за месяц
по изначальным договорам ОРД и отдаём ей файл — она по нему сдаёт свой акт в ОРД.

Площадка ничего не присылает: показы — согласованный итог сверки, уже разнесённый по размещениям (`publisher_recon_line`,
доля нашего факта); договор — привязка сделки в ОРД (изначальный, у прямого рекламодателя — доходный); CPM — площадки,
зафиксированный при согласовании. Цены клиента здесь нет и быть не должно.
Форма — по образцу владельца (разаллокация площадки за 08.2026).
"""
from __future__ import annotations

from decimal import ROUND_HALF_UP, Decimal
import re
from io import BytesIO
from typing import Dict, List, Optional

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.ord.enums import ACTION_TYPES, CONTRACT_TYPES, SUBJECT_TYPES, label_by_code
from app.xlsx_safe import save_workbook

COLUMNS: List[str] = [
    "Сайт", "ID ИД медиаскаут", "ИНН РД", "Наименование РД", "ИНН РА1", "Наименование РА1", "Номер договора",
    "Дата договора", "Тип договора", "Вид договора/Предмет договора", "Вид деятельности", "РА1 действует в интересах",
    "Налоговая форма", "Показы", "CPM", "НДС,%", "Расчетные показы", "Бюджет до НДС", "Бюджет с НДС",
]

# Доходный договор с прямым рекламодателем — договор на распространение рекламы (так в образце владельца);
# в зеркале доходных предмета нет.
DIRECT_SUBJECT = "Distribution"


def _money(v) -> float:
    return float(Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _interests(c: dict) -> str:
    if c.get("type") != "MediationContract":
        return ""
    return "В интересах РР" if c.get("is_agent_acting_for_publisher") else "В интересах РД"


def compose(lines: List[dict], site: str, cpm, vat_rate) -> dict:
    """Строки разаллокации: одна на договор ОРД, показы — сумма согласованных по его сделкам. Сделка без договора в ОРД
    не теряется молча — уходит в `missing`, чтобы её привязали до отправки площадке."""
    vat = Decimal(str(vat_rate or 0))
    groups: Dict[str, dict] = {}
    missing = []
    for ln in lines:
        volume = int(ln["agreed_volume"] or 0)
        if volume <= 0:
            continue
        c = ln.get("contract")
        if not c or not c.get("ord_id"):
            missing.append({"deal_code": ln["deal_code"], "volume": volume})
            continue
        g = groups.setdefault(c["ord_id"], {"contract": c, "volume": 0})
        g["volume"] += volume
    rows = []
    for key in sorted(groups, key=lambda k: -groups[k]["volume"]):
        c, volume = groups[key]["contract"], groups[key]["volume"]
        net = Decimal(volume) * Decimal(str(cpm or 0)) / 1000
        subject = c.get("subject_type") or (DIRECT_SUBJECT if c.get("kind") == "final" else None)
        rows.append({
            "Сайт": site, "ID ИД медиаскаут": c["ord_id"],
            "ИНН РД": c.get("advertiser_inn"), "Наименование РД": c.get("advertiser_name"),
            "ИНН РА1": c.get("contractor_inn"), "Наименование РА1": c.get("contractor_name"),
            "Номер договора": c.get("number"), "Дата договора": c.get("date"),
            "Тип договора": label_by_code(CONTRACT_TYPES, c.get("type")),
            "Вид договора/Предмет договора": label_by_code(SUBJECT_TYPES, subject),
            "Вид деятельности": label_by_code(ACTION_TYPES, c.get("action_type")),
            "РА1 действует в интересах": _interests(c),
            "Налоговая форма": "С НДС" if vat > 0 else "Без НДС",
            "Показы": volume, "CPM": float(cpm or 0), "НДС,%": float(vat / 100),
            "Расчетные показы": volume,
            "Бюджет до НДС": _money(net), "Бюджет с НДС": _money(net * (1 + vat / 100)),
        })
    return {"rows": rows, "missing": missing}


# ── данные из базы ──────────────────────────────────────────────────────────────────────────────────

def lines_of(db: Session, request_id: int) -> List[dict]:
    """Строки согласованной сверки с договором ОРД сделки: изначальный, а у сделки без него — доходный (прямой
    рекламодатель; исполнитель — наше юрлицо договора)."""
    from app import own_company
    from app.models import Contract
    rows = db.execute(text("""
        SELECT d.code AS deal_code, l.agreed_volume,
               ic.ord_id AS i_ord_id, ic.advertiser_inn AS i_adv_inn, ic.advertiser_name AS i_adv_name,
               ic.contractor_inn AS i_con_inn, ic.contractor_name AS i_con_name, ic.number AS i_number, ic.date AS i_date,
               ic.type AS i_type, ic.subject_type AS i_subject, ic.action_type AS i_action,
               ic.is_agent_acting_for_publisher AS i_for_rr,
               k.id AS contract_id, k.ord_contract_id AS f_ord_id, coalesce(m.number, k.contract_number) AS f_number,
               coalesce(m.date, k.contract_date) AS f_date, m.type AS f_type,
               coalesce(m.client_inn, cp.inn, k.inn) AS f_adv_inn, coalesce(m.client_name, cp.name, k.counterparty_name) AS f_adv_name
          FROM publisher_recon_line l
          JOIN launch_prep_target t ON t.id = l.target_id
          JOIN sales_deals d ON d.id = t.deal_id
          LEFT JOIN ord_initial_contracts ic ON ic.id = d.ord_initial_contract_id
          LEFT JOIN contracts k ON k.id = d.ord_final_contract_id
          LEFT JOIN counterparties cp ON cp.id = k.counterparty_id
          -- Зеркало уникально по (ord_id, контур): без контура один id из demo и prod удвоил бы строку и показы.
          LEFT JOIN ord_final_mirror m ON m.ord_id = k.ord_contract_id AND m.ord_env = coalesce(k.ord_env, m.ord_env)
         WHERE l.request_id = :r
         ORDER BY d.code"""), {"r": request_id}).mappings().all()
    out = []
    for r in rows:
        contract = None
        if r["i_ord_id"]:
            contract = {"kind": "initial", "ord_id": r["i_ord_id"], "advertiser_inn": r["i_adv_inn"],
                        "advertiser_name": r["i_adv_name"], "contractor_inn": r["i_con_inn"],
                        "contractor_name": r["i_con_name"], "number": r["i_number"], "date": r["i_date"],
                        "type": r["i_type"], "subject_type": r["i_subject"], "action_type": r["i_action"],
                        "is_agent_acting_for_publisher": r["i_for_rr"]}
        elif r["f_ord_id"]:
            us = own_company.of_contract(db.get(Contract, r["contract_id"])) or own_company.sole(db)
            contract = {"kind": "final", "ord_id": r["f_ord_id"], "advertiser_inn": r["f_adv_inn"],
                        "advertiser_name": r["f_adv_name"], "contractor_inn": us.inn if us else None,
                        "contractor_name": us.name if us else None, "number": r["f_number"], "date": r["f_date"],
                        "type": r["f_type"] or "ServiceAgreement", "subject_type": None, "action_type": None,
                        "is_agent_acting_for_publisher": False}
        out.append({"deal_code": r["deal_code"], "agreed_volume": r["agreed_volume"], "contract": contract})
    return out


def publisher_vat(db: Session, publisher_id: int) -> Optional[float]:
    """Ставка НДС площадки — расходная ставка юрлица её действующего договора «с площадкой». Нет договора — None."""
    return db.execute(text("""
        SELECT cp.vat_rate_expense FROM sales_publisher_contracts pc
          JOIN contracts k ON k.id = pc.contract_id JOIN counterparties cp ON cp.id = k.counterparty_id
         WHERE pc.publisher_id = :p AND pc.role = 'с площадкой' AND NOT pc.is_archived
         ORDER BY pc.id DESC LIMIT 1"""), {"p": publisher_id}).scalar()


def of_request(db: Session, req) -> dict:
    """Разаллокация согласованной сверки; `vat_known` — нашлась ли ставка площадки (иначе считаем без НДС и предупреждаем)."""
    site = db.execute(text("SELECT coalesce(domain, name) FROM sales_publishers WHERE id = :p"),
                      {"p": req.publisher_id}).scalar()
    # CPM и НДС — из строки итога, зафиксированные при согласовании, а не из `req.recon`: у сверки, согласованной в
    # этом же сеансе, связь ещё не перечитана. НДС пуст только у согласованных до 29.10.2026 — тогда текущая ставка.
    cpm, vat = db.execute(text("SELECT cpm, vat_rate FROM publisher_recon WHERE request_id = :r"), {"r": req.id}).first()
    if vat is None:
        vat = publisher_vat(db, req.publisher_id)
    out = compose(lines_of(db, req.id), site, cpm, None if vat is None else float(vat))
    return {**out, "site": site, "period": req.period, "vat_known": vat is not None}


def response(db: Session, reqs):
    """Файл согласованных сверок — общий для нашей стороны и шлюза кабинета; площадка на своей вкладке (владелец 29.10.2026).
    До согласования итога нет — 409, и тогда файла нет целиком: половина площадок месяца выглядела бы как вся разаллокация."""
    from urllib.parse import quote

    from fastapi import HTTPException
    from fastapi.responses import Response

    from app.publisher_recon.board import AGREED
    if any(r.verdict != AGREED for r in reqs):
        raise HTTPException(409, "Разаллокация появится после согласования сверки")
    data = [of_request(db, r) for r in reqs]
    period = "_".join(sorted({d["period"] for d in data}))
    human = f"{data[0]['site'] + ' ' if len(data) == 1 else ''}Разаллокация {period}.xlsx"
    return Response(xlsx([(d["site"], d["rows"]) for d in data]),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=\"razaloc_{period}.xlsx\"; "
                                                    f"filename*=UTF-8''{quote(human)}"})


# ── файл ────────────────────────────────────────────────────────────────────────────────────────────

def _sheet_title(site: str, used: set) -> str:
    """Имя вкладки по правилам Excel: до 31 знака, без скобок, двоеточия, звёздочки, вопроса и слешей, без повторов."""
    base = re.sub(r"[\[\]:*?/\\]", "_", site or "Разаллокация")[:31]
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(title.lower())
    return title


def xlsx(sheets: List[tuple]) -> bytes:
    """`sheets` — [(площадка, строки)]: площадка на своей вкладке."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used: set = set()
    for site, rows in sheets:
        _fill(wb.create_sheet(_sheet_title(site, used)), rows)
    buf = BytesIO()
    save_workbook(wb, buf)          # имена договоров и юрлиц — из ОРД; текст с «=» уходит текстом, не формулой
    return buf.getvalue()


def _fill(ws, rows: List[dict]) -> None:
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    formats = {"Дата договора": "DD.MM.YYYY", "Показы": "# ##0", "Расчетные показы": "# ##0", "НДС,%": "0%",
               "Бюджет до НДС": "# ##0.00", "Бюджет с НДС": "# ##0.00"}
    for r in rows:
        ws.append([r[c] for c in COLUMNS])
        for i, c in enumerate(COLUMNS, 1):
            if c in formats:
                ws.cell(ws.max_row, i).number_format = formats[c]
    for i, c in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(c) + 4))
    ws.freeze_panes = "A2"
