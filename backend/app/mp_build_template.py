"""Одноразовый генератор СТАРТОВОГО Excel-шаблона медиаплана (лист «МП») с плейсхолдерами.

Шаблон — источник дизайна: аккаунт правит его в Excel (шрифты/цвета/ширины/расположение/
логотип), НЕ трогая токены `{{...}}` и оставляя РОВНО ОДНУ строку-образец размещения
(с токенами `{{r.*}}`) и одну строку-образец доп.услуги (`{{e.*}}`). Рендер (media_plans.py
→ _render_from_template) грузит этот файл, подставляет значения и клонирует строки-образцы
под фактическое число позиций.

Запуск:  docker exec finance_backend python /app/app/mp_build_template.py
Пишет:   backend/app/templates/mp_template.xlsx
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

BLUE = "FF7299E8"
NOTES = [
    '1. В случае размещения по модели CPM (единица закупки — "1000 показов"), гарантированными показателями являются количество показов и стоимость закупки за 1000 показов, все остальные показатели являются прогнозными.',
    '2. Плановые показатели прогнозных значений основываются на данных и опыте Команды. Фактические показатели (результаты) рекламной кампании могут измениться как в меньшую, так и в большую сторону. Команда несет ответственность только за ключевые параметры закупки: количество кликов и стоимость за клик (для модели CPC) и кол-во показов и стоимость за 1000 показов (для модели CPM). Команда не дает 100% гарантии достижения прогнозируемых результатов рекламной кампании, указанных в данном медиаплане.',
    '3. Дедлайн предоставления согласованных креативных материалов под все форматы — не менее чем за 2 рабочих дня до старта рекламной кампании. В случае нарушения сроков предоставления исполнитель не может гарантировать своевременный старт рекламной кампании.',
    '4. Медиаплан актуален в течение 14 календарных дней с даты составления.',
    '5. В случае, если после согласования медиаплана и старта РК будут внесены изменения в ключевые KPI и модель закупки медиа-инвентаря, данный медиаплан подлежит перерасчету и согласованию повторно.',
    '6. KPI по CTR подтверждается при условии закупки по CPM и наличии всех форматов из медиаплана. Иначе прогнозные показатели должны быть обновлены.',
]
BONUS = ('*Бонусом при условии: от 500 000 рублей до НДС на бренд в месяц при первом размещении бренда '
         'или при размещении от 1,5М рублей до НДС в месяц на бренд.')

MONEY, MONEY0, INT, PCT2, ROI = "#,##0.00", "#,##0", "#,##0", "0.00", "0%"


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "МП"
    ws.sheet_view.showGridLines = False

    F10 = Font(name="Calibri", size=10)
    F10B = Font(name="Calibri", size=10, bold=True)
    HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
    FILL = PatternFill("solid", fgColor=BLUE)
    thin = Side(style="thin", color="FFBFBFBF")
    BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
    CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
    RIGHT = Alignment(horizontal="right", vertical="center")

    def put(r, c, v=None, font=F10, fill=None, align=None, border=None, numfmt=None):
        cell = ws.cell(r, c)
        if v is not None:
            cell.value = v
        cell.font = font
        if fill:
            cell.fill = fill
        cell.alignment = align or LEFT
        if border:
            cell.border = border
        if numfmt:
            cell.number_format = numfmt
        return cell

    def hdrcell(r, c, v):
        return put(r, c, v, font=HDR, fill=FILL, align=CTR, border=BORD)

    def span(rng, v=None, font=HDR, fill=FILL, align=CTR, border=BORD, numfmt=None):
        for row in ws[rng]:
            for cell in row:
                cell.font = font
                if fill:
                    cell.fill = fill
                cell.alignment = align
                if border:
                    cell.border = border
                if numfmt:
                    cell.number_format = numfmt
        ws.merge_cells(rng)
        anchor = ws[rng.split(":")[0]]
        if v is not None:
            anchor.value = v
        return anchor

    # ── Шапка: реквизиты + Итого ────────────────────────────────────────────
    span("B2:C4", "SIMB-AD", font=Font(name="Calibri", size=16, bold=True, color="FF2F5496"),
         fill=None, align=Alignment(horizontal="left", vertical="center"), border=None)
    reqs = [("Агентство", "{{agency}}"), ("Рекламодатель", "{{advertiser}}"),
            ("Бренд", "{{brand}}"), ("Название РК", "{{title}}"),
            ("Период размещения", "{{period}}"), ("Дата", "{{date}}")]
    for i, (label, tok) in enumerate(reqs):
        rr = 6 + i
        put(rr, 2, label, font=F10B)
        put(rr, 3, tok)
    span("B12:B13", "Таргетинг", font=F10B, fill=None, align=Alignment(horizontal="left", vertical="center"), border=None)
    span("C12:H12", "{{geo}}", font=F10, fill=None, align=LEFT, border=None)
    span("C13:H13", "ЦА: {{ca}}", font=F10, fill=None, align=LEFT, border=None)
    span("E6:H6", "Итого", font=F10B, fill=None, align=Alignment(horizontal="center", vertical="center"), border=None)
    for i, (label, tok) in enumerate([("Стоимость до НДС", "{{total_net}}"), ("НДС", "{{total_vat}}"), ("Стоимость с НДС", "{{total_gross}}")]):
        rr = 7 + i
        span(f"E{rr}:F{rr}", label, font=F10B, fill=None, align=Alignment(horizontal="left", vertical="center"), border=None)
        span(f"G{rr}:H{rr}", font=F10B, fill=None, align=RIGHT, border=None, numfmt=MONEY)
        ws.cell(rr, 7).value = tok

    # ── Шапка таблицы (15–16) ───────────────────────────────────────────────
    H1 = 15
    base = [("B", "Место размещения"), ("C", "Позиция"), ("D", "Гео"), ("E", "Формат"),
            ("F", "Девайс"), ("G", "Тип ротации (для медийных форматов)"), ("H", "Модель закупки"),
            ("K", "Период"), ("L", "Сезонный коэффициент"), ("M", "Стоимость за единицу закупки"),
            ("N", "Стоимость без скидки"), ("O", "Скидка,%"), ("P", "Скидка,руб"),
            ("Q", "Итоговая стоимость размещения до НДС"), ("R", "НДС 22%"),
            ("S", "Итоговая стоимость размещения с НДС")]
    for col, name in base:
        span(f"{col}{H1}:{col}{H1+1}", name)
    span(f"I{H1}:J{H1+1}", "Объем размещения")
    span(f"T{H1}:AH{H1}", "Прогнозные показатели")
    fc_heads = ["Частота\n(max)", "Охват", "Показы", "CTR, %", "Клики", "CPM", "CPC", "CPU",
                "CR,%", "CR, кол-во чеков", "CPО", "Цена", "Доход", "ROI", "SOV,%"]
    for i, name in enumerate(fc_heads):
        hdrcell(H1 + 1, 20 + i, name)

    # ── Одна строка-образец размещения (токены {{r.*}}) ─────────────────────
    r = H1 + 2
    rowmap = [(2, "{{r.place}}", LEFT, None), (3, "{{r.position}}", LEFT, None), (4, "{{r.geo}}", CTR, None),
              (5, "{{r.format}}", CTR, None), (6, "{{r.device}}", CTR, None), (7, "{{r.rotation}}", CTR, None),
              (8, "{{r.model}}", CTR, None), (9, "{{r.volume}}", RIGHT, INT), (10, "{{r.unit_name}}", CTR, None),
              (11, "{{r.period}}", CTR, None), (12, "{{r.season}}", CTR, INT), (13, "{{r.unit_price}}", RIGHT, MONEY),
              (14, "{{r.net_nodisc}}", RIGHT, MONEY), (15, "{{r.disc_pct}}", CTR, INT), (16, "{{r.disc_rub}}", RIGHT, MONEY),
              (17, "{{r.net}}", RIGHT, MONEY), (18, "{{r.vat}}", RIGHT, MONEY), (19, "{{r.gross}}", RIGHT, MONEY),
              (20, "{{r.freq}}", RIGHT, INT), (21, "{{r.reach}}", RIGHT, INT), (22, "{{r.imp}}", RIGHT, INT),
              (23, "{{r.ctr}}", CTR, PCT2), (24, "{{r.clicks}}", RIGHT, INT), (25, "{{r.cpm}}", RIGHT, MONEY0),
              (26, "{{r.cpc}}", RIGHT, MONEY0), (27, "{{r.cpu}}", RIGHT, MONEY), (28, "{{r.cr}}", CTR, PCT2),
              (29, "{{r.checks}}", RIGHT, INT), (30, "{{r.cpo}}", RIGHT, MONEY0), (31, "{{r.price}}", RIGHT, MONEY0),
              (32, "{{r.revenue}}", RIGHT, MONEY0), (33, "{{r.roi}}", RIGHT, ROI), (34, "{{r.sov}}", CTR, PCT2)]
    for c, tok, al, nf in rowmap:
        put(r, c, tok, align=al, border=BORD, numfmt=nf)
    # ИТОГО по размещениям
    r += 1
    put(r, 2, "ИТОГО:", font=F10B, align=LEFT, border=BORD)
    for c in (3, 4, 5, 6, 7, 8, 10, 11, 12, 13):
        put(r, c, border=BORD)
    for c, tok, nf in [(9, "{{t.volume}}", INT), (14, "{{t.net_nodisc}}", MONEY), (16, "{{t.disc_rub}}", MONEY),
                       (17, "{{t.net}}", MONEY), (18, "{{t.vat}}", MONEY), (19, "{{t.gross}}", MONEY), (32, "{{t.revenue}}", MONEY0)]:
        put(r, c, tok, font=F10B, align=RIGHT, border=BORD, numfmt=nf)
    put(r, 15, border=BORD)

    # ── Доп. услуги ─────────────────────────────────────────────────────────
    r += 2
    put(r, 2, "ДОПОЛНИТЕЛЬНЫЕ УСЛУГИ:", font=F10B, align=LEFT)
    r += 1
    exhdr = {"B": (2, "Место размещения"), "C": (3, "Позиция"), "K": (11, "Период"),
             "L": (12, "Сезонный коэффициент"), "M": (13, "Стоимость за единицу закупки"),
             "N": (14, "Стоимость без скидки"), "O": (15, "Скидка,%"), "P": (16, "Скидка,руб"),
             "Q": (17, "Итоговая стоимость размещения до НДС"), "R": (18, "НДС 22%"),
             "S": (19, "Итоговая стоимость размещения с НДС")}
    for _, (ci, name) in exhdr.items():
        hdrcell(r, ci, name)
    span(f"I{r}:J{r}", "Объем размещения")
    r += 1
    exmap = [(2, "{{e.place}}", LEFT, None), (3, "{{e.name}}", LEFT, None), (9, "{{e.volume}}", RIGHT, INT),
             (10, "{{e.unit_name}}", CTR, None), (11, "{{e.period}}", CTR, None), (12, "{{e.season}}", CTR, INT),
             (13, "{{e.unit_price}}", RIGHT, MONEY), (14, "{{e.net_nodisc}}", RIGHT, MONEY), (15, "{{e.disc_pct}}", CTR, INT),
             (16, "{{e.disc_rub}}", RIGHT, MONEY), (17, "{{e.total}}", RIGHT, MONEY), (18, "{{e.vat}}", RIGHT, MONEY),
             (19, "{{e.gross}}", RIGHT, MONEY)]
    for c, tok, al, nf in exmap:
        put(r, c, tok, align=al, border=BORD, numfmt=nf)
    for c in (4, 5, 6, 7, 8):
        put(r, c, border=BORD)
    r += 1
    put(r, 2, "ИТОГО:", font=F10B, align=LEFT, border=BORD)
    for c in (3, 9, 10, 11, 12, 13):
        put(r, c, border=BORD)
    for c, tok in [(14, "{{te.net_nodisc}}"), (16, "{{te.disc_rub}}"), (17, "{{te.total}}"), (18, "{{te.vat}}"), (19, "{{te.gross}}")]:
        put(r, c, tok, font=F10B, align=RIGHT, border=BORD, numfmt=MONEY)
    put(r, 15, border=BORD)
    r += 1
    span(f"B{r}:S{r}", BONUS, font=Font(name="Calibri", size=9, italic=True), fill=None, align=LEFT, border=None)

    # ── Примечания ──────────────────────────────────────────────────────────
    r += 2
    put(r, 2, "ПРИМЕЧАНИЯ:", font=F10B, align=LEFT)
    r += 1
    for note in NOTES:
        span(f"B{r}:S{r}", note, font=Font(name="Calibri", size=9), fill=None, align=LEFT, border=None)
        ws.row_dimensions[r].height = 30
        r += 1

    widths = {"A": 3, "B": 24.3, "C": 45.8, "D": 8, "E": 14, "F": 12, "G": 15, "H": 13,
              "I": 12, "J": 10, "K": 11, "L": 12.4, "M": 15, "N": 13, "O": 9, "P": 12,
              "Q": 15, "R": 12, "S": 15, "T": 9.1, "U": 11, "V": 11, "W": 8, "X": 10,
              "Y": 9.1, "Z": 9, "AA": 9, "AB": 8, "AC": 10.8, "AD": 9.9, "AE": 9,
              "AF": 12.8, "AG": 9.9, "AH": 8}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[H1].height = 40
    ws.row_dimensions[H1 + 1].height = 28

    out_dir = os.path.join(os.path.dirname(__file__), "templates")
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "mp_template.xlsx")
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    build()
