# -*- coding: utf-8 -*-
"""Импорт суточного отчёта Adfox в факт РК (владелец 02.10.2026).

Площадки вне нашей DSP (Максавит, Здесь аптека, Аптечество web) крутятся в Adfox, и до
этого их факт был пуст — недокрут горел красным. Отчёт Adfox по дням: «День · Название
кампании · Показы · Переходы · Уникальные показы».

КЛЮЧ — НАЗВАНИЕ КАМПАНИИ. Трафик заводит кампанию в Adfox под именем файла креатива из
«↓ ADFOX» плюс свои инициалы (`traffic.files` + `offsite_export.adfox_name`):
`PFPYGX-MXV-cr3-01_AK` = сделка · код площадки · креатив № · файл · трафик. Имя вводится
руками, поэтому разбор терпит хвостовые пробелы и переводы строк, а всё, что не сошлось
однозначно, не угадывается, а показывается человеку с вариантами.

Запись — в `ad_campaign_stat` источником `adfox`, по размещению и дню: два креатива одной
площадки в один день складываются. Повторная загрузка того же дня ПЕРЕЗАПИСЫВАЕТ строку
(ключ `uq_ad_stat`), а не прибавляет. Нет строки в отчёте — «нет данных», а не ноль:
отчёт за утро бывает неполным (владелец 03.10.2026).
"""
import re
from datetime import date, datetime
from io import BytesIO
from typing import Dict, Iterable, List, Optional, Tuple

from sqlalchemy import text
from sqlalchemy.orm import Session

SOURCE = "adfox"

HEADERS = {"день": "day", "название кампании": "name", "показы": "shows",
           "переходы": "clicks", "уникальные показы": "uniques"}
NAME_RE = re.compile(r"^([A-Z0-9]{6})-([A-Z0-9]{2,6})-cr(\d+)(?:-(\d+))?(?:_([A-Za-z]{1,4}))?$")

NEW, UPDATE, SAME = "new", "update", "same"
# Предел строк отчёта: суточный отчёт — десятки строк; 5 МБ сжатого xlsx распаковываются в
# куда больший лист, и читать его до конца незачем (ревью 03.10.2026).
MAX_ROWS = 20000
MATCHED, AMBIGUOUS, UNMATCHED = "matched", "ambiguous", "unmatched"
SKIPPED = "skipped"   # служебная строка Adfox — не ошибка и не наша РК


class ImportError_(ValueError):
    """Файл не читается как отчёт Adfox; текст — человеку."""


def _int(v) -> int:
    if v is None or v == "":
        return 0
    try:
        return int(float(str(v).replace(" ", "").replace(" ", "").replace(",", ".")))
    except ValueError:
        raise ImportError_(f"не число: {v!r}")


def _day(v) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


# Excel хранит управляющие символы в ячейке как `_xHHHH_` (перевод строки — `_x000D_`), и
# openpyxl в режиме чтения оставляет их ТЕКСТОМ. Отчёт Adfox 01–02.10 с прода: 12 строк из
# 48 не сопоставились именно так (03.10.2026).
_EXCEL_ESCAPE = re.compile(r"_x[0-9A-Fa-f]{4}_")
# Служебные строки самого Adfox — не наши РК: показы его заглушки «по умолчанию».
SERVICE_NAMES = {"кампания по умолчанию"}


def clean_name(v) -> str:
    return " ".join(_EXCEL_ESCAPE.sub(" ", str(v or "")).split())


def parse(data: bytes) -> List[dict]:
    """Строки отчёта: [{line, day, name, shows, clicks, uniques}]. «Всего» и пустые — мимо."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — любой сбой разбора = «не тот файл»
        raise ImportError_("Файл не читается как Excel (.xlsx)")
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = next(it, None) or ()
    cols = {}
    for i, h in enumerate(head):
        key = HEADERS.get(clean_name(h).lower())
        if key:
            cols[key] = i
    missing = [k for k in ("day", "name", "shows") if k not in cols]
    if missing:
        raise ImportError_("Нет колонок отчёта Adfox: нужны «День», «Название кампании», «Показы»")
    out = []
    for n, row in enumerate(it, start=2):
        if n - 1 > MAX_ROWS:
            raise ImportError_(f"В отчёте больше {MAX_ROWS} строк — это не суточный отчёт Adfox")

        def cell(k):
            i = cols.get(k)
            return row[i] if i is not None and i < len(row) else None
        name, day = clean_name(cell("name")), _day(cell("day"))
        if not name or day is None:
            continue            # «Всего», пустые строки хвоста листа
        try:
            out.append({"line": n, "day": day, "name": name, "shows": _int(cell("shows")),
                        "clicks": _int(cell("clicks")),
                        "uniques": _int(cell("uniques")) if "uniques" in cols else None})
        except ImportError_ as e:
            raise ImportError_(f"строка {n}: {e}")
    if not out:
        raise ImportError_("В отчёте нет ни одной строки с днём и названием кампании")
    return out


# ── сопоставление ─────────────────────────────────────────────────────────────

_CREATIVES_SQL = """
    SELECT cr.id AS creative_id, cr.creative_no, cr.status AS creative_status,
           pl.id AS placement_id, pl.campaign_id, pl.status AS placement_status,
           d.code AS deal, pb.code AS pub, pb.name AS pub_name, pb.domain
      FROM ad_campaign_creative cr
      JOIN ad_campaign_placement pl ON pl.id = cr.placement_id
      JOIN ad_campaign c ON c.id = pl.campaign_id
      JOIN sales_deals d ON d.id = c.deal_id
      JOIN sales_publishers pb ON pb.id = pl.publisher_id
"""


def _label(c) -> str:
    return f"{c['deal']} · {c['pub_name'] or c['pub']} · креатив №{c['creative_no']}"


def _cand(c, why: str) -> dict:
    return {"creative_id": c["creative_id"], "placement_id": c["placement_id"],
            "campaign_id": c["campaign_id"], "label": _label(c), "why": why,
            "placement_status": c["placement_status"]}


def _creatives_of(db: Session, deals: Iterable[str], allowed: Optional[set]) -> List[dict]:
    deals = sorted({d for d in deals if d})
    if not deals:
        return []
    rows = [dict(r) for r in db.execute(text(_CREATIVES_SQL + " WHERE d.code = ANY(:d)"),
                                        {"d": deals}).mappings()]
    return [r for r in rows if allowed is None or r["campaign_id"] in allowed]


def resolve(db: Session, rows: List[dict], allowed: Optional[set] = None) -> List[dict]:
    """Каждой строке — статус сопоставления и креатив либо варианты.

    `matched`   — ровно один креатив с той же сделкой, площадкой и номером;
    `ambiguous` — сделка нашлась, но креатив не однозначен (номер не тот, у площадки web и
                  app, площадки нет в РК): варианты, самые подходящие сверху;
    `unmatched` — имя не разбирается или сделки нет (или она вне области видимости)."""
    parsed = {}
    for r in rows:
        m = NAME_RE.match(r["name"])
        parsed[r["line"]] = m
    pool = _creatives_of(db, [m.group(1) for m in parsed.values() if m], allowed)
    by_deal: Dict[str, List[dict]] = {}
    for c in pool:
        by_deal.setdefault(c["deal"], []).append(c)

    out = []
    for r in rows:
        m = parsed[r["line"]]
        res = dict(r)
        if r["name"].lower() in SERVICE_NAMES:
            res.update(status=SKIPPED, reason="служебная строка Adfox — не наша РК", candidates=[])
            out.append(res)
            continue
        if not m:
            res.update(status=UNMATCHED, reason="имя не по шаблону «СДЕЛКА-ПЛОЩАДКА-crN-NN_ИНИЦИАЛЫ»",
                       candidates=[])
            out.append(res)
            continue
        deal, pub, no = m.group(1), m.group(2), int(m.group(3))
        cands = by_deal.get(deal, [])
        exact = [c for c in cands if c["pub"] == pub and c["creative_no"] == no]
        if len(exact) == 1:
            c = exact[0]
            res.update(status=MATCHED, reason="", candidates=[], **{
                k: c[k] for k in ("creative_id", "placement_id", "campaign_id")}, label=_label(c))
        elif exact:
            res.update(status=AMBIGUOUS, candidates=[_cand(c, "та же сделка, площадка и креатив — "
                                                          "несколько размещений (web / app)")
                                                    for c in exact],
                       reason="у площадки несколько размещений с этим креативом")
        elif cands:
            same_pub = sorted((c for c in cands if c["pub"] == pub),
                              key=lambda c: abs(c["creative_no"] - no))
            other = [c for c in cands if c["pub"] != pub]
            ranked = ([_cand(c, f"та же сделка и площадка, креатив №{c['creative_no']} вместо №{no}")
                       for c in same_pub]
                      + [_cand(c, "та же сделка, другая площадка") for c in other])
            res.update(status=AMBIGUOUS, candidates=ranked[:8],
                       reason=(f"у площадки {pub} в РК нет креатива №{no}" if same_pub
                               else f"площадки {pub} нет в РК сделки"))
        else:
            res.update(status=UNMATCHED, candidates=[],
                       reason=f"сделка {deal} не найдена или у неё нет РК")
        out.append(res)
    return out


def search(db: Session, q: str, allowed: Optional[set] = None, limit: int = 20) -> List[dict]:
    """Ручной поиск креатива: по коду сделки, коду площадки, имени или домену площадки."""
    q = clean_name(q)
    if len(q) < 2:
        return []
    rows = [dict(r) for r in db.execute(text(_CREATIVES_SQL + """
        WHERE d.code ILIKE :q OR pb.code ILIKE :q OR pb.name ILIKE :q OR pb.domain ILIKE :q
        ORDER BY d.code, pb.name, cr.creative_no LIMIT :n"""),
        {"q": f"%{q}%", "n": limit * 3}).mappings()]
    rows = [r for r in rows if allowed is None or r["campaign_id"] in allowed]
    return [_cand(c, "найдено поиском") for c in rows[:limit]]


# ── сверка с загруженным и запись ────────────────────────────────────────────

def aggregate(rows: Iterable[dict]) -> Dict[Tuple[int, int, date], dict]:
    """(кампания, размещение, день) → сумма. Два креатива площадки за день — одна строка."""
    acc: Dict[Tuple[int, int, date], dict] = {}
    for r in rows:
        k = (int(r["campaign_id"]), int(r["placement_id"]), r["day"])
        a = acc.setdefault(k, {"shows": 0, "clicks": 0, "uniques": None, "lines": [], "n": 0,
                               "creatives": {}})
        a["shows"] += int(r.get("shows") or 0)
        a["clicks"] += int(r.get("clicks") or 0)
        # Разбивка по креативам (05.10.2026) — для отчёта клиенту; пишется рядом с суммой.
        if r.get("creative_id") is not None:
            cr = a["creatives"].setdefault(int(r["creative_id"]),
                                           {"shows": 0, "clicks": 0, "uniques": None, "n": 0})
            cr["shows"] += int(r.get("shows") or 0)
            cr["clicks"] += int(r.get("clicks") or 0)
            cr["n"] += 1
            cr["uniques"] = (None if cr["n"] > 1 or r.get("uniques") is None
                             else int(r["uniques"]))
        a["n"] += 1
        # Уникальные НЕ складываются (ревью 03.10.2026): один человек мог видеть оба
        # креатива площадки. Одна строка — её число; несколько — сумма неизвестна, NULL.
        a["uniques"] = (None if a["n"] > 1 or r.get("uniques") is None
                        else int(r["uniques"]))
        if r.get("line") is not None:
            a["lines"].append(r["line"])
    return acc


def diff(db: Session, agg: Dict[Tuple[int, int, date], dict]) -> Dict[Tuple[int, int, date], dict]:
    """Что станет с каждой строкой факта: новые данные, обновление (было → станет) или без изменений."""
    if not agg:
        return {}
    pls = sorted({k[1] for k in agg})
    have = {(r["placement_id"], r["date"]): r for r in db.execute(text("""
        SELECT placement_id, date, shows, clicks, uniques FROM ad_campaign_stat
         WHERE source = :s AND placement_id = ANY(:p)"""), {"s": SOURCE, "p": pls}).mappings()}
    out = {}
    for k, a in agg.items():
        was = have.get((k[1], k[2]))
        if was is None:
            state = NEW
        elif (was["shows"], was["clicks"], was["uniques"]) == (a["shows"], a["clicks"], a["uniques"]):
            state = SAME
        else:
            state = UPDATE
        out[k] = {"state": state, "was": None if was is None else
                  {"shows": was["shows"], "clicks": was["clicks"], "uniques": was["uniques"]}}
    return out


def write(db: Session, agg: Dict[Tuple[int, int, date], dict]) -> int:
    """Записать суммы. Повтор того же дня перезаписывает строку источника `adfox`."""
    n = 0
    for (camp, pl, day), a in agg.items():
        db.execute(text("""
            INSERT INTO ad_campaign_stat (campaign_id, placement_id, date, shows, clicks, uniques,
                                          source, imported_at)
            VALUES (:c, :p, :d, :s, :k, :u, :src, now())
            ON CONFLICT ON CONSTRAINT uq_ad_stat DO UPDATE
               SET shows = EXCLUDED.shows, clicks = EXCLUDED.clicks,
                   uniques = EXCLUDED.uniques, imported_at = now()"""),
            {"c": camp, "p": pl, "d": day, "s": a["shows"], "k": a["clicks"], "u": a["uniques"],
             "src": SOURCE})
        _write_split(db, pl, day, a.get("creatives") or {})
        n += 1
    return n


def _write_split(db: Session, pl: int, day: date, creatives: dict) -> None:
    """Разбивка дня площадки по креативам ЗАМЕНЯЕТСЯ целиком (как и сумма): креатив, которого
    в новом отчёте нет, не остаётся старыми показами (05.10.2026)."""
    db.execute(text("""
        DELETE FROM adfox_creative_stat WHERE date = :d AND creative_id IN (
            SELECT id FROM ad_campaign_creative WHERE placement_id = :p)"""), {"d": day, "p": pl})
    for cid, c in creatives.items():
        db.execute(text("""
            INSERT INTO adfox_creative_stat (creative_id, date, shows, clicks, uniques, imported_at)
            VALUES (:c, :d, :s, :k, :u, now())"""),
            {"c": cid, "d": day, "s": c["shows"], "k": c["clicks"], "u": c["uniques"]})
