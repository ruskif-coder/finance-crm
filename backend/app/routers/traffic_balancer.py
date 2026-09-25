"""Контур «Траффики» → Админка → вкладка «Балансировщик».

Расчётная ёмкость площадки на поверхность (показов/мес) + ввод замеров трафика с листа.
Разбор формулы и зерна строки — в `app/ad/balance.py`. Замеры пишутся в
`sales_publisher_traffic` — те же данные, что в карточке площадки: правятся и там, и здесь.

Право то же, что у каталога блоков — `traffic_catalog` (общая админка трафика), поэтому
отдельного ключа НЕ заводим: раздача доступа к двум вкладкам одна.
Монтируется тем же префиксом `/api/traffic-catalog`.
"""
import io
from typing import List, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.ad import balance
from app.audit import log_action
from app.database import get_db
from app.models import User
from app.permissions import require_permission
from app.sales.models import PUBLISHER_ARCHIVE_STATUS, SalesPublisher

router = APIRouter()

VIEW = require_permission("traffic_catalog", "view")
EDIT = require_permission("traffic_catalog", "edit")


class BalanceRowIn(BaseModel):
    volume: Optional[float] = None          # объём замера поверхности
    depth: Optional[float] = None           # глубина просмотра
    requests: Optional[float] = None        # запросы рекламного кода
    index_manual: Optional[float] = None    # ручная правка индекса
    is_locked: Optional[bool] = None        # не перетирать пересчётом
    note: Optional[str] = None


class CoefficientsIn(BaseModel):
    k: Optional[float] = None
    depth_default: Optional[float] = None


# ── чтение и правка ───────────────────────────────────────────────────────

@router.get("/balancer")
def balancer_rows(db: Session = Depends(get_db), user: User = Depends(VIEW)):
    return {"rows": balance.rows(db), "coefficients": balance.get_coefficients(db),
            "month": balance.month_start().isoformat()}


@router.put("/balancer/row/{publisher_id}/{scope}")
def balancer_save_row(publisher_id: int, scope: str, payload: BalanceRowIn,
                      db: Session = Depends(get_db), user: User = Depends(EDIT)):
    if scope not in balance.SCOPES:
        raise HTTPException(400, f"Поверхность бывает {balance.SCOPES}")
    if not db.query(SalesPublisher).get(publisher_id):
        raise HTTPException(404, "Площадка не найдена")

    if payload.volume is not None or payload.depth is not None:
        balance.upsert_measurement(db, publisher_id, scope, payload.volume, payload.depth)
    if payload.requests is not None:
        balance.upsert_measurement(db, publisher_id, balance.REQ_SCOPE[scope], payload.requests)

    db.execute(text("""
        INSERT INTO publisher_balance_index
            (publisher_id, scope, index_manual, is_locked, note, updated_at, updated_by)
        VALUES (:p, :s, :im, COALESCE(:lk, FALSE), :n, now(), :u)
        ON CONFLICT (publisher_id, scope) DO UPDATE
        SET index_manual = EXCLUDED.index_manual,
            is_locked = COALESCE(:lk, publisher_balance_index.is_locked),
            note = EXCLUDED.note, updated_at = now(), updated_by = EXCLUDED.updated_by
    """), {"p": publisher_id, "s": scope, "im": payload.index_manual,
           "lk": payload.is_locked, "n": payload.note, "u": user.id})
    db.commit()
    log_action(db, user, "balancer_row_edit", "sales_publisher", publisher_id,
               f"{scope}: объём={payload.volume} глубина={payload.depth} "
               f"запросы={payload.requests} индекс_рука={payload.index_manual}")
    return {"ok": True, "rows": balance.rows(db)}


@router.post("/balancer/recalc")
def balancer_recalc(db: Session = Depends(get_db), user: User = Depends(EDIT)):
    res = balance.recalc(db, user.id)
    log_action(db, user, "balancer_recalc", "sales_publisher", None,
               f"пересчёт индексов: {res['updated']}, заперто {res['locked_skipped']}, "
               f"без данных {res['no_data']}")
    return {**res, "rows": balance.rows(db)}


@router.put("/balancer/settings")
def balancer_settings(payload: CoefficientsIn, db: Session = Depends(get_db),
                      user: User = Depends(EDIT)):
    coef = balance.set_coefficients(db, payload.k, payload.depth_default)
    log_action(db, user, "balancer_settings", "sales_publisher", None,
               f"K={coef['k']} глубина_по_умолчанию={coef['depth_default']}")
    return {"coefficients": coef, "rows": balance.rows(db)}


# ── Excel: выгрузка и загрузка ────────────────────────────────────────────

BALANCE_COLS = [
    ("publisher_id", "ID"), ("scope", "Поверхность (код)"), ("code", "Наш код"),
    ("name", "Площадка"), ("domain", "Домен"), ("ms_publisher_id", "ID в МС"),
    ("scope_label", "Поверхность"), ("services", "Услуги"),
    ("volume", "Объём"), ("depth", "Глубина"), ("requests", "Запросы кода"),
    ("index_auto", "Индекс расчётный"), ("index_manual", "Индекс ручной"),
    ("source", "Источник"), ("is_locked", "Заперт"), ("note", "Примечание"),
]
# Импортом правятся только эти; остальные колонки справочные (из каталога паблишеров).
BALANCE_EDITABLE = ("volume", "depth", "requests", "index_manual", "is_locked", "note")

BLOCK_COLS = [("publisher", "Площадка"), ("code", "Наш код"), ("surface", "Поверхность"),
              ("ms_publisher_id", "ID паблишера в МС"), ("ms_block_id", "ID блока"),
              ("name", "Название в xoalt"), ("page_type", "Раздел"),
              ("network", "Сеть"), ("is_active", "Активен")]


def _xlsx(sheet_title: str, headers: List[str], rows: List[list]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, min(38, len(str(h)) + 6))
    buf = io.BytesIO()
    from app.xlsx_safe import save_workbook   # формулы только наши (аудит, 1.L7)
    save_workbook(wb, buf)
    return buf.getvalue()


def _xlsx_response(data: bytes, filename: str) -> Response:
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


def _num(v):
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(" ", "").replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


@router.get("/balancer/export")
def balancer_export(db: Session = Depends(get_db), user: User = Depends(VIEW)):
    out = []
    for r in balance.rows(db):
        row = []
        for k, _ in BALANCE_COLS:
            if k == "services":
                row.append(", ".join(r["services"]))
            elif k == "is_locked":
                row.append("да" if r[k] else "")
            else:
                row.append(r.get(k))
        out.append(row)
    return _xlsx_response(_xlsx("Балансировщик", [t for _, t in BALANCE_COLS], out),
                          "Балансировщик.xlsx")


@router.post("/balancer/import")
def balancer_import(file: UploadFile = File(...), db: Session = Depends(get_db),
                    user: User = Depends(EDIT)):
    """Загрузка правок из выгруженного файла. Ключ строки — колонки «ID» и «Поверхность (код)»;
    берём только редактируемые поля, справочные колонки игнорируем."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file.file.read()), data_only=True)
    ws = wb.active
    head = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    title_to_key = {t: k for k, t in BALANCE_COLS}
    pos = {title_to_key[t]: i for i, t in enumerate(head) if t in title_to_key}
    if "publisher_id" not in pos or "scope" not in pos:
        raise HTTPException(400, "В файле нет колонок «ID» и «Поверхность (код)» — "
                                 "загружайте файл, полученный выгрузкой")

    applied = skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        def cell(key, _row=row):
            i = pos.get(key)
            return _row[i] if i is not None and i < len(_row) else None

        pid, scope = cell("publisher_id"), cell("scope")
        if not pid or scope not in balance.SCOPES:
            skipped += 1
            continue
        pid = int(pid)
        vol, dep, req = _num(cell("volume")), _num(cell("depth")), _num(cell("requests"))
        if vol is not None or dep is not None:
            balance.upsert_measurement(db, pid, scope, vol, dep, source="import")
        if req is not None:
            balance.upsert_measurement(db, pid, balance.REQ_SCOPE[scope], req, source="import")
        locked_raw = str(cell("is_locked") or "").strip().lower()
        db.execute(text("""
            INSERT INTO publisher_balance_index
                (publisher_id, scope, index_manual, is_locked, note, updated_at, updated_by)
            VALUES (:p, :s, :im, :lk, :n, now(), :u)
            ON CONFLICT (publisher_id, scope) DO UPDATE
            SET index_manual = EXCLUDED.index_manual, is_locked = EXCLUDED.is_locked,
                note = EXCLUDED.note, updated_at = now(), updated_by = EXCLUDED.updated_by
        """), {"p": pid, "s": scope, "im": _num(cell("index_manual")),
               "lk": locked_raw in ("да", "yes", "true", "1", "y"),
               "n": (cell("note") or None), "u": user.id})
        applied += 1
    db.commit()
    log_action(db, user, "balancer_import", "sales_publisher", None,
               f"импорт балансировщика: применено {applied}, пропущено {skipped}")
    return {"applied": applied, "skipped": skipped, "rows": balance.rows(db)}


@router.get("/blocks/export")
def blocks_export(db: Session = Depends(get_db), user: User = Depends(VIEW)):
    rows = db.execute(text("""
        SELECT p.name AS publisher, p.code, s.kind AS surface, s.ms_publisher_id,
               b.ms_block_id, b.name, b.page_type, b.network, b.is_active
        FROM publisher_block b
        JOIN sales_publisher_surfaces s ON s.id = b.surface_id
        JOIN sales_publishers p ON p.id = b.publisher_id
        WHERE p.status <> :arch
        ORDER BY lower(p.name), s.kind, b.page_type, b.ms_block_id
    """), {"arch": PUBLISHER_ARCHIVE_STATUS}).mappings().all()
    out = [[("да" if r["is_active"] else "") if k == "is_active" else r[k]
            for k, _ in BLOCK_COLS] for r in rows]
    return _xlsx_response(_xlsx("Блоки", [t for _, t in BLOCK_COLS], out), "Каталог блоков.xlsx")
