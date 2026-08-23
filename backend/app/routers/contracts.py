import os
import re
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse, FileResponse

UPLOADS_DIR = "/app/uploads/contracts"
MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # 20 МБ — максимальный размер прикреплённого файла
ALLOWED_EXTENSIONS = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".jpg", ".jpeg", ".png", ".zip"}
from sqlalchemy.orm import Session
from app.database import get_db
from app.xlsx_safe import xlsx_safe
from app.models import Contract, Counterparty, User
from app.permissions import require_permission
from app.audit import log_action
from pydantic import BaseModel
from typing import List, Optional
from datetime import date

router = APIRouter()


# ===================== Реестр договоров (Настройки → Справочники) =====================
# Counterparty — единый источник данных для всех полей "контрагент" в системе (см.
# CLAUDE.md). Contract.counterparty_id — FK на Counterparty; при создании договора
# обязателен (см. create_contract), при редактировании может оставаться NULL для
# исторических строк, ещё не сопоставленных с реестром (см.
# link_contracts_to_counterparties.py). Когда counterparty_id указан,
# counterparty_name/inn — это денормализованный снимок, который сервер ВСЕГДА
# перезаписывает из канонической записи Counterparty (см. _resolve_counterparty) —
# клиентский свободный текст в этих полях для привязанных строк игнорируется.
# Для НЕпривязанных строк (counterparty_id is None) counterparty_name/inn остаются
# обычным свободным текстом, как было раньше.

class ContractCreate(BaseModel):
    contract_number: Optional[str] = None
    contract_date: Optional[date] = None
    inn: Optional[str] = None
    counterparty_id: Optional[int] = None
    counterparty_name: Optional[str] = None
    marketing_name: Optional[str] = None
    cooperation_format: Optional[str] = None
    services: Optional[str] = None
    end_date_text: Optional[str] = None
    prolongation: Optional[str] = None
    payment_form: Optional[str] = None
    payment_term_days: Optional[int] = None
    payment_term_condition: Optional[str] = None
    note: Optional[str] = None
    document_link: Optional[str] = None


# Списки для выпадающих списков на фронте (settings.js) — здесь НЕ валидируются
# строго (Pydantic Literal), чтобы не сломать сохранение старых строк, чьи
# значения cooperation_format/prolongation не входят в текущий список (см.
# CLAUDE.md: то же решение принято для vat_rate/VAT_OPTIONS). Ограничение —
# только на фронте через <select>.
class ContractBulkUpdate(BaseModel):
    ids: List[int]
    cooperation_format: Optional[str] = None
    prolongation: Optional[str] = None
    payment_term_days: Optional[int] = None
    payment_term_condition: Optional[str] = None


def _clean(s: Optional[str]) -> Optional[str]:
    if s is None:
        return None
    s = s.strip()
    return s or None



from app.links import validate_link as _validate_link  # общая проверка, см. app/links.py


def _resolve_counterparty(db: Session, counterparty_id: Optional[int]) -> Optional[Counterparty]:
    if not counterparty_id:
        return None
    cp = db.query(Counterparty).filter(Counterparty.id == counterparty_id).first()
    if not cp:
        raise HTTPException(status_code=400, detail="Контрагент не найден в реестре контрагентов")
    return cp


def _serialize(c: Contract) -> dict:
    return {
        "id": c.id,
        "contract_number": c.contract_number,
        "contract_date": c.contract_date.isoformat() if c.contract_date else None,
        "inn": c.inn,
        "counterparty_id": c.counterparty_id,
        "counterparty_name": c.counterparty_name,
        "linked": c.counterparty_id is not None,
        # НДС контрагента по направлениям — read-only производные из реестра контрагентов
        # (как contracts_count в обратную сторону), редактируются в карточке контрагента.
        "vat_rate_income": c.counterparty.vat_rate_income if c.counterparty else None,
        "vat_rate_expense": c.counterparty.vat_rate_expense if c.counterparty else None,
        "marketing_name": c.marketing_name,
        "cooperation_format": c.cooperation_format,
        "services": c.services,
        "end_date_text": c.end_date_text,
        "prolongation": c.prolongation,
        "payment_form": c.payment_form,
        "payment_term_days": c.payment_term_days,
        "payment_term_condition": c.payment_term_condition,
        "payment_term_legacy": c.payment_term,  # старое текстовое поле — только для справки, не редактируется
        "note": c.note,
        "document_link": c.document_link,
        "attached_filename": c.attached_filename,
    }


@router.get("/registry")
def get_contracts_registry(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "view"))
):
    items = db.query(Contract).order_by(Contract.contract_date.desc().nullslast(), Contract.id.desc()).all()
    return {"items": [_serialize(c) for c in items]}


@router.patch("/bulk")
def bulk_update_contracts(
    payload: ContractBulkUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "edit"))
):
    if not payload.ids:
        raise HTTPException(status_code=400, detail="Не выбраны договоры")

    fields = payload.dict(exclude={"ids"}, exclude_unset=True)
    if "cooperation_format" in fields:
        fields["cooperation_format"] = _clean(fields["cooperation_format"])
    if "prolongation" in fields:
        fields["prolongation"] = _clean(fields["prolongation"])
    if "payment_term_condition" in fields:
        fields["payment_term_condition"] = _clean(fields["payment_term_condition"])
    if not fields:
        raise HTTPException(status_code=400, detail="Не указаны поля для изменения")

    contracts = db.query(Contract).filter(Contract.id.in_(payload.ids)).all()
    if not contracts:
        raise HTTPException(status_code=404, detail="Договоры не найдены")

    for c in contracts:
        for key, value in fields.items():
            setattr(c, key, value)
    db.commit()

    labels = {
        "cooperation_format": "формат сотрудничества", "prolongation": "пролонгация",
        "payment_term_days": "срок оплаты, дни", "payment_term_condition": "условие",
    }
    changes_str = "; ".join(f"{labels[k]} → {v or '—'}" for k, v in fields.items())
    ids_str = ", ".join(str(c.id) for c in contracts)
    log_action(db, current_user, "bulk_update_contract", entity_type="contract", entity_id=None,
               details=f"Массовое изменение договоров [{ids_str}]: {changes_str}")
    return {"message": f"Обновлено договоров: {len(contracts)}"}


class ContractBulkDelete(BaseModel):
    ids: List[int]


@router.delete("/bulk")
def bulk_delete_contracts(
    payload: ContractBulkDelete,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "edit"))
):
    # Массовое удаление — отдельно от точечного DELETE /{id} (которое разрешено любому
    # с правом "edit" на раздел) — здесь дополнительно жёстко требуем роль admin, по
    # запросу пользователя: массовое удаление договоров — это admin-only операция.
    if current_user.role.key != "admin":
        raise HTTPException(status_code=403, detail="Массовое удаление договоров доступно только администратору")
    if not payload.ids:
        raise HTTPException(status_code=400, detail="Не выбраны договоры")

    contracts = db.query(Contract).filter(Contract.id.in_(payload.ids)).all()
    if not contracts:
        raise HTTPException(status_code=404, detail="Договоры не найдены")

    labels_str = "; ".join(f"{c.contract_number or '—'} ({c.counterparty_name or '—'})" for c in contracts)
    ids_str = ", ".join(str(c.id) for c in contracts)
    count = len(contracts)
    for c in contracts:
        db.delete(c)
    db.commit()

    log_action(db, current_user, "bulk_delete_contract", entity_type="contract", entity_id=None,
               details=f"Массовое удаление договоров [{ids_str}]: {labels_str}")
    return {"message": f"Удалено договоров: {count}"}


@router.post("/")
def create_contract(
    data: ContractCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "edit"))
):
    # Контрагент обязателен при создании НОВОГО договора — единый источник данных
    # (см. CLAUDE.md), никаких новых свободно-текстовых контрагентов через эту форму.
    # Проверка здесь, а не Pydantic-полем Field(...), т.к. ContractCreate общая модель
    # с update_contract, где counterparty_id может оставаться NULL у старых строк.
    if not data.counterparty_id:
        raise HTTPException(status_code=400, detail="Выберите контрагента из реестра контрагентов")
    cp = _resolve_counterparty(db, data.counterparty_id)

    contract = Contract(
        contract_number=_clean(data.contract_number),
        contract_date=data.contract_date,
        inn=cp.inn,
        counterparty_name=cp.name,
        counterparty_id=cp.id,
        marketing_name=_clean(data.marketing_name),
        cooperation_format=_clean(data.cooperation_format),
        services=_clean(data.services),
        end_date_text=_clean(data.end_date_text),
        prolongation=_clean(data.prolongation),
        payment_form=_clean(data.payment_form),
        payment_term_days=data.payment_term_days,
        payment_term_condition=_clean(data.payment_term_condition),
        note=_clean(data.note),
        document_link=_validate_link(data.document_link),
    )
    db.add(contract)
    db.commit()
    db.refresh(contract)

    log_action(db, current_user, "create_contract", entity_type="contract", entity_id=contract.id,
               details=f"Создан договор № {contract.contract_number or '—'} ({contract.counterparty_name or '—'})")
    return {"id": contract.id, "message": "Договор создан"}


@router.put("/{contract_id}")
def update_contract(
    contract_id: int,
    data: ContractCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "edit"))
):
    contract = db.query(Contract).filter(Contract.id == contract_id).first()
    if not contract:
        raise HTTPException(status_code=404, detail="Договор не найден")

    # Если counterparty_id передан — контрагент привязан/перепривязан к реестру, и
    # counterparty_name/inn ВСЕГДА берутся из канонической записи Counterparty (клиентский
    # свободный текст в этих двух полях для привязанных строк игнорируется). Если
    # counterparty_id не передан (None) — это правка исторической непривязанной строки,
    # counterparty_name/inn остаются обычным свободным текстом, как раньше.
    if data.counterparty_id:
        cp = _resolve_counterparty(db, data.counterparty_id)
        resolved_name, resolved_inn, resolved_cp_id = cp.name, cp.inn, cp.id
    else:
        resolved_name, resolved_inn, resolved_cp_id = _clean(data.counterparty_name), _clean(data.inn), None

    new_values = {
        "contract_number": _clean(data.contract_number),
        "contract_date": data.contract_date,
        "inn": resolved_inn,
        "counterparty_name": resolved_name,
        "counterparty_id": resolved_cp_id,
        "marketing_name": _clean(data.marketing_name),
        "cooperation_format": _clean(data.cooperation_format),
        "services": _clean(data.services),
        "end_date_text": _clean(data.end_date_text),
        "prolongation": _clean(data.prolongation),
        "payment_form": _clean(data.payment_form),
        "payment_term_days": data.payment_term_days,
        "payment_term_condition": _clean(data.payment_term_condition),
        "note": _clean(data.note),
        "document_link": _validate_link(data.document_link),
    }

    changes = []
    labels = {
        "contract_number": "№ договора", "contract_date": "дата договора", "inn": "ИНН",
        "counterparty_name": "контрагент", "counterparty_id": "привязка к реестру",
        "marketing_name": "маркетинговое название",
        "cooperation_format": "формат сотрудничества", "services": "услуги",
        "end_date_text": "дата окончания", "prolongation": "пролонгация",
        "payment_form": "форма оплаты", "payment_term_days": "срок оплаты, дни",
        "payment_term_condition": "условие", "note": "примечание",
        "document_link": "ссылка на документ",
    }
    for key, val in new_values.items():
        old = getattr(contract, key)
        if val != old:
            changes.append(f"{labels[key]}: {old or '—'} → {val or '—'}")
            setattr(contract, key, val)

    db.commit()
    if changes:
        log_action(db, current_user, "update_contract", entity_type="contract", entity_id=contract.id,
                   details="; ".join(changes))
    return {"message": "Договор обновлён"}


@router.delete("/{contract_id}")
def delete_contract(
    contract_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "edit"))
):
    contract = db.query(Contract).filter(Contract.id == contract_id).first()
    if not contract:
        raise HTTPException(status_code=404, detail="Договор не найден")
    label = f"{contract.contract_number or '—'} ({contract.counterparty_name or '—'})"
    db.delete(contract)
    db.commit()

    log_action(db, current_user, "delete_contract", entity_type="contract", entity_id=contract_id,
               details=f"Удалён договор {label}")
    return {"message": "Договор удалён"}


# ===================== Прикреплённые документы =====================
# Файлы хранятся в /app/uploads/contracts/ (volume-mount из docker-compose.yml:
# ./uploads:/app/uploads). Имя файла на диске: {contract_id}_{sanitized_original_name}.
# attached_filename в БД содержит это имя; при скачивании клиент получает оригинальное
# имя (без префикса contract_id_). Замена файла — upload просто перезаписывает старый.

@router.post("/{contract_id}/upload")
async def upload_contract_document(
    contract_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "edit"))
):
    contract = db.query(Contract).filter(Contract.id == contract_id).first()
    if not contract:
        raise HTTPException(status_code=404, detail="Договор не найден")

    os.makedirs(UPLOADS_DIR, exist_ok=True)

    # Удаляем старый файл, если был
    if contract.attached_filename:
        old_path = os.path.join(UPLOADS_DIR, contract.attached_filename)
        if os.path.exists(old_path):
            os.remove(old_path)

    # Санитизация имени файла: оставляем только безопасные символы
    original_name = file.filename or "document"
    ext = os.path.splitext(original_name)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=415,
            detail=f"Недопустимый тип файла. Разрешены: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )
    safe_name = re.sub(r'[^\w\.\-]', '_', original_name)
    stored_name = f"{contract_id}_{safe_name}"

    content = await file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"Файл слишком большой (максимум {MAX_UPLOAD_BYTES // 1024 // 1024} МБ)"
        )
    with open(os.path.join(UPLOADS_DIR, stored_name), "wb") as f:
        f.write(content)

    contract.attached_filename = stored_name
    db.commit()

    log_action(db, current_user, "upload_contract_document", entity_type="contract",
               entity_id=contract_id, details=f"Прикреплён документ: {original_name}")
    return {"filename": stored_name, "message": "Документ прикреплён"}


@router.get("/{contract_id}/download")
def download_contract_document(
    contract_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "view"))
):
    contract = db.query(Contract).filter(Contract.id == contract_id).first()
    if not contract or not contract.attached_filename:
        raise HTTPException(status_code=404, detail="Документ не найден")

    file_path = os.path.join(UPLOADS_DIR, contract.attached_filename)
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="Файл не найден на сервере")

    # Отдаём с оригинальным именем (без префикса contract_id_)
    original_name = contract.attached_filename
    prefix = f"{contract_id}_"
    if original_name.startswith(prefix):
        original_name = original_name[len(prefix):]

    return FileResponse(file_path, filename=original_name,
                        media_type="application/octet-stream")


@router.delete("/{contract_id}/document")
def delete_contract_document(
    contract_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "edit"))
):
    contract = db.query(Contract).filter(Contract.id == contract_id).first()
    if not contract:
        raise HTTPException(status_code=404, detail="Договор не найден")

    filename = contract.attached_filename
    if filename:
        file_path = os.path.join(UPLOADS_DIR, filename)
        if os.path.exists(file_path):
            os.remove(file_path)
        contract.attached_filename = None
        db.commit()
        log_action(db, current_user, "delete_contract_document", entity_type="contract",
                   entity_id=contract_id, details=f"Удалён прикреплённый документ: {filename}")

    return {"message": "Документ удалён"}


# ===================== Экспорт / Импорт договоров (Excel) =====================

# Значения выпадающих списков — дублируют COOPERATION_FORMATS / PROLONGATION_OPTIONS /
# PAYMENT_TERM_CONDITIONS из frontend/pages/directories.js; при изменении менять в обоих местах.
_COOPERATION_FORMATS    = ['Агентство КЛ', 'Агентство ПД', 'Клиент', 'Подрядчик', 'Аптека', 'Паблишер', 'Рекламная система']
_PROLONGATION_OPTIONS   = ['АВТО на год', 'По соглашению', 'Нет']
_PAYMENT_TERM_CONDS     = ['С даты УПД', 'С даты АКТ', 'По периоду']

# Колонки Excel: (поле_модели, заголовок, ширина, только_чтение)
_EXPORT_COLS = [
    ('id',                    'ID (не менять)',          8,  True),
    ('contract_number',       '№ договора',             18,  False),
    ('contract_date',         'Дата договора',          14,  False),
    ('counterparty_name',     'Контрагент (не менять)', 28,  True),
    ('inn',                   'ИНН (не менять)',        14,  True),
    ('marketing_name',        'Маркетинговое название', 24,  False),
    ('cooperation_format',    'Формат сотрудничества',  22,  False),  # dropdown G
    ('services',              'Услуги',                 20,  False),
    ('end_date_text',         'Дата окончания',         14,  False),
    ('prolongation',          'Пролонгация',            16,  False),  # dropdown J
    ('payment_form',          'Форма оплаты',           16,  False),
    ('payment_term_days',     'Срок оплаты, дни',       16,  False),
    ('payment_term_condition','Условие оплаты',         18,  False),  # dropdown M
    ('note',                  'Примечание',             28,  False),
    ('document_link',         'Ссылка на документ',     32,  False),
]

# Поля, которые МОЖНО менять через импорт (read-only колонки игнорируются)
_EDITABLE_FIELDS = {
    'contract_number', 'contract_date', 'marketing_name', 'cooperation_format',
    'services', 'end_date_text', 'prolongation', 'payment_form',
    'payment_term_days', 'payment_term_condition', 'note', 'document_link',
}

_FIELD_LABELS = {
    'contract_number':       '№ договора',
    'contract_date':         'Дата договора',
    'marketing_name':        'Маркетинговое название',
    'cooperation_format':    'Формат сотрудничества',
    'services':              'Услуги',
    'end_date_text':         'Дата окончания',
    'prolongation':          'Пролонгация',
    'payment_form':          'Форма оплаты',
    'payment_term_days':     'Срок оплаты, дни',
    'payment_term_condition':'Условие оплаты',
    'note':                  'Примечание',
    'document_link':         'Ссылка на документ',
}


@router.get("/export")
def export_contracts(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "view"))
):
    import io
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation

    contracts_list = db.query(Contract).order_by(Contract.id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Договоры"

    # Скрытый лист со списками для DataValidation
    ws_ref = wb.create_sheet("Справочники")
    ws_ref.sheet_state = "hidden"
    ws_ref["A1"] = "Форматы сотрудничества"
    for i, v in enumerate(_COOPERATION_FORMATS, 2):
        ws_ref.cell(row=i, column=1, value=v)
    ws_ref["B1"] = "Пролонгация"
    for i, v in enumerate(_PROLONGATION_OPTIONS, 2):
        ws_ref.cell(row=i, column=2, value=v)
    ws_ref["C1"] = "Условие оплаты"
    for i, v in enumerate(_PAYMENT_TERM_CONDS, 2):
        ws_ref.cell(row=i, column=3, value=v)

    # Шапка
    header_font   = Font(bold=True, color="FFFFFF")
    fill_edit     = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    fill_readonly = PatternFill(start_color="6B7280", end_color="6B7280", fill_type="solid")

    for ci, (field, header, width, readonly) in enumerate(_EXPORT_COLS, 1):
        cell = ws.cell(row=1, column=ci, value=header)
        cell.font = header_font
        cell.fill = fill_readonly if readonly else fill_edit
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    # Данные
    fill_ro_data = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    for ri, c in enumerate(contracts_list, 2):
        vals = {
            'id': c.id, 'contract_number': c.contract_number, 'contract_date': c.contract_date,
            'counterparty_name': c.counterparty_name, 'inn': c.inn,
            'marketing_name': c.marketing_name, 'cooperation_format': c.cooperation_format,
            'services': c.services, 'end_date_text': c.end_date_text,
            'prolongation': c.prolongation, 'payment_form': c.payment_form,
            'payment_term_days': c.payment_term_days, 'payment_term_condition': c.payment_term_condition,
            'note': c.note, 'document_link': c.document_link,
        }
        for ci, (field, _, _, readonly) in enumerate(_EXPORT_COLS, 1):
            cell = ws.cell(row=ri, column=ci, value=xlsx_safe(vals[field]))
            if readonly:
                cell.fill = fill_ro_data
            if field == 'contract_date' and vals[field]:
                cell.number_format = 'DD.MM.YYYY'

    # DataValidation (dropdowns; errorStyle не задаём — разрешаем свободный ввод, как SelectWithOther)
    max_dv = max(len(contracts_list) + 1, 1000)

    dv_fmt = DataValidation(type="list",
        formula1=f"=Справочники!$A$2:$A${len(_COOPERATION_FORMATS)+1}",
        allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv_fmt)
    dv_fmt.sqref = f"G2:G{max_dv}"

    dv_prl = DataValidation(type="list",
        formula1=f"=Справочники!$B$2:$B${len(_PROLONGATION_OPTIONS)+1}",
        allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv_prl)
    dv_prl.sqref = f"J2:J{max_dv}"

    dv_cnd = DataValidation(type="list",
        formula1=f"=Справочники!$C$2:$C${len(_PAYMENT_TERM_CONDS)+1}",
        allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv_cnd)
    dv_cnd.sqref = f"M2:M{max_dv}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"dogovory_{ts}.xlsx"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"})


def _parse_import_file(content: bytes):
    """
    Парсит Excel-файл экспорта договоров.
    Возвращает список dict-строк с ключами по именам полей.
    """
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    # Строим маппинг заголовок → индекс
    header_row = rows[0]
    header_to_field = {}
    for idx, h in enumerate(header_row):
        for field, label, _, _ in _EXPORT_COLS:
            # заголовок содержит "(не менять)" — обрезаем для сравнения
            clean = (h or "").replace(" (не менять)", "").strip()
            if clean == label.replace(" (не менять)", "").strip():
                header_to_field[idx] = field
                break

    result = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        rec = {}
        for idx, val in enumerate(row):
            field = header_to_field.get(idx)
            if field:
                rec[field] = val
        if rec.get("id") is not None:
            result.append(rec)
    return result


def _coerce_import_value(field: str, raw):
    """Приводит значение из Excel к типу поля модели."""
    import datetime as dt
    if raw is None or raw == "":
        return None
    if field == "payment_term_days":
        try:
            return int(float(str(raw)))
        except Exception:
            return None
    if field == "contract_date":
        if isinstance(raw, (dt.date, dt.datetime)):
            return raw.date() if isinstance(raw, dt.datetime) else raw
        # строка в формате DD.MM.YYYY
        for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(str(raw).strip(), fmt).date()
            except Exception:
                pass
        return None
    return str(raw).strip() if raw is not None else None


@router.post("/import/preview")
async def preview_import_contracts(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "edit"))
):
    content = await file.read()
    try:
        rows = _parse_import_file(content)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Ошибка чтения файла: {e}")

    changes = []
    skipped = []

    for row_num, rec in enumerate(rows, 2):
        raw_id = rec.get("id")
        try:
            contract_id = int(raw_id)
        except (TypeError, ValueError):
            skipped.append({"row": row_num, "reason": f"Некорректный ID: {raw_id}"})
            continue

        contract = db.query(Contract).filter(Contract.id == contract_id).first()
        if not contract:
            skipped.append({"row": row_num, "reason": f"Договор ID {contract_id} не найден"})
            continue

        diffs = []
        for field in _EDITABLE_FIELDS:
            if field not in rec:
                continue
            new_val = _coerce_import_value(field, rec[field])
            if field == 'document_link':
                new_val = _validate_link(new_val, raise_on_bad=False)  # обход формы: санируем и здесь
            old_val = getattr(contract, field)
            # Нормализуем None / "" для сравнения
            old_norm = old_val if old_val is not None else None
            new_norm = new_val
            if str(old_norm or "") != str(new_norm or ""):
                diffs.append({
                    "field": field,
                    "label": _FIELD_LABELS.get(field, field),
                    "old": str(old_val) if old_val is not None else "",
                    "new": str(new_val) if new_val is not None else "",
                })

        if diffs:
            changes.append({
                "id": contract_id,
                "contract_number": contract.contract_number or "—",
                "counterparty_name": contract.counterparty_name or "—",
                "diffs": diffs,
            })

    return {"changes": changes, "skipped": skipped}


@router.post("/import/apply")
async def apply_import_contracts(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("contracts", "edit"))
):
    content = await file.read()
    try:
        rows = _parse_import_file(content)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Ошибка чтения файла: {e}")

    updated = 0
    skipped = []

    for row_num, rec in enumerate(rows, 2):
        raw_id = rec.get("id")
        try:
            contract_id = int(raw_id)
        except (TypeError, ValueError):
            skipped.append(f"Строка {row_num}: некорректный ID {raw_id}")
            continue

        contract = db.query(Contract).filter(Contract.id == contract_id).first()
        if not contract:
            skipped.append(f"Строка {row_num}: ID {contract_id} не найден")
            continue

        changes_log = []
        for field in _EDITABLE_FIELDS:
            if field not in rec:
                continue
            new_val = _coerce_import_value(field, rec[field])
            if field == 'document_link':
                new_val = _validate_link(new_val, raise_on_bad=False)  # обход формы: санируем и здесь
            old_val = getattr(contract, field)
            if str(old_val or "") != str(new_val or ""):
                changes_log.append(f"{_FIELD_LABELS.get(field, field)}: {old_val or '—'} → {new_val or '—'}")
                setattr(contract, field, new_val)

        if changes_log:
            db.commit()
            log_action(db, current_user, "import_update_contract", entity_type="contract",
                       entity_id=contract_id,
                       details=f"Импорт Excel, договор #{contract.contract_number or contract_id}: " + "; ".join(changes_log))
            updated += 1

    return {"updated": updated, "skipped": skipped}
