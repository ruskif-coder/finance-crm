# -*- coding: utf-8 -*-
"""«План продаж» в Excel (владелец 24.10.2026): оформленный отчёт и шаблон целей, загрузка правок обратно.

Лист «Выполнение» — отчёт, как экран: шапка, разделы «Продажники» / «Услуги» (итог группы), по строке — цель, выполнено и
прогноз по месяцам и за год, % к цели с заливкой тем же порогом, что на экране (< 80 / 80–99 / ≥ 100 / ≥ 110). Суммы —
тыс. ₽ без НДС. Лист «Цели» — шаблон правки в полных рублях: скрытый код строки (`разрез:id`), 12 месяцев, год формулой.
Загрузка не пишет сама: `preview` возвращает «было → станет» и ошибки строк, применяет экран через `plan.save_targets`
(журнал каждой правки). Пустая ячейка — «не менять»."""
from __future__ import annotations

import io
from datetime import datetime
from typing import Optional

from fastapi import HTTPException

from app.sales import plan

MONTHS = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]
EDITABLE = ("общий", "услуга", "продажник")
MAX_BYTES = 5 * 1024 * 1024
RUB = '#,##0'
THOUSANDS = '#,##0'
PCT = '0%'
# Палитра наших переменных (styles/globals.css): в Excel переменных нет, значения — те же, что у светлой темы
INK, MUTED, CAP = "1C2433", "79839A", "6B7488"
ACCENT_TINT, SUBTLE, BORDER = "ECEFFD", "F6F7FB", "E3E7F1"
HEAT = {"bad": "FBE4E6", "mid": "FBF0DE", "ok": "E6F5EF", "over": "CDEADF"}


def _heat(p: Optional[float]) -> Optional[str]:
    if p is None:
        return None
    return "over" if p >= 1.1 else "ok" if p >= 1 else "mid" if p >= 0.8 else "bad"


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color=BORDER)
    return {
        "title": Font(name="Calibri", size=16, bold=True, color=INK),
        "sub": Font(name="Calibri", size=10, color=MUTED),
        "head": Font(name="Calibri", size=9, bold=True, color=CAP),
        "head_fill": PatternFill("solid", fgColor=ACCENT_TINT),
        "sec": Font(name="Calibri", size=9, bold=True, color=CAP),
        "sec_fill": PatternFill("solid", fgColor=SUBTLE),
        "bold": Font(name="Calibri", size=11, bold=True, color=INK),
        "norm": Font(name="Calibri", size=11, color=INK),
        "muted": Font(name="Calibri", size=10, color=MUTED),
        "line": Border(bottom=thin),
        "right": Alignment(horizontal="right", vertical="center"),
        "left": Alignment(horizontal="left", vertical="center", indent=1),
        "edit_fill": PatternFill("solid", fgColor="FFFFFF"),
    }


def _sections(rows):
    """Строки отчёта в порядке экрана: общая → продажники (+ без) → услуги по группам с итогом (+ без)."""
    total = [r for r in rows if r["scope"] == "общий"]
    reps = [r for r in rows if r["scope"] in ("продажник", "без продажника")]
    svc, groups = [r for r in rows if r["scope"] == "услуга"], []
    for r in svc:
        g = next((x for x in groups if x["group"] == r.get("group")), None)
        if g is None:
            groups.append(g := {"group": r.get("group"), "rows": []})
        g["rows"].append(r)
    out = [("", total)]
    if reps:
        out.append(("Продажники", reps))
    if groups or any(r["scope"] == "без услуги" for r in rows):
        body = []
        for g in groups:
            sub = {"scope": "группа", "name": f"{g['group']} · итог группы", "months": [
                {k: (sum((x["months"][m][k] or 0) for x in g["rows"]) if k != "goal" else
                     (sum(x["months"][m]["goal"] for x in g["rows"] if x["months"][m]["goal"] is not None)
                      if any(x["months"][m]["goal"] is not None for x in g["rows"]) else None))
                 for k in ("goal", "done", "forecast")} for m in range(12)]}
            body += [sub] + [{**r, "indent": True} for r in g["rows"]]
        body += [r for r in rows if r["scope"] == "без услуги"]
        out.append(("Услуги", body))
    return out


def _report(wb, rows, year: int, today: str):
    from openpyxl.styles import PatternFill
    st = _styles()
    ws = wb.active
    ws.title = "Выполнение"
    ws.sheet_view.showGridLines = False
    cur = int(today[5:7]) - 1 if int(today[:4]) == year else (12 if int(today[:4]) > year else -1)
    ws["A1"] = f"План продаж {year}"
    ws["A1"].font = st["title"]
    ws["A2"] = (f"тыс. ₽ без НДС · выполнено = факт + в работе (со «Сборки»), прогноз = + планируемые (зелёный 100 %, "
                f"оранжевый 50 %) · прошлые и текущий месяц — выполнено к цели, будущие — прогноз к цели · на {today[8:10]}.{today[5:7]}.{today[:4]}")
    ws["A2"].font = st["sub"]
    hr = 4
    for c, v in enumerate(["Строка", "Показатель"] + MONTHS + ["Год"], start=1):
        cell = ws.cell(hr, c, v)
        cell.font, cell.fill, cell.border = st["head"], st["head_fill"], st["line"]
        cell.alignment = st["left"] if c <= 2 else st["right"]
        if 0 <= c - 3 == cur:
            cell.font = st["bold"]
    r = hr + 1
    for title, body in _sections(rows):
        if title:
            for c in range(1, 16):
                ws.cell(r, c).fill = st["sec_fill"]
            ws.cell(r, 1, title.upper()).font = st["sec"]
            r += 1
        for row in body:
            outside = row["scope"] in ("без продажника", "без услуги")
            name = ("    " if row.get("indent") else "") + row["name"] + (" (неактивен)" if row.get("inactive") else "")
            ws.cell(r, 1, name).font = st["bold"] if row["scope"] in ("общий", "группа") else st["norm"]
            ms = row["months"]
            y_goal = sum(m["goal"] for m in ms if m["goal"] is not None) if any(m["goal"] is not None for m in ms) else None
            y_done, y_fc = sum(m["done"] for m in ms), sum(m["forecast"] for m in ms)
            lines = [("цель", [m["goal"] for m in ms], y_goal), ("выполнено", [m["done"] for m in ms], y_done),
                     ("прогноз", [m["forecast"] for m in ms], y_fc), ("% к цели", None, None)]
            for i, (label, vals, yv) in enumerate(lines):
                rr = r + i
                ws.cell(rr, 2, label).font = st["muted"]
                for m in range(12):
                    cell = ws.cell(rr, 3 + m)
                    if label == "% к цели":
                        base = ms[m]["forecast"] if m > cur else ms[m]["done"]
                        p = (base / ms[m]["goal"]) if ms[m]["goal"] else None
                        cell.value = p
                        cell.number_format = PCT
                        h = None if outside else _heat(p)
                        if h:
                            cell.fill = PatternFill("solid", fgColor=HEAT[h])
                    else:
                        v = vals[m]
                        cell.value = None if v is None else round(v / 1000)
                        cell.number_format = THOUSANDS
                    cell.alignment = st["right"]
                yc = ws.cell(rr, 15)
                if label == "% к цели":
                    p = (y_fc / y_goal) if y_goal else None
                    yc.value, yc.number_format = p, PCT
                    if (h := None if outside else _heat(p)):
                        yc.fill = PatternFill("solid", fgColor=HEAT[h])
                else:
                    yc.value, yc.number_format = (None if yv is None else round(yv / 1000)), THOUSANDS
                yc.font, yc.alignment = st["bold"], st["right"]
            for c in range(1, 16):
                ws.cell(r + 3, c).border = st["line"]
            r += 4
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12
    for c in "CDEFGHIJKLMN":
        ws.column_dimensions[c].width = 10
    ws.column_dimensions["O"].width = 12
    ws.freeze_panes = ws.cell(hr + 1, 3)


GOALS_HEADER = ["Код", "Разрез", "Название"] + MONTHS + ["Год"]


def _goals(wb, rows, year: int):
    from openpyxl.styles import PatternFill
    st = _styles()
    ws = wb.create_sheet("Цели")
    ws.sheet_view.showGridLines = False
    ws["B1"] = f"Цели {year} — шаблон для правки"
    ws["B1"].font = st["title"]
    ws["B2"] = ("Полные рубли без НДС по месяцам, например 10 000 000. Пустая ячейка при загрузке — «не менять». Колонку «Год» "
                "не правьте: она считается. Загрузка — «План продаж» → «Цели» → кнопка загрузки из Excel; сначала покажется, что изменится.")
    ws["B2"].font = st["sub"]
    hr = header_row_default = 4
    for c, v in enumerate(GOALS_HEADER, start=1):
        cell = ws.cell(hr, c, v)
        cell.font, cell.fill, cell.border = st["head"], st["head_fill"], st["line"]
        cell.alignment = st["left"] if c <= 3 else st["right"]
    r = hr + 1
    for row in rows:
        if row["scope"] not in EDITABLE:
            continue
        ws.cell(r, 1, f"{row['scope']}:{row['key_id']}")
        ws.cell(r, 2, row["scope"]).font = st["muted"]
        ws.cell(r, 3, row["name"]).font = st["bold"] if row["scope"] == "общий" else st["norm"]
        for m in range(12):
            cell = ws.cell(r, 4 + m, row["months"][m]["goal"])
            cell.number_format, cell.alignment = RUB, st["right"]
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
            cell.border = st["line"]
        y = ws.cell(r, 16, f"=SUM(D{r}:O{r})")
        y.number_format, y.font, y.alignment = RUB, st["bold"], st["right"]
        y.fill = st["sec_fill"]
        r += 1
    ws.column_dimensions["A"].hidden = True          # код — для загрузки обратно, человеку не нужен
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 30
    for c in "DEFGHIJKLMNO":
        ws.column_dimensions[c].width = 14
    ws.column_dimensions["P"].width = 16
    ws.freeze_panes = ws.cell(header_row_default + 1, 4)


def export(db, year: int, own: Optional[list] = None) -> bytes:
    import openpyxl

    from app.xlsx_safe import save_workbook
    data = plan.screen(db, year, own=own)
    wb = openpyxl.Workbook()
    _report(wb, data["rows"], year, data["today"])
    _goals(wb, data["rows"], year)
    wb.properties.title = f"План продаж {year}"
    wb.properties.created = datetime.utcnow()
    buf = io.BytesIO()
    save_workbook(wb, buf)
    return buf.getvalue()


def header_row(ws) -> int:
    """Строка шапки листа «Цели»: ищется по «Код» в первой колонке (шапка ниже заголовка листа)."""
    for r in range(1, 21):
        if ws.cell(r, 1).value == "Код":
            return r
    return 1


def _num(v) -> Optional[float]:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).replace(" ", "").replace("\xa0", "").replace(",", "."))


def preview(db, year: int, body: bytes) -> dict:
    """Изменения «было → станет» из загруженного файла и ошибки строк. Ничего не пишет."""
    import openpyxl
    if len(body) > MAX_BYTES:
        raise HTTPException(413, "Файл больше 5 МБ")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(body), data_only=True)
        ws = wb["Цели"]
    except KeyError:
        raise HTTPException(400, "В файле нет листа «Цели» — загрузите файл, выгруженный с этого экрана")
    except Exception:
        raise HTTPException(400, "Не получилось прочитать файл — нужен .xlsx, выгруженный с этого экрана")
    goals = plan.targets_of(db, year)
    names = {(r["scope"], r["key_id"]): r["name"] for r in plan.screen(db, year)["rows"]}
    changes, errors = [], []
    for n, row in enumerate(ws.iter_rows(min_row=header_row(ws) + 1, values_only=True), start=header_row(ws) + 1):
        if not row or not row[0]:
            continue
        code = str(row[0]).strip()
        scope, _, key = code.partition(":")
        try:
            key_id = int(key)
        except ValueError:
            errors.append(f"строка {n}: непонятный код «{code}»")
            continue
        if scope not in EDITABLE or (scope, key_id) not in names:
            errors.append(f"строка {n}: «{row[2] or code}» ({code}) — такой строки плана нет")
            continue
        was = goals.get((scope, key_id)) or [None] * 12
        for m in range(12):
            cell = row[3 + m] if len(row) > 3 + m else None
            try:
                now = _num(cell)
            except ValueError:
                errors.append(f"строка {n}, {MONTHS[m]}: «{cell}» — не число")
                continue
            if now is None:
                continue
            if now < 0:
                errors.append(f"строка {n}, {MONTHS[m]}: цель не может быть отрицательной")
                continue
            now = round(now, 2)
            if was[m] != now:
                changes.append({"scope": scope, "key_id": key_id, "name": names[(scope, key_id)], "month": m + 1,
                                "was": was[m], "now": now})
    return {"year": year, "changes": changes, "errors": errors}
