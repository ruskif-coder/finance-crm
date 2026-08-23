"""
Перенос рабочей таблицы «Аптеки» в справочник паблишеров (разовый скрипт).

Запуск:
    docker cp "Аптеки - Рабочая.xlsx" finance_backend:/tmp/publishers.xlsx
    docker exec finance_backend python -m scripts.import_publishers /tmp/publishers.xlsx [--apply]

Без --apply только считает и печатает отчёт, в базу не пишет.

Почему сопоставление, а не создание: юрлица и договоры уже живут в реестрах финмодуля.
Ничего не угадываем — сопоставляем по ИНН и по номеру договора буквально; 0 или больше
одного кандидата означает строку в отчёте на ручную привязку, а не выбор наугад. Тот же
принцип, что в sync_contracts.py и link_contracts_to_counterparties.py.

Строк в исходнике 47, площадок 41: шесть сайтов заведены дважды — отдельной строкой на
WEB и на APP, причём в APP-строках все поля поверхностей пустые, вся фактура лежит в
WEB-строке. Поэтому группа строк по домену сливается «первым непустым значением».
"""
import re
import sys

from app.database import SessionLocal
from app.models import Counterparty, Contract
from app.sales.models import (SalesPublisher, SalesPublisherSurface, SalesPublisherContact,
                              SalesPublisherCounterparty, SalesPublisherContract,
                              normalize_domain)

# Колонки исходника (1-based), как в файле «Аптеки - Рабочая».
COL = {"site": 2, "status": 3, "network": 4, "deal_type": 5, "exclusive": 6, "dsp": 7,
       "surface": 8, "legal": 9, "inn": 10, "contract": 11, "contract_ag": 12, "cpm": 13,
       "connect": 14, "self_promo": 20, "self_promo_note": 21,
       "figma_web": 24, "status_web": 25, "coverage_web": 26,
       "figma_app": 29, "status_app": 30, "coverage_app": 31, "note_app": 32,
       "priority": 35, "basket": 36, "chat": 37, "messenger": 38, "contact": 39,
       "comment": 40, "tasks": 41, "comment2": 42,
       "contact_name": 43, "contact_email": 44, "contact_email2": 45, "contact_name2": 46}

STATUS_MAP = {"СОТРУДНИЧАЕМ": "СОТРУДНИЧАЕМ", "НА ПАУЗЕ": "НА ПАУЗЕ",
              "ПЕРЕГОВОРЫ": "ПЕРЕГОВОРЫ"}
# «ОТСТУТСТВУЕТ» (в исходнике с опечаткой) и «не обсуждали» стояли вперемешку об одном
# и том же — поверхности нет. Разными значениями это не является.
# Статусы, при которых поверхность считается взятой в работу.
WORKING_STATUSES = {"ПОДКЛЮЧЕНО", "СОГЛАСОВАНИЕ", "ПРАВКИ", "ПОДГОТОВКА"}

SURFACE_STATUS_MAP = {"ПОДКЛЮЧЕНО": "ПОДКЛЮЧЕНО", "ПОДГОТОВКА": "ПОДГОТОВКА",
                      "СОГЛАСОВАНИЕ": "СОГЛАСОВАНИЕ", "ПРАВКИ": "ПРАВКИ",
                      "ОТЛОЖЕНО": "ОТЛОЖЕНО", "ОТСТУТСТВУЕТ": "НЕТ",
                      "ОТСУТСТВУЕТ": "НЕТ", "НЕ ОБСУЖДАЛИ": "НЕТ"}


def txt(v):
    return str(v).strip() if v not in (None, "") else None


def pct(v):
    """«80 %» → 80.0. Проценты записаны текстом и с неразрывным пробелом."""
    s = txt(v)
    if not s:
        return None
    s = s.replace("\xa0", " ").replace("%", "").replace(",", ".").strip()
    try:
        return float(s)
    except ValueError:
        return None


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = txt(v)
    if not s:
        return None
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return None


def norm_inn(v):
    """ИНН из Excel приходит числом, и ведущий ноль в нём теряется: у «Дагфарм+» в
    таблице 571018796, в реестре 0571018796 — девять цифр против десяти. Поэтому
    короткое на один знак значение дополняется нулём слева, а не отбрасывается."""
    s = txt(v)
    if not s:
        return None
    digits = re.sub(r"\D", "", s.split(".")[0])
    if len(digits) in (9, 11):
        digits = digits.zfill(len(digits) + 1)
    return digits if len(digits) in (10, 12) else None


_LEGAL_FORMS = ("ооо", "оао", "зао", "пао", "ао", "ип", "нао")


def norm_legal_name(v):
    """Имя юрлица для сравнения. В таблице пишут «ООО ДПД Медиа», в реестре
    контрагентов — «ДПД МЕДИА ООО»: организационная форма стоит с разных сторон,
    поэтому она отбрасывается вместе с кавычками, а не участвует в сравнении."""
    s = txt(v)
    if not s:
        return None
    s = re.sub(r"[\"«»'`]", " ", s).lower()
    words = [w for w in re.split(r"[\s,.]+", s) if w and w not in _LEGAL_FORMS]
    return " ".join(words) or None


def norm_contract(v):
    """Номер договора для сравнения: без №, пробелов и регистра. В исходнике один и
    тот же номер встречается как «№25/08/23», «РМ 09-01-25» и «АГ-18-02-25 »."""
    s = txt(v)
    if not s:
        return None
    return re.sub(r"[\s№\"«»]", "", s).lower()


def split_people(names, emails):
    """Контакты в ячейке идут списком через «;», запятую или перенос строки, и имена с
    почтами лежат в разных колонках параллельными списками. Когда длины совпадают —
    сшиваем попарно; когда нет — не выдумываем пары, отдаём порознь."""
    def parts(v):
        s = txt(v)
        if not s:
            return []
        return [p.strip() for p in re.split(r"[;\n]|,\s(?=[А-ЯA-Z])", s) if p.strip()]

    ns, es = parts(names), parts(emails)
    if ns and es and len(ns) == len(es):
        return list(zip(ns, es))
    return [(n, None) for n in ns] + [(None, e) for e in es]


def load_rows(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = []
    for r in range(2, ws.max_row + 1):
        site = txt(ws.cell(r, COL["site"]).value)
        if not site:
            continue
        rows.append({k: ws.cell(r, c).value for k, c in COL.items()})
    return rows


def merge(group):
    """Первое непустое значение по колонке. APP-строки дублей пусты во всех полях
    поверхностей, так что порядок строк на результат не влияет."""
    out = {}
    for key in COL:
        for row in group:
            if txt(row.get(key)) is not None:
                out[key] = row[key]
                break
        else:
            out[key] = None
    return out


def add_contracts(db, publisher_id, row, contracts, domain, report):
    for key, role in [("contract", "с площадкой"), ("contract_ag", "агентский")]:
        raw = txt(row[key])
        number = norm_contract(row[key])
        if not number:
            continue
        exists = (db.query(SalesPublisherContract)
                  .filter(SalesPublisherContract.publisher_id == publisher_id,
                          SalesPublisherContract.role == role).first())
        if exists:
            continue
        cands = contracts.get(number)
        if cands and len(cands) == 1:
            db.add(SalesPublisherContract(publisher_id=publisher_id, contract_id=cands[0].id,
                                          number_raw=raw, role=role))
        else:
            db.add(SalesPublisherContract(publisher_id=publisher_id, number_raw=raw, role=role))
            report["contract_unmatched"].append((domain, raw, role, len(cands or [])))


def main(path, apply=False):
    db = SessionLocal()
    report = {"created": 0, "skipped": 0, "cp_unmatched": [], "contract_unmatched": [],
              "cp_ambiguous": [], "contacts": 0, "surfaces": 0, "connect_pct": []}

    # Индексы реестров: ИНН — надёжный ключ, имя — запасной.
    cps = db.query(Counterparty).all()
    by_inn, by_name = {}, {}
    for c in cps:
        if c.inn:
            by_inn.setdefault(re.sub(r"\D", "", c.inn), []).append(c)
        key = norm_legal_name(c.name)
        if key:
            by_name.setdefault(key, []).append(c)
    contracts = {}
    for c in db.query(Contract).all():
        key = norm_contract(c.contract_number)
        if key:
            contracts.setdefault(key, []).append(c)

    groups = {}
    for row in load_rows(path):
        groups.setdefault(normalize_domain(row["site"]), []).append(row)

    for domain, group in groups.items():
        row = merge(group)
        existing = db.query(SalesPublisher).filter(SalesPublisher.domain == domain).first()
        if existing:
            # Площадка уже заведена: до этой правки номера без пары в реестре
            # отбрасывались, поэтому добираем их отдельно.
            add_contracts(db, existing.id, row, contracts, domain, report)
            report["skipped"] += 1
            continue

        deal = txt(row["deal_type"])
        network = txt(row["network"])
        notes = [f"{label}: {txt(row[key])}" for label, key in
                 [("Комментарии", "comment"), ("Задачи", "tasks"),
                  ("Доп. комментарии", "comment2")] if txt(row[key])]
        chat = txt(row["chat"])
        connect = pct(row["connect"])
        if connect is not None:
            report["connect_pct"].append((domain, connect))

        p = SalesPublisher(
            name=txt(row["site"]), domain=domain, kind="Аптеки",
            status=STATUS_MAP.get((txt(row["status"]) or "").upper(), "ПЕРЕГОВОРЫ"),
            # «НЕЗАВИСИМЫЕ» — не сеть, а её отсутствие.
            network=None if (network or "").upper() == "НЕЗАВИСИМЫЕ" else network,
            deal_type="прямой" if (deal or "").upper() == "ПРЯМОЙ" else ("посредник" if deal else None),
            is_exclusive=(txt(row["exclusive"]) or "").upper() == "ДА",
            has_dsp=(txt(row["dsp"]) or "").upper() == "ДА",
            is_priority=(txt(row["priority"]) or "").upper() == "ДА",
            self_promo={"ДА": "ДА", "НЕТ": "НЕТ", "ЗАПРОСИТЬ": "ЗАПРОСИТЬ"}.get(
                (txt(row["self_promo"]) or "").upper()),
            self_promo_note=txt(row["self_promo_note"]),
            cpm_contract=num(row["cpm"]),
            basket_note=txt(row["basket"]),
            note="\n".join(notes) or None,
            chat_title=chat if chat and not chat.startswith("http") else None,
            chat_url=chat if chat and chat.startswith("http") else None,
            messenger_note=txt(row["messenger"]))
        db.add(p)
        db.flush()
        report["created"] += 1

        # Поверхности. Пустая строка «не обсуждали» превращается в «НЕТ», а поверхность
        # без единого заполненного поля не заводится вовсе — её просто нет.
        for kind, f, s, c, n in [("web", "figma_web", "status_web", "coverage_web", None),
                                 ("app", "figma_app", "status_app", "coverage_app", "note_app")]:
            status = SURFACE_STATUS_MAP.get((txt(row[s]) or "").upper())
            figma, cover = txt(row[f]), pct(row[c])
            note = txt(row[n]) if n else None
            if not any([figma, cover, note]) and status in (None, "НЕТ"):
                continue
            # «Работаем» выводится из статуса тем же правилом, что и в миграции v2:
            # подключённая или готовящаяся поверхность — это работа, которая уже идёт.
            # Правило продублировано намеренно: миграция правит СУЩЕСТВУЮЩИЕ строки, а на
            # чистой базе (прод) сначала идут миграции, потом импорт — и флаг остался бы
            # снятым у всех, отчего реестр показал бы нули услуг и бледные поверхности.
            db.add(SalesPublisherSurface(publisher_id=p.id, kind=kind, figma_url=figma,
                                         integration_status=status or "НЕТ",
                                         we_work=(status in WORKING_STATUSES),
                                         coverage_percent=cover, note=note))
            report["surfaces"] += 1

        # Юрлицо: по ИНН, затем по имени. Неоднозначное и ненайденное — в отчёт.
        inn, legal = norm_inn(row["inn"]), txt(row["legal"])
        found = by_inn.get(inn) if inn else None
        if not found and legal:
            found = by_name.get(norm_legal_name(legal))
        if found and len(found) == 1:
            db.add(SalesPublisherCounterparty(publisher_id=p.id, counterparty_id=found[0].id))
        elif found:
            report["cp_ambiguous"].append((domain, legal, [c.id for c in found]))
        elif legal or inn:
            report["cp_unmatched"].append((domain, legal, inn))

        # Договоры: два вида, у одиннадцати площадок заполнены оба. Ряд заводится
        # всегда — номер известен из таблицы; ссылка в реестр проставляется, только
        # если кандидат ровно один, иначе остаётся номер текстом.
        add_contracts(db, p.id, row, contracts, domain, report)

        # Контакты: телеграм-ник из «КОНТАКТНОЕ ЛИЦО» и две пары «имя/почта».
        seen = set()
        tg = txt(row["contact"])
        if tg:
            name = tg.split("-")[0].strip() if "-" in tg else None
            handle = re.search(r"@[\w\d_]+", tg)
            db.add(SalesPublisherContact(publisher_id=p.id, name=name or (None if handle else tg),
                                         telegram=handle.group(0) if handle else None,
                                         is_primary=True))
            report["contacts"] += 1
        for names, emails in [(row["contact_name"], row["contact_email"]),
                              (row["contact_name2"], row["contact_email2"])]:
            for nm, em in split_people(names, emails):
                key = ((nm or "").lower(), (em or "").lower())
                if key in seen or not any(key):
                    continue
                seen.add(key)
                db.add(SalesPublisherContact(publisher_id=p.id, name=nm, email=em))
                report["contacts"] += 1

    if apply:
        db.commit()
    else:
        db.rollback()

    print(f"{'ЗАПИСАНО' if apply else 'ПРОГОН БЕЗ ЗАПИСИ'}")
    print(f"площадок создано: {report['created']}, пропущено (уже есть): {report['skipped']}")
    print(f"поверхностей: {report['surfaces']}, контактов: {report['contacts']}")
    print(f"\nюрлица без пары в реестре контрагентов ({len(report['cp_unmatched'])}):")
    for d, legal, inn in report["cp_unmatched"]:
        print(f"   {d}: {legal or '—'} / ИНН {inn or '—'}")
    if report["cp_ambiguous"]:
        print(f"\nюрлица с несколькими кандидатами ({len(report['cp_ambiguous'])}):")
        for d, legal, ids in report["cp_ambiguous"]:
            print(f"   {d}: {legal} → {ids}")
    print(f"\nдоговоры без пары в реестре ({len(report['contract_unmatched'])}):")
    for d, number, role, cnt in report["contract_unmatched"]:
        print(f"   {d}: {number} ({role}) — кандидатов {cnt}")
    if report["connect_pct"]:
        print(f"\nколонка «ПОДКЛЮЧЕНИЕ (% из доступных мест)» не перенесена "
              f"({len(report['connect_pct'])} значений) — она не совпадает с покрытием "
              f"поверхности и своего поля пока не имеет:")
        for d, v in report["connect_pct"]:
            print(f"   {d}: {v:g} %")
    db.close()


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    main(args[0] if args else "/tmp/publishers.xlsx", apply="--apply" in sys.argv)
