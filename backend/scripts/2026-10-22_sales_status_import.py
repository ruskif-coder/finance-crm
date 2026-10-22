# -*- coding: utf-8 -*-
"""Разовый перенос встреч и комментариев из Excel продаж в «Статусы продаж» (владелец 22.10.2026).

    docker exec finance_backend python -m scripts.2026-10-22_sales_status_import /tmp/sales.xlsx \
        [--sheet "Статусы_Q4 26"] [--apply] [--replace] [--report /tmp/unmatched.xlsx]

Без `--apply` — сухой прогон. Лист — раскладки Q3/Q4 26: D — клиент, B — встреча с брендом, C — встреча с РА,
Q — комментарий Sales, R — РА, T — Sales, U — комментарий руководителя (галочка встречи → «подтверждена»).
Сопоставление — `app.sales.status_match` (написания рекламодателя, транслитерация, бренды в скобках, опечатки).
Строка со многими брендами ложится в каждый; бренд не найден — в строку рекламодателя; ничего — клиентом без сделки.
Несколько строк Excel на один бренд сливаются: комментарии через « · », встреча — подтверждённая.
`--replace` — сначала убрать прошлый перенос: заметки, которые никто не правил руками (`updated_by` пуст).
Существующие (правленные) заметки не перезаписываются. `--report` — Excel несопоставленных строк с подсказками."""
import argparse
import sys

from sqlalchemy import text

import app.model_registry  # noqa: F401
from app.database import SessionLocal
from app.sales.status_match import Directory

COLS = {"client": "D", "meet_brand": "B", "meet_agency": "C", "comment_sales": "Q", "agency": "R", "sales": "T", "comment_head": "U"}
FIELDS = ("meet_brand", "meet_agency", "comment_sales", "comment_head")


def rows_of(path: str, sheet: str):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    for r in range(5, ws.max_row + 1):
        client = ws[f"{COLS['client']}{r}"].value
        if not client or not str(client).strip():
            continue

        def txt(c):
            v = ws[f"{COLS[c]}{r}"].value
            return (str(v).strip() or None) if v not in (None, True, False) else None
        yield {"line": r, "client": " ".join(str(client).split())[:300],
               "meet_brand": "подтверждена" if ws[f"{COLS['meet_brand']}{r}"].value is True else None,
               "meet_agency": "подтверждена" if ws[f"{COLS['meet_agency']}{r}"].value is True else None,
               "comment_sales": txt("comment_sales"), "comment_head": txt("comment_head"),
               "agency": txt("agency"), "sales": txt("sales")}


def directory(db) -> tuple:
    advs = db.execute(text("SELECT id, name, short_name, name_en, name_ru FROM sales_advertisers WHERE is_active")).all()
    brands = db.execute(text("SELECT id, advertiser_id, name FROM sales_brands WHERE is_active AND advertiser_id IS NOT NULL")).all()
    names = {a[0]: (a[2] or a[1]) for a in advs}
    return Directory([(a[0], [a[1], a[2], a[3], a[4]]) for a in advs], brands), names


def merge(acc: dict, key, r: dict) -> None:
    cur = acc.setdefault(key, {f: None for f in FIELDS})
    for f in ("meet_brand", "meet_agency"):
        cur[f] = cur[f] or r[f]
    for f in ("comment_sales", "comment_head"):
        if r[f] and r[f] not in (cur[f] or ""):
            cur[f] = f"{cur[f]} · {r[f]}" if cur[f] else r[f]


def report(path: str, rows: list, names: dict, d: Directory) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Не сопоставлено"
    ws.append(["Строка Excel", "Клиент / Производитель", "Встреча с брендом", "Встреча с РА", "Комментарий Sales",
               "Комментарий руководителя", "РА", "Sales", "Похожие рекламодатели в справочнике"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["line"], r["client"], "да" if r["meet_brand"] else "", "да" if r["meet_agency"] else "", r["comment_sales"],
                   r["comment_head"], r["agency"], r["sales"], ", ".join(d.suggest(r["client"], names))])
    for col, w in zip("ABCDEFGHI", (10, 40, 12, 12, 60, 40, 16, 16, 40)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--sheet", default="Статусы_Q4 26")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--replace", action="store_true")
    ap.add_argument("--report")
    a = ap.parse_args()
    db = SessionLocal()
    try:
        d, names = directory(db)
        by_brand, by_client, unmatched = {}, {}, []
        stats = {"rows": 0, "empty": 0, "adv": 0, "brands": 0, "client": 0}
        for r in rows_of(a.path, a.sheet):
            stats["rows"] += 1
            if not any(r[f] for f in FIELDS):
                stats["empty"] += 1
                continue
            m = d.match(r["client"])
            if m is None:
                merge(by_client, r["client"].lower(), {**r})
                by_client[r["client"].lower()]["client"] = r["client"]
                unmatched.append(r)
                stats["client"] += 1
                continue
            stats["brands" if m.brand_ids else "adv"] += 1
            for bid in (m.brand_ids or [None]):
                merge(by_brand, (m.advertiser_id, bid), r)
        inserted = 0
        if a.apply:
            if a.replace:
                db.execute(text("DELETE FROM sales_status_notes WHERE updated_by IS NULL AND cardinality(from_excel) > 0"))
            for (aid, bid), v in by_brand.items():
                inserted += db.execute(text(
                    "INSERT INTO sales_status_notes (advertiser_id, brand_id, meet_brand, meet_agency, comment_sales, comment_head, from_excel) "
                    "VALUES (:a, :b, :meet_brand, :meet_agency, :comment_sales, :comment_head, :fx) "
                    "ON CONFLICT (advertiser_id, coalesce(brand_id, 0)) WHERE advertiser_id IS NOT NULL DO NOTHING"),
                    {"a": aid, "b": bid, "fx": [f for f in FIELDS if v[f]], **v}).rowcount or 0
            for v in by_client.values():
                inserted += db.execute(text(
                    "INSERT INTO sales_status_notes (client_name, meet_brand, meet_agency, comment_sales, comment_head, from_excel) "
                    "VALUES (:c, :meet_brand, :meet_agency, :comment_sales, :comment_head, :fx) "
                    "ON CONFLICT (lower(client_name)) WHERE client_name IS NOT NULL DO NOTHING"),
                    {"c": v["client"], "fx": [f for f in FIELDS if v[f]], **{f: v[f] for f in FIELDS}}).rowcount or 0
            db.commit()
        if a.report:
            report(a.report, unmatched, names, d)
        print(("ЗАПИСАНО" if a.apply else "СУХОЙ ПРОГОН") + f": строк {stats['rows']}, пустых {stats['empty']}; "
              f"сопоставлено — с брендами {stats['brands']}, только рекламодатель {stats['adv']}; не сопоставлено {stats['client']}; "
              f"заметок по брендам {len(by_brand)}, клиентов {len(by_client)}, записано новых {inserted}")
        return 0
    finally:
        db.close()


if __name__ == "__main__":
    sys.exit(main())
