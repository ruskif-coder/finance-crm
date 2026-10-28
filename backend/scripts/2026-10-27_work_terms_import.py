# -*- coding: utf-8 -*-
"""«Условия работы» из Excel владельца (27.10.2026): лист с колонками «РА · РД · Коммит Руб до НДС · Условия · ТЕМА почта ·
… · Доп комментарий» → `sk_work_terms`.

Сопоставление строгое, без угадывания (как привязка договоров к реестру): РА — агентство по любому из его названий (основное,
короткое, ENG, рус.) без учёта регистра и «ё»; «прямой» — без агентства; РД — рекламодатель по названиям, иначе бренд.
Чего не нашлось — в отчёт, строка без единого признака не пишется. Файл с данными в репозиторий не кладётся.

    docker exec finance_backend python -m scripts.2026-10-27_work_terms_import --file /tmp/terms.xlsx          # пробный
    … --apply                                                                                                     # записать
    … --alias "group4m=G4M" --alias "Оккам общие условия=OKKAM"                                                   # подсказки"""
import argparse
import re

import app.model_registry  # noqa: F401
from sqlalchemy import text

from app.agency_terms import service as S
from app.audit import log_action
from app.database import SessionLocal

DIRECT = {"прямой", "прямые", "прямая"}


def norm(s) -> str:
    return re.sub(r"[\s«»\"'.,]+", " ", str(s or "").lower().replace("ё", "е")).strip()


def index(db, table, cols):
    out = {}
    for r in db.execute(text(f"SELECT id, {', '.join(cols)} FROM {table}")).all():
        for v in r[1:]:
            if v and norm(v):
                out.setdefault(norm(v), set()).add(r[0])
    return out


def one(ix, name):
    ids = ix.get(norm(name)) or set()
    return next(iter(ids)) if len(ids) == 1 else None


def read(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).worksheets[0]
    head = [norm(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col = lambda *names: next((i for i, h in enumerate(head) if any(h.startswith(n) for n in names)), None)  # noqa: E731
    ci = {"ra": col("ра"), "rd": col("рд"), "commit": col("коммит руб"), "body": col("условия"), "mail": col("тема почта"),
          "comment": col("доп комментарий")}
    rows = []
    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        get = lambda k: row[ci[k]] if ci[k] is not None and ci[k] < len(row) else None  # noqa: E731
        if not get("body") or not str(get("body")).strip():
            continue
        rows.append({"line": n, "ra": get("ra"), "rd": get("rd"), "commit": get("commit"), "body": str(get("body")).strip(),
                     "mail": get("mail"), "comment": get("comment")})
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--from", dest="valid_from", default="2026-01-01")
    # конца срока по умолчанию нет (владелец 27.10.2026): в файле его нет, актуальные даты заполнят на вкладке «Условия работы»
    ap.add_argument("--to", dest="valid_to", default=None)
    ap.add_argument("--alias", action="append", help="«как в файле=как в справочнике»")
    ap.add_argument("--apply", action="store_true")
    a = ap.parse_args()
    alias = {norm(k): v for k, _, v in (x.partition("=") for x in (a.alias or []))}
    db = SessionLocal()
    try:
        ag = index(db, "sales_agencies", ["name", "short_name", "name_en", "name_ru"])
        adv = index(db, "sales_advertisers", ["name", "short_name", "name_en", "name_ru"])
        br = index(db, "sales_brands", ["name"])
        plan, skipped = [], []
        last_ra = None
        for r in read(a.file):
            ra = r["ra"] if r["ra"] not in (None, "") else last_ra      # пустая РА — продолжение блока агентства выше
            last_ra = ra
            feat, miss = {}, []
            if ra and norm(ra) not in DIRECT:
                aid = one(ag, alias.get(norm(ra), ra))
                (feat.__setitem__("agency_id", aid) if aid else miss.append(f"РА «{ra}»"))
            if r["rd"]:
                name = alias.get(norm(r["rd"]), r["rd"])
                vid = one(adv, name)
                bid = None if vid else one(br, name)
                if vid:
                    feat["advertiser_id"] = vid
                elif bid:
                    feat["brand_id"] = bid
                elif norm(r["rd"]) not in {"разные"}:
                    miss.append(f"РД «{r['rd']}»")
            commit = r["commit"] if isinstance(r["commit"], (int, float)) else None
            item = {**feat, "body": r["body"], "mail_subject": r["mail"], "commit_net": commit, "comment": r["comment"],
                    "valid_from": a.valid_from, "valid_to": a.valid_to}
            if not feat:
                skipped.append((r["line"], miss or ["нет признаков"], r["body"][:60]))
            else:
                plan.append((r["line"], item, miss))
        print(f"к записи: {len(plan)}, пропущено: {len(skipped)}")
        for line, item, miss in plan:
            if miss:
                print(f"  строка {line}: записывается частично, не найдено {', '.join(miss)}")
        for line, miss, body in skipped:
            print(f"  строка {line}: ПРОПУСК — {', '.join(miss)} · «{body}»")
        if not a.apply:
            print("пробный прогон — ничего не записано; запись — с --apply")
            return
        for _line, item, _miss in plan:
            w = S.save_work(db, item, None)
            db.commit()
            log_action(db, None, "agency_terms_work_save", "sk_work_term", w.id, f"Условие работы из файла: {w.body[:120]}")
        print(f"записано: {len(plan)}")
    finally:
        db.close()


if __name__ == "__main__":
    main()
